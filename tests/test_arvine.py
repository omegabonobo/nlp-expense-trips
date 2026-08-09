from __future__ import annotations

import csv
import tempfile
import unittest
from collections import Counter
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from nlp_expenses.accounting import (
    builtin_accounting_profile,
    save_default_accounting_profile,
    trip_accounting_profile,
)
from nlp_expenses.cli import main
from nlp_expenses.extraction.arvine import heuristic_parse_arvine_receipt, normalize_arvine_line_items
from nlp_expenses.generator import generate_review
from nlp_expenses.matching import match_normalized_transactions
from nlp_expenses.models import Expense, LineItem, NormalizedTransaction
from nlp_expenses.statement_normalizer import (
    normalize_statement_files,
    preflight_statement_files,
    set_statement_date_convention,
)
from nlp_expenses.trips import ensure_trip, trip_mode
from nlp_expenses.workbook import build_arvine_workbook


def write_csv(path: Path, headers: list[str], rows: list[list[object]], delimiter: str = ",") -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerow(headers)
        writer.writerows(rows)


class ArvineTests(unittest.TestCase):
    def test_openai_tax_rows_collapse_into_canonical_structured_lines(self):
        expense = Expense(
            source_file=Path("receipt.pdf"),
            expense_id="",
            expense_type="meal",
            amount=18.29,
            line_items=[
                LineItem(description="Pizza", amount=15.90),
                LineItem(description="GST/HST", amount=0.80),
                LineItem(description="QST", amount=1.59),
            ],
        )

        normalize_arvine_line_items(expense)

        self.assertEqual(expense.gst_hst, 0.80)
        self.assertEqual(expense.qst, 1.59)
        self.assertEqual(
            [(item.description, item.amount, item.line_type) for item in expense.line_items],
            [
                ("Pizza", 15.90, "purchase"),
                ("GST/HST", 0.80, "gst_hst"),
                ("QST", 1.59, "qst"),
            ],
        )

    def test_trip_mode_is_saved_and_legacy_trip_defaults_to_ivado(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            arvine = ensure_trip(root, "202607_arvine", mode="arvine")
            self.assertEqual(trip_mode(arvine), "arvine")
            legacy = root / "trips" / "202607_legacy"
            legacy.mkdir(parents=True)
            self.assertEqual(trip_mode(legacy), "ivado")

    def test_new_trips_snapshot_default_accounting_profile(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            profile = builtin_accounting_profile()
            profile["company_legal_name"] = "Example Inc."
            profile["counter_account"] = "Corporate Card Payable"
            save_default_accounting_profile(root, profile)
            first = ensure_trip(root, "202607_first", mode="arvine")
            self.assertEqual(trip_accounting_profile(root, first)["counter_account"], "Corporate Card Payable")

            changed = dict(profile)
            changed["counter_account"] = "New Default Payable"
            save_default_accounting_profile(root, changed)
            self.assertEqual(trip_accounting_profile(root, first)["counter_account"], "Corporate Card Payable")
            second = ensure_trip(root, "202608_second", mode="arvine")
            self.assertEqual(trip_accounting_profile(root, second)["counter_account"], "New Default Payable")

    def test_arvine_receipt_retains_meal_lines_and_extracts_canadian_tax_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "montreal-meal.pdf"
            path.write_bytes(b"fixture")
            text = "\n".join(
                [
                    "Bistro Exemple Montréal QC",
                    "Invoice date 2026-07-02",
                    "Dinner 100.00 CAD",
                    "Shiraz 20.00",
                    "Subtotal 100.00",
                    "GST 5.00",
                    "QST 9.98",
                    "Total CAD 134.98",
                    "GST No 123456789 RT 0001",
                    "QST No 1234567890 TQ 0001",
                ]
            )
            expense = heuristic_parse_arvine_receipt(path, text)
            self.assertEqual(expense.expense_type, "meal")
            self.assertEqual(expense.country, "Canada")
            self.assertEqual(expense.province, "QC")
            self.assertEqual(expense.gst_hst, 5.0)
            self.assertEqual(expense.qst, 9.98)
            self.assertEqual(expense.gst_hst_number, "123456789RT0001")
            self.assertEqual(expense.qst_number, "1234567890TQ0001")
            self.assertEqual(
                [(item.description, item.amount) for item in expense.line_items],
                [
                    ("Dinner", 100.0),
                    ("Shiraz", 20.0),
                    ("GST/HST", 5.0),
                    ("QST", 9.98),
                ],
            )
            self.assertTrue(all(item.included for item in expense.line_items))
            self.assertTrue(all(not item.is_alcohol for item in expense.line_items))
            self.assertTrue(
                all(not item.alcohol_reason for item in expense.line_items)
            )
            self.assertEqual(expense.corrected_amount_in_currency, expense.amount)
            self.assertEqual(expense.tax_documentation_status, "ok")

    def test_four_provider_formats_normalize_to_acceptance_counts(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            amex = folder / "amex.csv"
            write_csv(
                amex,
                ["Date", "Date Processed", "Description", "Card Member", "Account", "Amount"],
                [["2026-07-01", "2026-07-02", f"MERCHANT {i}", "A Person", "10001234", i + 1] for i in range(24)],
            )

            bmo = folder / "bmo.csv"
            with bmo.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(["Statement metadata"])
                writer.writerow(["Generated", "2026-07-10"])
                writer.writerow(["First Bank Card", "Transaction Type", "Date Posted", "Transaction Amount", "Description"])
                for i in range(29):
                    writer.writerow(["55556666", "DEBIT", "2026-07-03", -(i + 1), f"BMO PURCHASE {i}"])

            bnc = folder / "bnc.csv"
            write_csv(
                bnc,
                ["Date", "Card Number", "Description", "Category", "Debit", "Credit"],
                [["2026-07-04", "99990000", f"BNC PURCHASE {i}", "Travel", i + 1, ""] for i in range(17)],
                delimiter=";",
            )

            wise = folder / "wise.csv"
            wise_headers = [
                "ID",
                "Status",
                "Direction",
                "Created on",
                "Finished on",
                "Source amount after fees",
                "Source fee amount",
                "Source currency",
                "Target amount after fees",
                "Target currency",
                "Source name",
                "Target name",
                "Category",
            ]
            wise_rows: list[list[object]] = []
            for i in range(370):
                wise_rows.append(
                    [f"CARD_TRANSACTION-{i}", "COMPLETED", "OUT", "2026-07-05", "2026-07-05", 10, 0, "CAD", 7, "USD", "Wise", f"WISE MERCHANT {i}", "card"]
                )
            for group in range(370, 374):
                wise_rows.extend(
                    [
                        [f"CARD_TRANSACTION-{group}", "COMPLETED", "OUT", "2026-07-05", "2026-07-05", 5, 0, "CAD", 3, "USD", "Wise", f"SPLIT {group}", "card"],
                        [f"CARD_TRANSACTION-{group}", "COMPLETED", "OUT", "2026-07-05", "2026-07-05", 4, 0, "EUR", 2, "USD", "Wise", f"SPLIT {group}", "card"],
                    ]
                )
            wise_rows.append(["TRANSFER-1", "COMPLETED", "OUT", "2026-07-05", "2026-07-05", 100, 0, "CAD", 70, "USD", "Wise", "Excluded transfer", "transfer"])
            write_csv(wise, wise_headers, wise_rows)

            result = normalize_statement_files([amex, bmo, bnc, wise])
            self.assertEqual(result.errors, [])
            counts = Counter(transaction.provider for transaction in result.transactions)
            self.assertEqual(counts, {"amex": 24, "bmo": 29, "bnc": 17, "wise": 378})
            wise_groups = {transaction.transaction_group_id for transaction in result.transactions if transaction.provider == "wise"}
            self.assertEqual(len(wise_groups), 374)

    def test_amex_foreign_spend_amount_retains_purchase_currency_and_cad_settlement(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "amex.csv"
            write_csv(
                path,
                ["Date", "Date Processed", "Description", "Card Member", "Account", "Amount", "Foreign Spend Amount"],
                [["2026-07-01", "2026-07-02", "FOREIGN HOTEL", "A Person", "10001234", 130, "100 USD"]],
            )
            result = normalize_statement_files([path])
            self.assertEqual(result.errors, [])
            transaction = result.transactions[0]
            self.assertEqual(transaction.purchase_amount, 100)
            self.assertEqual(transaction.purchase_currency, "USD")
            self.assertEqual(transaction.cad_amount, 130)
            self.assertEqual(transaction.cad_completeness, "complete")

    def test_amex_two_digit_month_name_dates_are_normalized(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "amex.csv"
            write_csv(
                path,
                ["Date", "Date Processed", "Description", "Card Member", "Account #", "Amount"],
                [["28-Jul-26", "29-Jul-26", "MONTREAL HOTEL", "A Person", "-41003", 125.5]],
            )
            result = normalize_statement_files([path])
            self.assertEqual(result.errors, [])
            self.assertEqual(result.transactions[0].transaction_date, "2026-07-28")
            self.assertEqual(result.transactions[0].posted_date, "2026-07-29")

    def test_bnc_whole_row_quoted_credit_card_export_is_unwrapped(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bnc-card.csv"
            path.write_text(
                "\n".join(
                    [
                        '"Date;""card Number"";Description;Category;Debit;Credit"',
                        '"2026-07-28;""************2739"";""Yul Hurleys"";Restaurants;""19.94"";""0"""',
                        '"2026-07-14;""************2739"";""Payment received thank you!"";""Credit card payment"";""0"";""542.08"""',
                        '"2026-07-10;""************2739"";""Cashback program"";Finances;""0"";""3.88"""',
                    ]
                )
                + "\n",
                encoding="utf-8",
            )
            result = normalize_statement_files([path])
            self.assertEqual(result.errors, [])
            self.assertEqual(len(result.transactions), 3)
            self.assertEqual(
                [transaction.transaction_type for transaction in result.transactions],
                ["cashback", "payment", "purchase"],
            )
            purchase = next(
                transaction
                for transaction in result.transactions
                if transaction.transaction_type == "purchase"
            )
            self.assertEqual(purchase.account_label, "••••2739")
            self.assertEqual(purchase.cad_amount, 19.94)

    def test_bnc_debit_export_uses_bnc_adapter_and_audits_non_purchases(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bnc-debit.csv"
            path.write_text(
                "\n".join(
                    [
                        "Date;Description;Category;Debit;Credit;Balance",
                        '"2026-07-28;""Restaurant Burg"";Restaurants;""19.99"";""0"";""11782.56"""',
                        '"2026-07-24;""Fixed monthly fees"";Fees;""3.95"";""0"";""11825.43"""',
                        '"2026-07-15;""Mastercard payment"";""Credit card payment"";""542.08"";""0"";""11890.01"""',
                        '"2026-07-09;""ABM withdrawal"";Cash;""100.0"";""0"";""12432.09"""',
                    ]
                )
                + "\n",
                encoding="utf-8",
            )
            result = normalize_statement_files([path])
            self.assertEqual(result.errors, [])
            self.assertTrue(
                all(transaction.provider == "bnc" for transaction in result.transactions)
            )
            by_type = {
                transaction.transaction_type: transaction
                for transaction in result.transactions
            }
            self.assertEqual(set(by_type), {"purchase", "fee", "payment", "cash"})
            self.assertTrue(by_type["purchase"].match_eligible)
            self.assertTrue(
                all(
                    not by_type[transaction_type].match_eligible
                    for transaction_type in ("fee", "payment", "cash")
                )
            )
            self.assertEqual(by_type["purchase"].account_label, "BNC debit")

    def test_standard_csv_normalizes_exact_cad_refunds_and_audit_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "standard-statement.csv"
            write_csv(
                path,
                [
                    "transaction_date",
                    "description",
                    "purchase_amount",
                    "purchase_currency",
                    "cad_amount",
                    "transaction_type",
                    "posted_date",
                    "account",
                    "cardholder",
                    "category",
                ],
                [
                    [
                        "2026-07-02",
                        "USD HOTEL",
                        100,
                        "USD",
                        136,
                        "purchase",
                        "2026-07-03",
                        "Visa ••••1234",
                        "A Person",
                        "Hotel",
                    ],
                    [
                        "2026-07-04",
                        "MERCHANT REFUND",
                        -20,
                        "CAD",
                        "",
                        "",
                        "",
                        "Visa ••••1234",
                        "",
                        "",
                    ],
                    [
                        "2026-07-05",
                        "CARD PAYMENT",
                        500,
                        "CAD",
                        "",
                        "payment",
                        "",
                        "Visa ••••1234",
                        "",
                        "",
                    ],
                ],
            )
            result = normalize_statement_files([path])
            self.assertEqual(result.errors, [])
            by_description = {
                transaction.description: transaction
                for transaction in result.transactions
            }
            hotel = by_description["USD HOTEL"]
            self.assertEqual(hotel.provider, "standard")
            self.assertEqual(hotel.purchase_amount, 100)
            self.assertEqual(hotel.purchase_currency, "USD")
            self.assertEqual(hotel.cad_amount, 136)
            self.assertEqual(
                hotel.cad_conversion_method,
                "statement_exact_cad_settlement",
            )
            refund = by_description["MERCHANT REFUND"]
            self.assertEqual(refund.transaction_type, "refund")
            self.assertEqual(refund.purchase_amount, -20)
            self.assertEqual(refund.cad_amount, -20)
            payment = by_description["CARD PAYMENT"]
            self.assertFalse(payment.match_eligible)
            self.assertEqual(payment.cad_completeness, "not_applicable")

    def test_normalization_signs_audit_rows_and_wise_cad_completeness(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "wise.csv"
            headers = [
                "ID",
                "Status",
                "Direction",
                "Created on",
                "Finished on",
                "Source amount after fees",
                "Source fee amount",
                "Source currency",
                "Target amount after fees",
                "Target currency",
                "Source name",
                "Target name",
                "Category",
            ]
            rows = [
                ["CARD_TRANSACTION-1", "COMPLETED", "OUT", "2026-07-01", "2026-07-01", 80, 0, "CAD", 60, "USD", "Wise", "Hotel", "card"],
                ["CARD_TRANSACTION-1", "COMPLETED", "OUT", "2026-07-01", "2026-07-01", 20, 0, "EUR", 15, "USD", "Wise", "Hotel", "card"],
                ["CARD_TRANSACTION-2", "COMPLETED", "IN", "2026-07-02", "2026-07-02", 12, 0, "CAD", 9, "USD", "Wise", "Merchant refund", "card"],
                ["CARD_TRANSACTION-3", "CANCELLED", "OUT", "2026-07-03", "2026-07-03", 5, 0, "CAD", 5, "CAD", "Wise", "Cancelled", "card"],
                ["TRANSFER-4", "COMPLETED", "OUT", "2026-07-03", "2026-07-03", 5, 0, "CAD", 5, "CAD", "Wise", "Transfer", "transfer"],
            ]
            write_csv(path, headers, rows)
            result = normalize_statement_files([path])
            self.assertEqual(len(result.transactions), 4)
            split = [item for item in result.transactions if item.transaction_group_id == "CARD_TRANSACTION-1"]
            self.assertTrue(all(item.cad_completeness == "partial" for item in split))
            self.assertTrue(all(item.cad_amount is None or item.cad_amount == 80 for item in split))
            refund = next(item for item in result.transactions if item.transaction_group_id == "CARD_TRANSACTION-2")
            self.assertEqual(refund.transaction_type, "refund")
            self.assertEqual(refund.purchase_amount, -9)
            self.assertEqual(refund.cad_amount, -12)
            cancelled = next(item for item in result.transactions if item.transaction_group_id == "CARD_TRANSACTION-3")
            self.assertFalse(cancelled.match_eligible)

    def test_payments_deposits_cashback_refunds_and_generic_debit_credit_are_auditable(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            amex = folder / "amex.csv"
            write_csv(
                amex,
                ["Date", "Date Processed", "Description", "Card Member", "Account", "Amount"],
                [
                    ["2026-07-01", "2026-07-01", "PAYMENT RECEIVED", "A", "1234", -100],
                    ["2026-07-02", "2026-07-02", "MERCHANT REFUND", "A", "1234", -15],
                    ["2026-07-03", "2026-07-03", "CASHBACK REWARD", "A", "1234", -2],
                ],
            )
            bmo = folder / "bmo.csv"
            write_csv(
                bmo,
                ["First Bank Card", "Transaction Type", "Date Posted", "Transaction Amount", "Description"],
                [["5678", "CREDIT", "2026-07-04", 250, "ACCOUNT DEPOSIT"]],
            )
            generic = folder / "generic.csv"
            write_csv(
                generic,
                ["Date", "Description", "Debit", "Credit", "Currency"],
                [
                    ["2026-07-05", "GENERIC PURCHASE", 12, "", "CAD"],
                    ["2026-07-06", "GENERIC REFUND", "", 5, "CAD"],
                    ["2026-07-07", "CARD PAYMENT", "", 100, "CAD"],
                ],
            )
            result = normalize_statement_files([amex, bmo, generic])
            self.assertEqual(result.errors, [])
            by_type = {transaction.transaction_type: transaction for transaction in result.transactions}
            for nonmatch_type in ("payment", "deposit", "cashback"):
                self.assertFalse(by_type[nonmatch_type].match_eligible)
            amex_refund = next(
                transaction
                for transaction in result.transactions
                if transaction.provider == "amex" and transaction.transaction_type == "refund"
            )
            self.assertEqual(amex_refund.purchase_amount, -15)
            generic_rows = [transaction for transaction in result.transactions if transaction.provider == "generic"]
            self.assertEqual([transaction.settlement_amount for transaction in generic_rows], [12, -5, -100])
            self.assertEqual([transaction.match_eligible for transaction in generic_rows], [True, True, False])

    def test_deduplication_only_removes_stable_wise_legs(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            headers = [
                "ID", "Status", "Direction", "Created on", "Source amount after fees", "Source currency",
                "Target amount after fees", "Target currency", "Target name",
            ]
            row = ["CARD_TRANSACTION-42", "COMPLETED", "OUT", "2026-07-01", 14, "CAD", 10, "USD", "Merchant"]
            first = folder / "wise-one.csv"
            second = folder / "wise-two.csv"
            write_csv(first, headers, [row])
            write_csv(second, headers, [row])
            result = normalize_statement_files([first, second])
            self.assertEqual(len(result.transactions), 1)
            self.assertTrue(any("Removed duplicate" in warning for warning in result.warnings))

            amex = folder / "amex.csv"
            write_csv(
                amex,
                ["Date", "Date Processed", "Description", "Card Member", "Account", "Amount"],
                [
                    ["2026-07-01", "2026-07-02", "SAME MERCHANT", "A", "1234", 10],
                    ["2026-07-01", "2026-07-02", "SAME MERCHANT", "A", "1234", 10],
                ],
            )
            ambiguous = normalize_statement_files([amex])
            self.assertEqual(len(ambiguous.transactions), 2)
            self.assertTrue(all(item.normalization_status == "possible_duplicate" for item in ambiguous.transactions))

    def test_unsupported_and_malformed_statement_files_fail_visibly(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            unsupported = folder / "statement.txt"
            unsupported.write_text("not a statement", encoding="utf-8")
            malformed = folder / "malformed.csv"
            write_csv(
                malformed,
                ["Date", "Date Processed", "Description", "Card Member", "Account", "Amount"],
                [["bad", "bad", "Merchant", "A", "1234", "not-a-number"]],
            )
            reports = preflight_statement_files([unsupported, malformed])
            self.assertIn("statement.txt", reports[0].errors[0])
            self.assertIn("malformed.csv", reports[1].errors[0])
            result = normalize_statement_files([unsupported])
            self.assertTrue(result.errors)

    def test_generic_numeric_dates_require_evidence_or_saved_convention(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_dates", mode="arvine")
            folder = trip / "card_statements"

            day_first = folder / "day-first.csv"
            write_csv(
                day_first,
                ["Date", "Description", "Amount", "Currency"],
                [["13/07/2026", "Proof", 10, "CAD"], ["07/01/2026", "Ambiguous", 12, "CAD"]],
            )
            day_report = preflight_statement_files([day_first])[0]
            self.assertEqual(day_report.date_convention, "day_first")
            self.assertFalse(day_report.errors)
            day_transactions = normalize_statement_files([day_first]).transactions
            self.assertEqual(
                {transaction.transaction_date for transaction in day_transactions},
                {"2026-07-13", "2026-01-07"},
            )

            month_first = folder / "month-first.csv"
            write_csv(
                month_first,
                ["Date", "Description", "Amount", "Currency"],
                [["07/13/2026", "Proof", 10, "CAD"]],
            )
            month_report = preflight_statement_files([month_first])[0]
            self.assertEqual(month_report.date_convention, "month_first")
            self.assertIn("2026-07-13", month_report.date_samples[0])

            ambiguous = folder / "ambiguous.csv"
            write_csv(
                ambiguous,
                ["Date", "Description", "Amount", "Currency"],
                [["07/01/2026", "Ambiguous", 10, "CAD"]],
            )
            blocked = preflight_statement_files([ambiguous])[0]
            self.assertTrue(blocked.date_convention_required)
            self.assertTrue(blocked.errors)
            self.assertTrue(normalize_statement_files([ambiguous]).errors)

            set_statement_date_convention(trip, ambiguous.name, "month_first")
            selected = preflight_statement_files([ambiguous])[0]
            self.assertFalse(selected.date_convention_required)
            self.assertEqual(selected.date_convention, "month_first")
            parsed = normalize_statement_files([ambiguous]).transactions
            self.assertEqual(parsed[0].transaction_date, "2026-07-01")

    def test_matching_prefers_exact_purchase_currency_and_groups_wise_legs(self):
        expense = Expense(
            source_file=Path("hotel.pdf"),
            expense_id="EXP-1",
            date="2026-07-01",
            supplier_name="Example Hotel",
            expense_type="hotel",
            amount=75,
            currency="USD",
        )
        transactions = [
            NormalizedTransaction(
                source_file=Path("wise.csv"), source_row=2, provider="wise", transaction_group_id="GROUP-1",
                funding_leg_id="GROUP-1:a", transaction_date="2026-07-01", description="EXAMPLE HOTEL",
                match_eligible=True, purchase_amount=60, purchase_currency="USD", settlement_amount=80,
                settlement_currency="CAD", cad_amount=80, cad_completeness="partial",
            ),
            NormalizedTransaction(
                source_file=Path("wise.csv"), source_row=3, provider="wise", transaction_group_id="GROUP-1",
                funding_leg_id="GROUP-1:b", transaction_date="2026-07-01", description="EXAMPLE HOTEL",
                match_eligible=True, purchase_amount=15, purchase_currency="USD", settlement_amount=20,
                settlement_currency="EUR", cad_amount=None, cad_completeness="partial",
            ),
        ]
        match_normalized_transactions([expense], transactions)
        self.assertTrue(all(item.expense_id == "EXP-1" for item in transactions))
        self.assertTrue(all(item.match_status == "auto" for item in transactions))
        self.assertTrue(all(item.cad_completeness == "partial" for item in transactions))

    def test_matching_uses_shared_receipt_weekly_fx_cad_estimate(self):
        expense = Expense(
            source_file=Path("shared-dinner.pdf"),
            expense_id="EXP-QAR",
            date="2026-07-10",
            supplier_name="Dinner",
            expense_type="meal",
            amount=364,
            currency="QAR",
            number_of_people=2,
        )
        transaction = NormalizedTransaction(
            source_file=Path("new-card.csv"),
            source_row=8,
            provider="generic",
            account_label="Corporate ••••4412",
            transaction_group_id="GROUP-CAD",
            funding_leg_id="GROUP-CAD:1",
            transaction_date="2026-07-10",
            description="DINNER CHARGE",
            match_eligible=True,
            purchase_amount=69.80,
            purchase_currency="CAD",
            settlement_amount=69.80,
            settlement_currency="CAD",
            cad_amount=69.80,
            cad_completeness="complete",
        )

        match_normalized_transactions(
            [expense],
            [transaction],
            estimated_cad_by_expense={"shared-dinner.pdf": 69.16},
        )

        self.assertEqual(transaction.expense_id, "EXP-QAR")
        self.assertEqual(transaction.match_status, "auto")
        self.assertGreaterEqual(transaction.match_confidence, 0.72)

    def test_arvine_workbook_has_review_sheets_formulas_and_editable_matching(self):
        with tempfile.TemporaryDirectory() as tmp:
            trip = Path(tmp) / "202607_montreal"
            receipts = trip / "expenses_receipts"
            statements = trip / "card_statements"
            receipts.mkdir(parents=True)
            statements.mkdir()
            meal_file = receipts / "meal.pdf"
            travel_file = receipts / "taxi.pdf"
            statement_file = statements / "card.csv"
            meal_file.touch()
            travel_file.touch()
            statement_file.touch()
            expenses = [
                Expense(
                    source_file=meal_file, expense_id="MEAL-1", date="2026-07-01", supplier_name="Bistro",
                    description="Client dinner", expense_type="meal", amount=115, currency="CAD", country="Canada",
                    province="QC", gst_hst=5, qst=10, tax_documentation_status="ok",
                ),
                Expense(
                    source_file=travel_file, expense_id="TRAVEL-1", date="2026-07-02", supplier_name="Taxi",
                    description="Airport taxi", expense_type="transport", amount=210, currency="CAD", country="Canada",
                    province="QC", gst_hst=10, qst=0, tax_documentation_status="ok",
                ),
            ]
            transactions = [
                NormalizedTransaction(
                    source_file=statement_file, source_row=2, provider="amex", transaction_group_id="G1", funding_leg_id="G1:1",
                    transaction_date="2026-07-01", description="BISTRO", match_eligible=True, purchase_amount=115,
                    purchase_currency="CAD", settlement_amount=115, settlement_currency="CAD", cad_amount=115,
                    cad_completeness="complete",
                ),
                NormalizedTransaction(
                    source_file=statement_file, source_row=3, provider="amex", transaction_group_id="G2", funding_leg_id="G2:1",
                    transaction_date="2026-07-02", description="TAXI", match_eligible=True, purchase_amount=210,
                    purchase_currency="CAD", settlement_amount=210, settlement_currency="CAD", cad_amount=210,
                    cad_completeness="complete",
                ),
            ]
            output = build_arvine_workbook(trip, expenses, transactions)
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["expense_detail", "expense_summary", "card_statements", "expense_line_items"],
            )
            detail = workbook["expense_detail"]
            summary = workbook["expense_summary"]
            cards = workbook["card_statements"]
            self.assertEqual(detail["P2"].value, 0.5)
            self.assertEqual(detail["Q2"].value, 0.5)
            self.assertEqual(detail["P3"].value, 1.0)
            self.assertIn("expense_line_items", detail["AL2"].value)
            self.assertIn("COUNTIFS(expense_line_items!", detail["AO2"].value)
            self.assertIn("$AM2/$AL2", detail["AO2"].value)
            self.assertIn("COUNTIFS(expense_line_items!", detail["AO3"].value)
            self.assertIn("card_statements", detail["R2"].value)
            self.assertEqual(detail["AQ1"].value, "statement_purchase_amount")
            self.assertEqual(detail["AR1"].value, "statement_purchase_currency")
            self.assertIn("'card_statements'!$N$2", detail["AQ2"].value)
            self.assertEqual(detail["AS1"].value, "accounting_original_basis")
            self.assertEqual(detail["AT1"].value, "accounting_basis_status")
            self.assertIn("$R2/$AS2", detail["U2"].value)
            self.assertEqual(detail["AF2"].hyperlink.target, "expenses_receipts/meal.pdf")
            self.assertEqual([summary.cell(row, 4).value for row in range(10, 15)], [
                "Travel – Non-meal",
                "Meals – Deductible (50%)",
                "Meals – Non-deductible (50%)",
                "GST/HST Receivable",
                "QST Receivable",
            ])
            self.assertEqual(summary["O15"].value, "=SUM(O10:O14)")
            self.assertIn("manual override", summary["A24"].value)
            self.assertTrue(cards.protection.sheet)
            self.assertFalse(cards["T2"].protection.locked)
            self.assertFalse(cards["V2"].protection.locked)
            self.assertEqual(cards["X2"].hyperlink.target, "card_statements/card.csv")
            # Representative journal arithmetic: 200 travel + 53.75 + 53.75 meals + 12.50 GST + 5 QST.
            self.assertAlmostEqual(200 + 53.75 + 53.75 + 12.50 + 5, 325.0)

    def test_accounting_profile_controls_tax_recovery_deduction_and_accounts(self):
        with tempfile.TemporaryDirectory() as tmp:
            trip = Path(tmp) / "202607_profile"
            receipts = trip / "expenses_receipts"
            statements = trip / "card_statements"
            receipts.mkdir(parents=True)
            statements.mkdir()
            receipt = receipts / "meal.pdf"
            receipt.touch()
            expense = Expense(
                source_file=receipt,
                expense_id="MEAL-1",
                date="2026-07-01",
                supplier_name="Bistro",
                description="Client dinner",
                expense_type="meal",
                amount=115,
                currency="CAD",
                country="Canada",
                province="QC",
                gst_hst=5,
                qst=10,
                tax_documentation_status="ok",
            )
            profile = builtin_accounting_profile()
            profile.update(
                {
                    "company_legal_name": "Arvine Corp.",
                    "traveller_reimbursement_type": "corporate_card",
                    "counter_account": "Corporate Card Payable",
                    "gst_hst_registrant": False,
                    "qst_registrant": False,
                    "meal_deduction_pct": 0.75,
                    "meal_tax_recovery_pct": 0.8,
                }
            )
            profile["account_mapping"] = {
                **profile["account_mapping"],
                "meal_deductible": "Meals Custom Deductible",
            }
            output = build_arvine_workbook(trip, [expense], [], accounting_profile=profile)
            workbook = load_workbook(output, data_only=False)
            detail = workbook["expense_detail"]
            summary = workbook["expense_summary"]
            self.assertEqual(detail["P2"].value, 0.75)
            self.assertEqual(detail["Q2"].value, 0)
            self.assertIn("*0.000000", detail["Y2"].value)
            self.assertIn("*0.000000", detail["Z2"].value)
            self.assertEqual(summary["B3"].value, "Arvine Corp.")
            self.assertEqual(summary["B5"].value, "Corporate Card Payable")
            self.assertEqual(summary["D11"].value, "Meals Custom Deductible")
            self.assertEqual(summary["T4"].value, "corporate_card")
            self.assertEqual(summary["C22"].value, '=IF(ABS(B22)<0.01,"OK","REVIEW")')

    def test_statement_confirmation_no_preserves_existing_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_trip", mode="arvine")
            statement = trip / "card_statements" / "amex.csv"
            write_csv(statement, ["Date", "Date Processed", "Description", "Card Member", "Account", "Amount"], [])
            output = trip / f"expense_review_{trip.name}.xlsx"
            output.write_bytes(b"existing")
            result = generate_review(trip, root, llm_mode="off", confirm_input=lambda _prompt: "n")
            self.assertIsNone(result)
            self.assertEqual(output.read_bytes(), b"existing")

    def test_noninteractive_cli_requires_statement_confirmation_flag(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_trip", mode="arvine")
            with patch("nlp_expenses.cli.Path.cwd", return_value=root), patch("nlp_expenses.cli.sys.stdin.isatty", return_value=False):
                with self.assertRaisesRegex(SystemExit, "--statements-complete"):
                    main(["generate", str(trip), "--llm", "off"])


if __name__ == "__main__":
    unittest.main()
