from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from nlp_expenses.accounting import builtin_accounting_profile
from nlp_expenses.consolidation import (
    calculate_accounting_summary,
    calculate_expense_result,
    consolidation_view,
    finalize_consolidation,
)
from nlp_expenses.generator import generate_review
from nlp_expenses.line_items import (
    add_line_item,
    line_item_review_view,
    remove_line_item,
    save_line_item_review,
    set_expense_review,
    set_line_item_review,
)
from nlp_expenses.models import Expense, LineItem
from nlp_expenses.reconciliation import (
    set_manual_match,
    set_statement_basis,
    set_transaction_decision,
    sync_reconciliation,
)
from nlp_expenses.trip_manifest import receipt_manifest_record
from nlp_expenses.trip_metadata import save_trip_metadata
from nlp_expenses.trips import ensure_trip, trip_mode


def complete_metadata(trip: Path) -> None:
    save_trip_metadata(
        trip,
        {
            "claim_program": "ivado_sponsored" if trip_mode(trip) == "ivado" else "arvine_only",
            "traveller": "Florent",
            "company": "Example Corp.",
            "sponsor": "IVADO Labs" if trip_mode(trip) == "ivado" else "",
            "start_date": "2026-07-01",
            "end_date": "2026-07-02",
            "business_purpose": "Client workshop",
            "approver": "Manager",
            "payment_method": "Personal card reimbursement",
            "default_paid_by": "employee_personal",
            "payer_confirmed": True,
        },
    )


def meal_expense(receipt: Path, currency: str = "CAD", amount: float = 120.0) -> Expense:
    return Expense(
        source_file=receipt,
        expense_id="",
        date="2026-07-01",
        supplier_name="Bistro",
        expense_type="meal",
        amount=amount,
        currency=currency,
        line_items=[
            LineItem(description="Dinner", amount=100.0),
            LineItem(
                description="Shiraz",
                amount=20.0,
                is_alcohol=True,
                alcohol_confidence=0.97,
                alcohol_reason="recognized wine",
            ),
        ],
    )


class ConsolidationTests(unittest.TestCase):
    def test_existing_melbourne_fixture_uses_reviewed_share_and_exact_cad(self):
        root = Path(__file__).resolve().parents[1]
        trip = root / "trips" / "202607_melb-flo-test"
        if not trip.exists():
            self.skipTest("Melbourne reconciliation fixture is not present")

        preview = consolidation_view(root, trip)
        by_source = {expense["source_file"]: expense for expense in preview["expenses"]}

        self.assertEqual(by_source["Garuda - DPS to MEL.pdf"]["ivado_claimable_cad"], 1_659.90)
        self.assertEqual(
            by_source["Scanned_20260530 - Receipt - May 30 2026 - 9-27 PM.pdf"][
                "ivado_claimable_cad"
            ],
            83.87,
        )
        self.assertEqual(
            by_source["Scanned_20260531 - Receipt - May 31 2026 - 8-42 PM.pdf"][
                "ivado_claimable_cad"
            ],
            106.83,
        )
        self.assertEqual(
            by_source["Scanned_20260604-2259.pdf"]["ivado_claimable_cad"],
            154.00,
        )
        # The exact statement charge for this non-shared receipt is 19.86,
        # while the reference CSV contains 19.85. Keeping exact statement CAD
        # authoritative makes the app total one cent above the CSV row sum.
        self.assertEqual(
            by_source["Scanned_20260603-1216.pdf"]["ivado_claimable_cad"],
            19.86,
        )
        self.assertEqual(preview["summary"]["ivado_claim_total_cad"], 7_532.51)
        self.assertEqual(preview["summary"]["employee_reimbursement_total_cad"], 7_532.51)

    def test_accounting_rounding_residual_is_identified_and_balanced(self):
        expense = {
            "expense_type": "meal",
            "currency": "CAD",
            "included_in_arvine": True,
            "arvine_reimbursable_cad": 1.01,
            "arvine_claimable_ratio": 1.0,
            "fx_rate": 1.0,
            "gst_hst": 0.0,
            "qst": 0.0,
        }

        accounting = calculate_accounting_summary(
            [expense],
            builtin_accounting_profile(),
            "company",
        )

        self.assertEqual(accounting["pre_adjustment_difference"], 0.01)
        self.assertEqual(accounting["rounding_adjustment_cad"], -0.01)
        self.assertEqual(
            accounting["rounding_adjustment_account"],
            "Meals – Non-deductible (50%)",
        )
        self.assertEqual(accounting["journal_total"], 1.01)
        self.assertEqual(accounting["balance_difference"], 0.0)

    def test_app_review_controls_people_lines_totals_and_finalization(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_app-first", mode="ivado")
            complete_metadata(trip)
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"fixture")
            save_line_item_review(trip, [meal_expense(receipt)])

            set_expense_review(
                trip,
                receipt.name,
                {"number_of_people": 2, "included": True, "business_purpose": "Client dinner"},
            )
            result = consolidation_view(root, trip)
            self.assertTrue(result["is_ready"])
            self.assertAlmostEqual(result["expenses"][0]["claimable_cad"], 50.0)
            self.assertEqual(result["expenses"][0]["number_of_people"], 2)

            finalized = finalize_consolidation(root, trip)
            self.assertTrue(finalized["current"])
            cocktail = next(
                item
                for item in line_item_review_view(trip)["receipts"][0]["line_items"]
                if item["description"] == "Shiraz"
            )
            set_line_item_review(
                trip,
                receipt.name,
                cocktail["line_id"],
                {"is_alcohol": False},
            )
            changed = consolidation_view(root, trip)
            self.assertFalse(changed["is_finalized"])
            self.assertAlmostEqual(changed["summary"]["claimable_cad"], 60.0)

    def test_manual_line_add_remove_and_automatic_reset_history(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_lines", mode="arvine")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"fixture")
            save_line_item_review(trip, [meal_expense(receipt)])

            added = add_line_item(trip, receipt.name, "Tip", 15, included=True)
            manual = next(
                item for item in added["receipts"][0]["line_items"] if item["description"] == "Tip"
            )
            self.assertTrue(manual["manual"])
            removed = remove_line_item(trip, receipt.name, manual["line_id"])
            self.assertNotIn(
                "Tip",
                [item["description"] for item in removed["receipts"][0]["line_items"]],
            )

    def test_ivado_sync_mapping_fx_preview_and_workbook_reuse_the_same_decision(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_ivado-fx", mode="ivado")
            complete_metadata(trip)
            receipt = trip / "expenses_receipts" / "bistro.pdf"
            receipt.write_bytes(b"fixture")
            statement = trip / "card_statements" / "card.csv"
            statement.write_text(
                "Date,Description,Amount,Foreign Spend Amount\n2026-07-01,BISTRO,75,50 AUD\n",
                encoding="utf-8",
            )
            expense = meal_expense(receipt, currency="AUD", amount=50)
            expense.line_items = [LineItem(description="Dinner", amount=50)]
            save_line_item_review(trip, [expense])

            with patch("nlp_expenses.generator.parse_receipt", return_value=expense):
                reconciliation = sync_reconciliation(trip, root, llm_mode="off")
            transaction = reconciliation["transactions"][0]
            self.assertEqual(transaction["expense_file"], receipt.name)
            preview = consolidation_view(root, trip)
            self.assertEqual(preview["expenses"][0]["cad_amount_used"], 75.0)
            self.assertEqual(preview["expenses"][0]["fx_rate"], 1.5)
            self.assertTrue(preview["is_ready"])

            set_manual_match(
                trip,
                transaction["group_id"],
                expense_file=receipt.name,
                use_auto=False,
            )
            with patch("nlp_expenses.generator.parse_receipt", return_value=expense):
                workbook_path = generate_review(trip, root, llm_mode="off")
            workbook = load_workbook(workbook_path, data_only=False)
            self.assertEqual(workbook["card_statements"]["D2"].value, "20260701_#1")
            self.assertEqual(workbook["expense_list"]["I2"].value, 1)
            self.assertEqual(workbook["expense_list"]["Q2"].value, True)

            set_manual_match(
                trip,
                transaction["group_id"],
                expense_file=None,
                use_auto=False,
            )
            unresolved = consolidation_view(root, trip)
            self.assertGreater(unresolved["summary"]["blocking_count"], 0)
            self.assertFalse(
                any(issue["kind"] == "statement_mapping" for issue in unresolved["issues"])
            )
            self.assertTrue(any(issue["kind"] == "cad_amount" for issue in unresolved["issues"]))
            self.assertEqual(len(unresolved["expenses"]), 1)

    def test_statement_purchase_amount_prevents_double_dividing_a_shared_meal(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_shared-meal", mode="ivado")
            complete_metadata(trip)
            receipt = trip / "expenses_receipts" / "shared.pdf"
            receipt.write_bytes(b"fixture")
            statement = trip / "card_statements" / "card.csv"
            statement.write_text(
                "Date,Description,Amount,Foreign Spend Amount\n2026-07-01,BISTRO,110,100 AUD\n",
                encoding="utf-8",
            )
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Bistro",
                expense_type="meal",
                amount=300,
                currency="AUD",
                line_items=[
                    LineItem(description="Dinner", amount=240),
                    LineItem(
                        description="Wine",
                        amount=60,
                        is_alcohol=True,
                        alcohol_confidence=0.99,
                    ),
                ],
            )
            save_line_item_review(trip, [expense])
            set_expense_review(trip, receipt.name, {"number_of_people": 3})

            with patch("nlp_expenses.generator.parse_receipt", return_value=expense):
                reconciliation = sync_reconciliation(trip, root, llm_mode="off")
            transaction = reconciliation["transactions"][0]
            set_manual_match(trip, transaction["group_id"], receipt.name, use_auto=False)

            preview = consolidation_view(root, trip)
            result = preview["expenses"][0]
            self.assertEqual(result["arvine_claimable_original"], 100.0)
            self.assertEqual(result["ivado_claimable_original"], 80.0)
            self.assertEqual(result["statement_purchase_amount_used"], 100.0)
            self.assertEqual(result["statement_purchase_currency"], "AUD")
            self.assertEqual(result["fx_rate"], 1.1)
            self.assertEqual(result["claimable_cad"], 88.0)

            with patch("nlp_expenses.generator.parse_receipt", return_value=expense):
                workbook_path = generate_review(trip, root, llm_mode="off")
            workbook = load_workbook(workbook_path, data_only=False)
            detail = workbook["expense_list"]
            self.assertIn("card_statements!$H:$H", detail["T2"].value)
            self.assertIn("$K2/$V2", detail["L2"].value)
            self.assertIn("$T2-$F2/$I2", detail["V2"].value)
            self.assertEqual(detail["W1"].value, "accounting_basis_status")
            self.assertIn('"statement_person_share"', detail["M2"].value)
            self.assertIn("$K2*", detail["M2"].value)

    def test_explicit_personal_share_statement_prevents_double_division(self):
        receipt = {
            "source_file": "shared.pdf",
            "expense_type": "meal",
            "amount": 300.0,
            "currency": "AUD",
            "included": True,
            "included_in_ivado": True,
            "number_of_people": 3,
            "line_total": 300.0,
            "arvine_included_total": 300.0,
            "ivado_included_total": 300.0,
            "arvine_excluded_total": 0.0,
            "ivado_excluded_total": 0.0,
        }
        reconciled = {
            "cad_amount_used": 110.0,
            "cad_source": "statement",
            "statement_amount_basis": "personal_share",
            "statement_amount_basis_explicit": True,
        }

        result = calculate_expense_result(receipt, reconciled, "ivado_reimbursed")

        self.assertEqual(result["arvine_reimbursable_cad"], 110.0)
        self.assertEqual(result["ivado_claimable_cad"], 110.0)
        self.assertEqual(result["statement_amount_basis"], "personal_share")

    def test_personal_share_alcohol_exclusion_uses_reviewed_line_ratio(self):
        receipt = {
            "source_file": "shared-dinner.pdf",
            "expense_type": "meal",
            "amount": 300.0,
            "currency": "AUD",
            "included": True,
            "included_in_ivado": True,
            "number_of_people": 3,
            "line_total": 300.0,
            "arvine_included_total": 300.0,
            "ivado_included_total": 240.0,
            "arvine_excluded_total": 0.0,
            "ivado_excluded_total": 60.0,
        }
        reconciled = {
            "cad_amount_used": 110.0,
            "cad_source": "statement",
            "statement_amount_basis": "personal_share",
            "statement_amount_basis_explicit": True,
        }

        result = calculate_expense_result(receipt, reconciled, "ivado_reimbursed")

        self.assertEqual(result["arvine_reimbursable_cad"], 88.0)
        self.assertEqual(result["ivado_claimable_cad"], 88.0)
        self.assertEqual(result["ivado_excluded_cad"], 22.0)

    def test_exact_idr_statement_cad_is_not_rebuilt_from_rounded_fx(self):
        receipt = {
            "source_file": "flight.pdf",
            "expense_type": "flight",
            "amount": 20_598_982.0,
            "currency": "IDR",
            "included": True,
            "included_in_ivado": True,
            "number_of_people": 1,
        }
        reconciled = {
            "cad_amount_used": 1_659.90,
            "cad_source": "statement",
            "statement_purchase_amount_used": 1_659.90,
            "statement_purchase_currency": "CAD",
            "fx_basis_amount_used": 20_598_982.0,
            "fx_basis_status": "receipt_fallback_currency",
        }

        result = calculate_expense_result(receipt, reconciled, "ivado_reimbursed")

        self.assertEqual(result["fx_rate"], 0.000081)
        self.assertEqual(result["arvine_reimbursable_cad"], 1_659.90)
        self.assertEqual(result["ivado_claimable_cad"], 1_659.90)
        manifest = receipt_manifest_record(Path("trip"), result, "ivado_sponsored")
        self.assertEqual(manifest["fx_rate"], 0.000081)

    def test_manual_cad_override_remains_authoritative_over_statement_cad(self):
        receipt = {
            "source_file": "flight.pdf",
            "expense_type": "flight",
            "amount": 20_598_982.0,
            "currency": "IDR",
            "included": True,
            "included_in_ivado": True,
            "number_of_people": 1,
            "manual_cad_override": 1_600.0,
            "manual_cad_note": "Card provider correction",
        }
        reconciled = {
            "cad_amount_used": 1_659.90,
            "cad_source": "statement",
            "statement_purchase_amount_used": 1_659.90,
            "statement_purchase_currency": "CAD",
        }

        result = calculate_expense_result(receipt, reconciled, "ivado_reimbursed")

        self.assertEqual(result["cad_source"], "manual")
        self.assertEqual(result["arvine_reimbursable_cad"], 1_600.0)
        self.assertEqual(result["ivado_claimable_cad"], 1_600.0)

    def test_statement_basis_decision_persists_without_resync(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_shared-cad-only", mode="ivado")
            complete_metadata(trip)
            receipt = trip / "expenses_receipts" / "shared.pdf"
            receipt.write_bytes(b"fixture")
            (trip / "card_statements" / "card.csv").write_text(
                "Date,Description,Amount,Foreign Spend Amount\n2026-07-01,BISTRO,110,\n",
                encoding="utf-8",
            )
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Bistro",
                expense_type="meal",
                amount=300,
                currency="AUD",
                line_items=[
                    LineItem(description="Dinner", amount=240),
                    LineItem(description="Wine", amount=60, is_alcohol=True),
                ],
            )
            save_line_item_review(trip, [expense])
            set_expense_review(trip, receipt.name, {"number_of_people": 3})
            reconciliation = sync_reconciliation(trip, root, llm_mode="off")
            transaction = reconciliation["transactions"][0]
            set_manual_match(trip, transaction["group_id"], receipt.name, use_auto=False)

            set_statement_basis(trip, receipt.name, "personal_share")
            preview = consolidation_view(root, trip)
            result = preview["expenses"][0]

            self.assertFalse(preview["issues"])
            self.assertEqual(result["fx_basis_status"], "statement_personal_share_explicit")
            self.assertEqual(result["fx_rate"], 1.1)
            self.assertEqual(result["ivado_claimable_cad"], 88.0)
            self.assertTrue(result["statement_amount_basis_explicit"])

    def test_malformed_statement_purchase_amount_falls_back_to_receipt_fx_basis(self):
        receipt = {
            "source_file": "flight.pdf",
            "amount": 3_649.43,
            "currency": "AUD",
            "included": True,
            "number_of_people": 1,
        }
        reconciled = {
            "cad_amount_used": 3_688.57,
            "cad_source": "statement",
            "statement_purchase_amount_used": 649.43,
            "statement_purchase_currency": "AUD",
        }

        result = calculate_expense_result(receipt, reconciled)

        self.assertEqual(result["fx_basis_amount_used"], 3_649.43)
        self.assertEqual(result["fx_basis_status"], "receipt_fallback_mismatch")
        self.assertEqual(result["fx_rate"], 1.010725)
        self.assertEqual(result["claimable_cad"], 3_688.57)

    def test_full_receipt_match_uses_exact_statement_cad_settlement(self):
        receipt = {
            "source_file": "breakfast.pdf",
            "expense_type": "meal",
            "amount": 11.50,
            "currency": "AUD",
            "included": True,
            "number_of_people": 1,
            "line_total": 11.50,
            "included_total": 11.50,
            "excluded_total": 0.0,
        }
        reconciled = {
            "cad_amount_used": 11.93,
            "cad_source": "statement",
            "statement_purchase_amount_used": 11.72,
            "statement_purchase_currency": "AUD",
        }

        result = calculate_expense_result(receipt, reconciled)

        self.assertEqual(result["fx_basis_status"], "statement_receipt_total")
        self.assertEqual(result["claimable_cad"], 11.93)

    def test_card_tip_gap_uses_full_settlement_without_inflating_receipt_taxes(self):
        receipt = {
            "source_file": "dinner.pdf",
            "expense_type": "meal",
            "amount": 57.0,
            "currency": "CAD",
            "included": True,
            "number_of_people": 1,
            "gst_hst": 2.48,
            "qst": 4.95,
            "line_total": 57.0,
            "included_total": 57.0,
            "excluded_total": 0.0,
        }
        reconciled = {
            "cad_amount_used": 67.0,
            "cad_source": "statement",
            "statement_purchase_amount_used": 67.0,
            "statement_purchase_currency": "CAD",
            "statement_receipt_difference": 10.0,
            "fx_basis_amount_used": 67.0,
            "fx_basis_status": "statement_includes_tip",
        }

        result = calculate_expense_result(receipt, reconciled)

        self.assertEqual(result["claimable_cad"], 67.0)
        self.assertEqual(result["arvine_reimbursable_cad"], 67.0)
        self.assertEqual(result["fx_rate"], 1.0)
        self.assertEqual(result["fx_basis_status"], "statement_includes_tip")
        self.assertEqual(result["statement_receipt_difference"], 10.0)

    def test_non_meal_receipt_correction_is_not_overridden_by_stale_synthetic_line(self):
        receipt = {
            "source_file": "flight.pdf",
            "expense_type": "flight",
            "amount": 3_649.43,
            "currency": "AUD",
            "included": True,
            "number_of_people": 1,
            "line_total": 3_579.43,
            "included_total": 3_579.43,
            "excluded_total": 0.0,
        }
        reconciled = {
            "cad_amount_used": 3_688.57,
            "cad_source": "statement",
            "statement_purchase_amount_used": 649.43,
            "statement_purchase_currency": "AUD",
        }

        result = calculate_expense_result(receipt, reconciled)

        self.assertEqual(result["claimable_original"], 3_649.43)
        self.assertEqual(result["claimable_cad"], 3_688.57)

    def test_unrelated_statement_transaction_can_be_excluded_and_restored(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_statement-disposition", mode="ivado")
            complete_metadata(trip)
            receipt = trip / "expenses_receipts" / "taxi.pdf"
            receipt.write_bytes(b"fixture")
            statement = trip / "card_statements" / "card.csv"
            statement.write_text(
                "Date,Description,Amount,Foreign Spend Amount\n"
                "2026-07-01,TAXI,75,50 AUD\n"
                "2026-07-01,SPOTIFY,20,\n",
                encoding="utf-8",
            )
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Taxi",
                expense_type="transport",
                amount=50,
                currency="AUD",
                line_items=[LineItem(description="Taxi", amount=50)],
            )
            save_line_item_review(trip, [expense])
            with patch("nlp_expenses.generator.parse_receipt", return_value=expense):
                reconciliation = sync_reconciliation(trip, root, llm_mode="off")
            unrelated = next(
                item for item in reconciliation["transactions"] if "SPOTIFY" in item["description"]
            )

            excluded = set_transaction_decision(
                trip,
                unrelated["group_id"],
                "ignore",
                "Personal subscription outside the trip",
            )
            excluded_transaction = next(
                item
                for item in excluded["transactions"]
                if item["group_id"] == unrelated["group_id"]
            )
            self.assertTrue(excluded_transaction["ignored"])
            self.assertEqual(excluded["summary"]["needs_review_count"], 0)
            self.assertTrue(consolidation_view(root, trip)["is_ready"])

            restored = set_transaction_decision(trip, unrelated["group_id"], "keep")
            restored_transaction = next(
                item
                for item in restored["transactions"]
                if item["group_id"] == unrelated["group_id"]
            )
            self.assertFalse(restored_transaction["ignored"])
            self.assertEqual(restored_transaction["normalization_status"], "ok")
            self.assertEqual(restored["summary"]["needs_review_count"], 0)
            self.assertIn(
                restored_transaction["group_id"],
                {item["group_id"] for item in restored["unmatched_transactions"]},
            )

    def test_manual_cad_override_is_shared_without_a_statement(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_manual-cad", mode="ivado")
            complete_metadata(trip)
            receipt = trip / "expenses_receipts" / "taxi.pdf"
            receipt.write_bytes(b"fixture")
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Taxi",
                expense_type="transport",
                amount=50,
                currency="AUD",
                line_items=[LineItem(description="Taxi", amount=50)],
            )
            save_line_item_review(trip, [expense])
            set_expense_review(
                trip,
                receipt.name,
                {
                    "manual_cad_override": 45,
                    "manual_cad_note": "Card portal confirmation",
                },
            )
            preview = consolidation_view(root, trip)
            self.assertTrue(preview["is_ready"])
            self.assertEqual(preview["expenses"][0]["cad_source"], "manual")
            self.assertEqual(preview["expenses"][0]["fx_rate"], 0.9)

    def test_legacy_line_review_defaults_to_included_one_person_and_existing_total(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_legacy", mode="ivado")
            receipt = trip / "expenses_receipts" / "taxi.pdf"
            receipt.write_bytes(b"fixture")
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Taxi",
                expense_type="transport",
                amount=35,
                currency="CAD",
                line_items=[LineItem(description="Taxi", amount=35)],
            )
            save_line_item_review(trip, [expense])
            path = trip / ".nlp-expenses-line-items.json"
            state = json.loads(path.read_text(encoding="utf-8"))
            legacy = state["receipts"][0]
            for field in (
                "amount",
                "included",
                "number_of_people",
                "manual_cad_override",
                "manual_cad_note",
                "extracted",
                "field_overrides",
                "extracted_line_items",
            ):
                legacy.pop(field, None)
            path.write_text(json.dumps(state), encoding="utf-8")

            migrated = line_item_review_view(trip)["receipts"][0]
            self.assertEqual(migrated["amount"], 35)
            self.assertTrue(migrated["included"])
            self.assertEqual(migrated["number_of_people"], 1)
            self.assertEqual(migrated["field_issues"], [])


if __name__ == "__main__":
    unittest.main()
