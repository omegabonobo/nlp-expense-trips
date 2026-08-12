from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from nlp_expenses.models import Expense, NormalizedTransaction
from nlp_expenses.reconciliation import (
    aggregate_transaction_groups,
    apply_allocations_to_groups,
    reconciliation_input_fingerprint,
    reconciliation_view,
    serialize_expense,
)
from nlp_expenses.trip_metadata import save_trip_metadata
from nlp_expenses.trips import ensure_trip
from nlp_expenses.ui import create_app
from nlp_expenses.workbook import build_arvine_workbook

EXPECTED_ROUTES = [
    ("/", ("GET",)),
    ("/api/jobs/<job_id>", ("GET",)),
    ("/api/settings/openai", ("POST",)),
    ("/api/trips", ("POST",)),
    ("/api/trips/<trip_name>", ("DELETE",)),
    ("/api/trips/<trip_name>", ("GET",)),
    ("/api/trips/<trip_name>/accounting-profile", ("POST",)),
    ("/api/trips/<trip_name>/approve", ("POST",)),
    ("/api/trips/<trip_name>/archive", ("POST",)),
    ("/api/trips/<trip_name>/claim-program", ("POST",)),
    ("/api/trips/<trip_name>/consolidation", ("GET",)),
    ("/api/trips/<trip_name>/download-manifest", ("GET",)),
    ("/api/trips/<trip_name>/download-package", ("GET",)),
    ("/api/trips/<trip_name>/download-workbook", ("GET",)),
    ("/api/trips/<trip_name>/export-package", ("POST",)),
    ("/api/trips/<trip_name>/file-state", ("GET",)),
    ("/api/trips/<trip_name>/finalize", ("POST",)),
    ("/api/trips/<trip_name>/generate", ("POST",)),
    ("/api/trips/<trip_name>/line-items/add", ("POST",)),
    ("/api/trips/<trip_name>/line-items/expense", ("POST",)),
    ("/api/trips/<trip_name>/line-items/item", ("POST",)),
    ("/api/trips/<trip_name>/line-items/remove", ("POST",)),
    ("/api/trips/<trip_name>/line-items/reset", ("POST",)),
    ("/api/trips/<trip_name>/line-items/sync", ("POST",)),
    ("/api/trips/<trip_name>/metadata", ("POST",)),
    ("/api/trips/<trip_name>/mode", ("POST",)),
    ("/api/trips/<trip_name>/open-workbook", ("POST",)),
    ("/api/trips/<trip_name>/policy-exception", ("POST",)),
    ("/api/trips/<trip_name>/receipt", ("GET",)),
    ("/api/trips/<trip_name>/reconcile", ("POST",)),
    ("/api/trips/<trip_name>/reconciliation", ("GET",)),
    ("/api/trips/<trip_name>/reconciliation/allocations", ("POST",)),
    ("/api/trips/<trip_name>/reconciliation/coverage-settings", ("POST",)),
    ("/api/trips/<trip_name>/reconciliation/invoice", ("POST",)),
    ("/api/trips/<trip_name>/reconciliation/mapping", ("POST",)),
    ("/api/trips/<trip_name>/reconciliation/transaction-decision", ("POST",)),
    ("/api/trips/<trip_name>/remove-file", ("POST",)),
    ("/api/trips/<trip_name>/reveal", ("POST",)),
    ("/api/trips/<trip_name>/statement-date-convention", ("POST",)),
    ("/api/trips/<trip_name>/upload/<kind>", ("POST",)),
    ("/static/<path:filename>", ("GET",)),
]


def route_contract(app) -> list[tuple[str, tuple[str, ...]]]:
    return sorted(
        (
            rule.rule,
            tuple(sorted(rule.methods - {"HEAD", "OPTIONS"})),
        )
        for rule in app.url_map.iter_rules()
    )


def serializable_cell_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def workbook_semantic_contract(path: Path) -> dict:
    """Capture user-visible workbook behavior while ignoring ZIP metadata."""

    workbook = load_workbook(path, data_only=False)
    sheets = []
    for worksheet in workbook.worksheets:
        cells = []
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None and not cell.hyperlink and not cell.has_style:
                    continue
                cells.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": serializable_cell_value(cell.value),
                        "number_format": cell.number_format,
                        "style_id": cell.style_id,
                        "locked": cell.protection.locked,
                        "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                    }
                )
        validations = [
            {
                "ranges": str(validation.sqref),
                "type": validation.type,
                "formula1": validation.formula1,
                "formula2": validation.formula2,
                "allow_blank": validation.allow_blank,
            }
            for validation in worksheet.data_validations.dataValidation
        ]
        sheets.append(
            {
                "title": worksheet.title,
                "state": worksheet.sheet_state,
                "freeze_panes": str(worksheet.freeze_panes or ""),
                "auto_filter": str(worksheet.auto_filter.ref or ""),
                "protected": worksheet.protection.sheet,
                "merged_cells": sorted(str(item) for item in worksheet.merged_cells.ranges),
                "validations": validations,
                "cells": cells,
            }
        )
    return {
        "sheet_order": workbook.sheetnames,
        "calculation_mode": workbook.calculation.calcMode,
        "sheets": sheets,
    }


def workbook_contract_digest(path: Path) -> str:
    payload = json.dumps(
        workbook_semantic_contract(path),
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


class CharacterizationTests(unittest.TestCase):
    maxDiff = None

    def test_http_route_and_method_contract(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(Path(tmp), access_token="contract-token")

        self.assertEqual(route_contract(app), EXPECTED_ROUTES)

    def test_reconciliation_view_contract_for_matches_allocations_and_policy(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_contract", mode="arvine")
            save_trip_metadata(
                trip,
                {
                    "claim_program": "arvine_only",
                    "start_date": "2026-07-01",
                    "end_date": "2026-07-03",
                    "expected_accounts": ["Corporate 1234"],
                    "policy": {"personal_expense_treatment": "review"},
                },
            )
            expenses = [
                Expense(
                    source_file=Path("hotel.pdf"),
                    expense_id="20260701_#1",
                    date="2026-07-01",
                    supplier_name="Hotel Example",
                    expense_type="hotel",
                    amount=100,
                    currency="USD",
                    tax_documentation_status="ok",
                ),
                Expense(
                    source_file=Path("taxi.pdf"),
                    expense_id="20260702_#2",
                    date="2026-07-02",
                    supplier_name="Airport Taxi",
                    expense_type="transport",
                    amount=40,
                    currency="USD",
                    tax_documentation_status="ok",
                ),
            ]
            transactions = [
                NormalizedTransaction(
                    source_file=Path("corporate.csv"),
                    source_row=2,
                    provider="amex",
                    transaction_group_id="G-HOTEL",
                    funding_leg_id="G-HOTEL:1",
                    transaction_date="2026-07-01",
                    account_label="Corporate 1234",
                    description="HOTEL EXAMPLE",
                    match_eligible=True,
                    purchase_amount=100,
                    purchase_currency="USD",
                    settlement_amount=130,
                    settlement_currency="CAD",
                    cad_amount=130,
                    cad_completeness="complete",
                    expense_id="20260701_#1",
                    suggested_expense_id="20260701_#1",
                    match_status="auto",
                    match_confidence=0.95,
                ),
                NormalizedTransaction(
                    source_file=Path("corporate.csv"),
                    source_row=3,
                    provider="amex",
                    transaction_group_id="G-SPLIT",
                    funding_leg_id="G-SPLIT:1",
                    transaction_date="2026-07-02",
                    account_label="Corporate 1234",
                    description="TAXI AND PERSONAL",
                    match_eligible=True,
                    purchase_amount=50,
                    purchase_currency="USD",
                    settlement_amount=65,
                    settlement_currency="CAD",
                    cad_amount=65,
                    cad_completeness="complete",
                ),
            ]
            groups = aggregate_transaction_groups(expenses, transactions)
            apply_allocations_to_groups(
                groups,
                {
                    "G-SPLIT": [
                        {
                            "allocation_id": "taxi",
                            "type": "purchase",
                            "invoice_file": "taxi.pdf",
                            "cad_amount": 52,
                            "original_amount": 40,
                            "note": "business taxi",
                        },
                        {
                            "allocation_id": "personal",
                            "type": "personal",
                            "invoice_file": None,
                            "cad_amount": 13,
                            "original_amount": 10,
                            "note": "personal portion",
                        },
                    ]
                },
            )
            state = {
                "version": 1,
                "mode": "arvine",
                "synced_at": "2026-07-04T12:00:00",
                "input_fingerprint": reconciliation_input_fingerprint(trip),
                "requires_resync": False,
                "expenses": [serialize_expense(expense) for expense in expenses],
                "invoice_overrides": {"hotel.pdf": {"description": "Reviewed hotel"}},
                "manual_cad_overrides": {},
                "manual_matches": {},
                "estimated_cad_by_expense": {},
                "transactions": groups,
                "coverage_confirmation": None,
                "warnings": ["fixture warning"],
            }

            view = reconciliation_view(trip, state)

        contract = {
            "available": view["available"],
            "stale": view["stale"],
            "synced_at": view["synced_at"],
            "summary": view["summary"],
            "expenses": [
                {
                    key: expense.get(key)
                    for key in (
                        "source_file",
                        "overridden_fields",
                        "cad_amount_used",
                        "cad_source",
                        "statement_purchase_amount_used",
                        "statement_purchase_currency",
                        "fx_rate",
                        "fx_basis_status",
                    )
                }
                for expense in view["expenses"]
            ],
            "transactions": [
                {
                    key: transaction.get(key)
                    for key in (
                        "group_id",
                        "expense_file",
                        "match_status",
                        "allocation_total",
                        "allocation_balance",
                        "allocation_status",
                        "cad_source",
                        "fx_rate",
                    )
                }
                for transaction in view["transactions"]
            ],
            "policy_warnings": view["policy_warnings"],
            "coverage": {
                "expected_accounts": view["coverage"]["expected_accounts"],
                "expected_start": view["coverage"]["expected_start"],
                "expected_end": view["coverage"]["expected_end"],
                "gaps": view["coverage"]["gaps"],
            },
        }
        self.assertEqual(
            contract,
            {
                "available": True,
                "stale": False,
                "synced_at": "2026-07-04T12:00:00",
                "summary": {
                    "invoice_count": 2,
                    "matched_invoice_count": 2,
                    "transaction_count": 2,
                    "matched_transaction_count": 2,
                    "audit_transaction_count": 0,
                    "incomplete_allocation_count": 0,
                    "unresolved_policy_warning_count": 1,
                    "needs_review_count": 1,
                },
                "expenses": [
                    {
                        "source_file": "hotel.pdf",
                        "overridden_fields": ["description"],
                        "cad_amount_used": 130.0,
                        "cad_source": "statement",
                        "statement_purchase_amount_used": 100.0,
                        "statement_purchase_currency": "USD",
                        "fx_rate": 1.3,
                        "fx_basis_status": "statement_receipt_total",
                    },
                    {
                        "source_file": "taxi.pdf",
                        "overridden_fields": [],
                        "cad_amount_used": 52.0,
                        "cad_source": "allocation",
                        "statement_purchase_amount_used": 40.0,
                        "statement_purchase_currency": "USD",
                        "fx_rate": 1.3,
                        "fx_basis_status": "statement_receipt_total",
                    },
                ],
                "transactions": [
                    {
                        "group_id": "G-HOTEL",
                        "expense_file": "hotel.pdf",
                        "match_status": "auto",
                        "allocation_total": 0.0,
                        "allocation_balance": 130,
                        "allocation_status": "none",
                        "cad_source": "statement",
                        "fx_rate": 1.3,
                    },
                    {
                        "group_id": "G-SPLIT",
                        "expense_file": None,
                        "match_status": "split",
                        "allocation_total": 65.0,
                        "allocation_balance": 0.0,
                        "allocation_status": "balanced",
                        "cad_source": "allocation",
                        "fx_rate": None,
                    },
                ],
                "policy_warnings": [
                    {
                        "id": "personal:G-SPLIT:personal",
                        "message": "TAXI AND PERSONAL: personal allocation of 13 CAD requires review.",
                        "exception_note": "",
                        "resolved": False,
                    }
                ],
                "coverage": {
                    "expected_accounts": ["Corporate 1234"],
                    "expected_start": "2026-07-01",
                    "expected_end": "2026-07-03",
                    "gaps": [
                        {
                            "code": "coverage_ends_early",
                            "message": (
                                "AMEX Corporate 1234 ends on 2026-07-02, before the expected "
                                "coverage end 2026-07-03."
                            ),
                        }
                    ],
                },
            },
        )

    def test_arvine_workbook_semantic_contract(self):
        with tempfile.TemporaryDirectory() as tmp:
            trip = Path(tmp) / "202607_contract"
            receipts = trip / "expenses_receipts"
            statements = trip / "card_statements"
            receipts.mkdir(parents=True)
            statements.mkdir()
            receipt = receipts / "meal.pdf"
            statement = statements / "corporate.csv"
            receipt.touch()
            statement.touch()
            expense = Expense(
                source_file=receipt,
                expense_id="20260701_#1",
                date="2026-07-01",
                supplier_name="Bistro Example",
                description="Client dinner",
                expense_type="meal",
                amount=115,
                currency="CAD",
                country="Canada",
                province="QC",
                gst_hst=5,
                qst=10,
                tax_documentation_status="ok",
            )
            transaction = NormalizedTransaction(
                source_file=statement,
                source_row=2,
                provider="amex",
                transaction_group_id="G1",
                funding_leg_id="G1:1",
                transaction_date="2026-07-01",
                account_label="Corporate 1234",
                description="BISTRO EXAMPLE",
                match_eligible=True,
                purchase_amount=115,
                purchase_currency="CAD",
                settlement_amount=115,
                settlement_currency="CAD",
                cad_amount=115,
                cad_completeness="complete",
            )

            output = build_arvine_workbook(trip, [expense], [transaction])

            self.assertEqual(
                workbook_contract_digest(output),
                "7980bd3051096ce42f1091e7b6858a322e8351a2cc3a5f7b1eea2a5525bca2de",
            )


if __name__ == "__main__":
    unittest.main()
