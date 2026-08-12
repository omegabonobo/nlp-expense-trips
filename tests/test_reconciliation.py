from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from nlp_expenses.fx_rates import FxRateUnavailable
from nlp_expenses.generator import generate_review
from nlp_expenses.line_items import line_item_review_view
from nlp_expenses.models import Expense
from nlp_expenses.reconciliation import (
    InvoiceValidationError,
    confirm_statement_coverage,
    ensure_reconciliation_ready,
    load_manual_matches,
    reconciliation_view,
    serialized_candidate_reason,
    serialized_candidate_score,
    set_coverage_settings,
    set_invoice_review,
    set_manual_match,
    set_transaction_allocations,
    set_transaction_decision,
    sync_reconciliation,
)
from nlp_expenses.statement_normalizer import set_statement_date_convention
from nlp_expenses.trips import ensure_trip


def write_generic_statement(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Date", "Description", "Amount", "Currency", "Foreign Amount"])
        writer.writerow(["2026-07-01", "FOREIGN HOTEL", 130, "CAD", "100 USD"])
        writer.writerow(["2026-07-02", "AIRPORT TAXI", 65, "CAD", "50 USD"])


class ReconciliationTests(unittest.TestCase):
    def test_receipt_matcher_ranks_same_date_merchant_with_tax_tip_gap(self):
        expense = {
            "date": "2026-07-10",
            "vendor": "Bistro Montreal",
            "amount": 70.0,
            "currency": "CAD",
            "number_of_people": 1,
        }
        transaction = {
            "transaction_date": "2026-07-10",
            "description": "SQ *BISTRO MONTREAL",
            "purchase_amount": 100.0,
            "purchase_currency": "CAD",
            "cad_amount": 100.0,
        }

        self.assertGreaterEqual(serialized_candidate_score(expense, transaction, None), 0.72)
        self.assertIn("30% below the card total", serialized_candidate_reason(expense, transaction))

    def test_nested_receipt_folders_are_recursive_and_duplicate_basenames_stay_distinct(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_nested", mode="arvine")
            hotel_file = trip / "expenses_receipts" / "meta ads" / "2026-06" / "invoice.pdf"
            taxi_file = trip / "expenses_receipts" / "travel" / "invoice.pdf"
            hotel_file.parent.mkdir(parents=True)
            taxi_file.parent.mkdir(parents=True)
            hotel_file.write_bytes(b"hotel")
            taxi_file.write_bytes(b"taxi")
            write_generic_statement(trip / "card_statements" / "card.csv")

            def parsed(path: Path, **_kwargs) -> Expense:
                if "meta ads" in path.parts:
                    return Expense(
                        source_file=path,
                        expense_id="",
                        date="2026-07-01",
                        supplier_name="Foreign Hotel",
                        expense_type="hotel",
                        amount=100,
                        currency="USD",
                    )
                return Expense(
                    source_file=path,
                    expense_id="",
                    date="2026-07-02",
                    supplier_name="Airport Taxi",
                    expense_type="transport",
                    amount=50,
                    currency="USD",
                )

            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                view = sync_reconciliation(trip, root, llm_mode="off")

            expected_files = {
                "meta ads/2026-06/invoice.pdf",
                "travel/invoice.pdf",
            }
            self.assertEqual(
                {expense["source_file"] for expense in view["expenses"]},
                expected_files,
            )
            self.assertEqual(
                {receipt["source_file"] for receipt in line_item_review_view(trip)["receipts"]},
                expected_files,
            )
            self.assertEqual(view["summary"]["matched_invoice_count"], 2)

            hotel_transaction = next(
                item for item in view["transactions"] if item["description"] == "FOREIGN HOTEL"
            )
            manually_mapped = set_manual_match(
                trip,
                hotel_transaction["group_id"],
                expense_file="travel/invoice.pdf",
            )
            changed = next(
                item
                for item in manually_mapped["transactions"]
                if item["group_id"] == hotel_transaction["group_id"]
            )
            self.assertEqual(changed["expense_file"], "travel/invoice.pdf")

            output = trip / "nested-receipts.xlsx"
            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                generate_review(
                    trip,
                    root,
                    llm_mode="off",
                    statements_complete=True,
                    output_path=output,
                )
            workbook = load_workbook(output, data_only=False)
            detail = workbook["expense_detail"]
            self.assertEqual(
                {detail.cell(row, 32).value for row in range(2, detail.max_row + 1)},
                expected_files,
            )

            later_folder = trip / "expenses_receipts" / "meta ads" / "2026-07"
            later_folder.mkdir(parents=True)
            (later_folder / "another.pdf").write_bytes(b"new")
            self.assertTrue(reconciliation_view(trip)["stale"])

    def test_sync_manual_override_fx_rate_and_workbook_reuse(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_montreal", mode="arvine")
            hotel_file = trip / "expenses_receipts" / "hotel.pdf"
            taxi_file = trip / "expenses_receipts" / "taxi.pdf"
            hotel_file.write_bytes(b"hotel")
            taxi_file.write_bytes(b"taxi")
            write_generic_statement(trip / "card_statements" / "card.csv")

            def parsed(path: Path, **_kwargs) -> Expense:
                if path.name == "hotel.pdf":
                    return Expense(
                        source_file=path,
                        expense_id="",
                        date="2026-07-01",
                        supplier_name="Foreign Hotel",
                        expense_type="hotel",
                        amount=100,
                        currency="USD",
                    )
                return Expense(
                    source_file=path,
                    expense_id="",
                    date="2026-07-02",
                    supplier_name="Airport Taxi",
                    expense_type="transport",
                    amount=50,
                    currency="USD",
                )

            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                view = sync_reconciliation(trip, root, llm_mode="off")

            self.assertEqual(view["summary"]["matched_invoice_count"], 2)
            hotel_transaction = next(
                item for item in view["transactions"] if item["description"] == "FOREIGN HOTEL"
            )
            self.assertEqual(hotel_transaction["expense_file"], "hotel.pdf")
            self.assertEqual(hotel_transaction["match_status"], "auto")
            self.assertAlmostEqual(hotel_transaction["fx_rate"], 1.3)

            manual = set_manual_match(trip, hotel_transaction["group_id"], expense_file="taxi.pdf")
            changed = next(
                item
                for item in manual["transactions"]
                if item["group_id"] == hotel_transaction["group_id"]
            )
            self.assertEqual(changed["match_status"], "manual")
            self.assertEqual(changed["expense_file"], "taxi.pdf")
            # The row-level rate uses the card transaction's own original and CAD
            # amounts instead of dividing by an unrelated invoice total.
            self.assertAlmostEqual(changed["fx_rate"], 1.3)
            taxi_review = next(
                item for item in manual["expenses"] if item["source_file"] == "taxi.pdf"
            )
            self.assertEqual(taxi_review["statement_purchase_amount_used"], 150.0)
            self.assertAlmostEqual(taxi_review["fx_rate"], 1.3)
            self.assertEqual(load_manual_matches(trip)[hotel_transaction["group_id"]], "taxi.pdf")

            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                resynced = sync_reconciliation(trip, root, llm_mode="off")
            preserved = next(
                item
                for item in resynced["transactions"]
                if item["group_id"] == hotel_transaction["group_id"]
            )
            self.assertEqual(preserved["expense_file"], "taxi.pdf")
            self.assertEqual(preserved["match_status"], "manual")

            output = trip / "manual-mapping.xlsx"
            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                generate_review(
                    trip,
                    root,
                    llm_mode="off",
                    statements_complete=True,
                    output_path=output,
                )
            workbook = load_workbook(output, data_only=False)
            statements = workbook["card_statements"]
            hotel_row = next(
                row
                for row in range(2, statements.max_row + 1)
                if statements.cell(row, 8).value == "FOREIGN HOTEL"
            )
            self.assertEqual(statements.cell(hotel_row, 20).value, "20260702_#2")
            self.assertEqual(statements.cell(hotel_row, 22).value, "manual")

            unmatched = set_manual_match(trip, hotel_transaction["group_id"], expense_file=None)
            cleared = next(
                item
                for item in unmatched["transactions"]
                if item["group_id"] == hotel_transaction["group_id"]
            )
            self.assertIsNone(cleared["expense_file"])
            self.assertEqual(cleared["match_status"], "unmatched")

            restored = set_manual_match(trip, hotel_transaction["group_id"], use_auto=True)
            automatic = next(
                item
                for item in restored["transactions"]
                if item["group_id"] == hotel_transaction["group_id"]
            )
            self.assertEqual(automatic["expense_file"], "hotel.pdf")
            self.assertEqual(automatic["match_status"], "auto")
            self.assertEqual(automatic["match_confidence"], hotel_transaction["match_confidence"])
            self.assertNotIn(hotel_transaction["group_id"], load_manual_matches(trip))

    def test_reconciliation_becomes_stale_when_inputs_change(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_montreal", mode="arvine")
            receipt = trip / "expenses_receipts" / "hotel.pdf"
            receipt.write_bytes(b"hotel")
            write_generic_statement(trip / "card_statements" / "card.csv")
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Foreign Hotel",
                expense_type="hotel",
                amount=100,
                currency="USD",
            )
            with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
                sync_reconciliation(trip, root, llm_mode="off")
            self.assertFalse(reconciliation_view(trip)["stale"])
            (trip / "expenses_receipts" / "new.pdf").write_bytes(b"new")
            self.assertTrue(reconciliation_view(trip)["stale"])
            self.assertEqual(load_manual_matches(trip), {})
            with self.assertRaisesRegex(ValueError, "Sync again"):
                ensure_reconciliation_ready(trip)
            with self.assertRaisesRegex(ValueError, "Sync again"):
                set_manual_match(
                    trip,
                    reconciliation_view(trip)["transactions"][0]["group_id"],
                    expense_file="hotel.pdf",
                )
            with self.assertRaisesRegex(ValueError, "Sync again"):
                generate_review(
                    trip,
                    root,
                    llm_mode="off",
                    statements_complete=True,
                    output_path=trip / "stale.xlsx",
                )
            self.assertFalse((trip / "stale.xlsx").exists())

    def test_changing_statement_date_convention_makes_reconciliation_stale(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_date-choice", mode="arvine")
            receipt = trip / "expenses_receipts" / "hotel.pdf"
            receipt.write_bytes(b"hotel")
            statement = trip / "card_statements" / "ambiguous.csv"
            statement.write_text(
                "Date,Description,Amount,Currency\n07/01/2026,HOTEL,100,CAD\n",
                encoding="utf-8",
            )
            set_statement_date_convention(trip, statement.name, "month_first")
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Hotel",
                expense_type="hotel",
                amount=100,
                currency="CAD",
            )
            with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
                sync_reconciliation(trip, root, llm_mode="off")
            self.assertFalse(reconciliation_view(trip)["stale"])
            set_statement_date_convention(trip, statement.name, "day_first")
            self.assertTrue(reconciliation_view(trip)["stale"])
            with self.assertRaisesRegex(ValueError, "Sync again"):
                ensure_reconciliation_ready(trip)

    def test_invoice_corrections_and_manual_cad_override_flow_into_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_review", mode="arvine")
            receipt = trip / "expenses_receipts" / "hotel.pdf"
            receipt.write_bytes(b"hotel")
            write_generic_statement(trip / "card_statements" / "card.csv")

            def parsed(_path: Path, **_kwargs) -> Expense:
                return Expense(
                    source_file=receipt,
                    expense_id="",
                    date="2026-07-01",
                    supplier_name="Foreign Hotel",
                    expense_type="hotel",
                    amount=100,
                    currency="USD",
                    country="US",
                    review_note="Check taxes.",
                )

            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                initial = sync_reconciliation(trip, root, llm_mode="off")
            initial_confidence = next(
                item["match_confidence"]
                for item in initial["transactions"]
                if item["description"] == "FOREIGN HOTEL"
            )

            corrected = set_invoice_review(
                trip,
                "hotel.pdf",
                fields={
                    "date": "2026-07-01",
                    "vendor": "Montreal Hotel",
                    "description": "Client accommodation",
                    "expense_type": "hotel",
                    "amount": "90",
                    "currency": "EUR",
                    "country": "CA",
                    "province": "QC",
                    "gst_hst": "4.50",
                    "qst": "8.98",
                    "gst_hst_number": "GST-1",
                    "qst_number": "QST-1",
                    "business_purpose": "Client planning",
                    "attendees_client": "",
                },
            )
            self.assertTrue(corrected["stale"])
            reviewed_invoice = corrected["expenses"][0]
            self.assertEqual(reviewed_invoice["amount"], 90.0)
            self.assertEqual(reviewed_invoice["currency"], "EUR")
            self.assertIn("amount", reviewed_invoice["overridden_fields"])

            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                resynced = sync_reconciliation(trip, root, llm_mode="off")
            self.assertFalse(resynced["stale"])
            self.assertEqual(resynced["expenses"][0]["vendor"], "Montreal Hotel")
            updated_confidence = next(
                item["match_confidence"]
                for item in resynced["transactions"]
                if item["description"] == "FOREIGN HOTEL"
            )
            self.assertNotEqual(updated_confidence, initial_confidence)

            overridden = set_invoice_review(
                trip,
                "hotel.pdf",
                update_manual_cad=True,
                manual_cad={"amount": "145.50", "note": "Confirmed in card provider portal"},
            )
            invoice = overridden["expenses"][0]
            self.assertEqual(invoice["cad_source"], "manual")
            self.assertEqual(invoice["cad_amount_used"], 145.5)
            self.assertAlmostEqual(invoice["fx_rate"], round(145.5 / 90, 6))

            output = trip / "reviewed.xlsx"
            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                generate_review(
                    trip,
                    root,
                    llm_mode="off",
                    statements_complete=True,
                    output_path=output,
                )
            workbook = load_workbook(output, data_only=False)
            detail = workbook["expense_detail"]
            self.assertEqual(detail["D2"].value, "Montreal Hotel")
            self.assertEqual(detail["I2"].value, 90)
            self.assertEqual(detail["J2"].value, "EUR")
            self.assertEqual(detail["T2"].value, 145.5)
            self.assertIn('IF($T2<>"","manual"', detail["S2"].value)
            self.assertIn("Confirmed in card provider portal", detail["AJ2"].value)

            cleared = set_invoice_review(
                trip,
                "hotel.pdf",
                update_manual_cad=True,
                manual_cad=None,
            )
            self.assertNotEqual(cleared["expenses"][0]["cad_source"], "manual")
            restored = set_invoice_review(trip, "hotel.pdf", restore_extracted=True)
            self.assertTrue(restored["stale"])
            self.assertEqual(restored["expenses"][0]["vendor"], "Foreign Hotel")
            self.assertEqual(restored["expenses"][0]["amount"], 100)
            self.assertEqual(restored["expenses"][0]["overridden_fields"], [])

    def test_invoice_review_validation_is_field_specific(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_validation", mode="arvine")
            receipt = trip / "expenses_receipts" / "receipt.pdf"
            receipt.write_bytes(b"receipt")
            write_generic_statement(trip / "card_statements" / "card.csv")
            expense = Expense(source_file=receipt, expense_id="", amount=100, currency="CAD")
            with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
                sync_reconciliation(trip, root, llm_mode="off")

            with self.assertRaises(InvoiceValidationError) as invalid:
                set_invoice_review(
                    trip,
                    "receipt.pdf",
                    fields={
                        "date": "2026-99-99",
                        "amount": "-1",
                        "currency": "dollars",
                        "gst_hst": "120",
                    },
                )
            self.assertEqual(
                {"date", "amount", "currency", "gst_hst"},
                set(invalid.exception.fields),
            )
            with self.assertRaises(InvoiceValidationError) as missing_note:
                set_invoice_review(
                    trip,
                    "receipt.pdf",
                    update_manual_cad=True,
                    manual_cad={"amount": "125", "note": ""},
                )
            self.assertIn("manual_cad_note", missing_note.exception.fields)

    def test_manual_cad_resolves_incomplete_wise_settlement_review(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_wise", mode="arvine")
            receipt = trip / "expenses_receipts" / "hotel.pdf"
            receipt.write_bytes(b"hotel")
            statement = trip / "card_statements" / "wise.csv"
            with statement.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(
                    [
                        "ID",
                        "Status",
                        "Direction",
                        "Created on",
                        "Finished on",
                        "Source amount after fees",
                        "Source fee amount",
                        "Source currency",
                        "Target amount after fees",
                        "Target currency",
                        "Source name",
                        "Target name",
                        "Category",
                    ]
                )
                writer.writerow(
                    [
                        "CARD_TRANSACTION-1",
                        "COMPLETED",
                        "OUT",
                        "2026-07-01",
                        "2026-07-01",
                        80,
                        0,
                        "CAD",
                        60,
                        "USD",
                        "Wise",
                        "Foreign Hotel",
                        "card",
                    ]
                )
                writer.writerow(
                    [
                        "CARD_TRANSACTION-1",
                        "COMPLETED",
                        "OUT",
                        "2026-07-01",
                        "2026-07-01",
                        20,
                        0,
                        "EUR",
                        15,
                        "USD",
                        "Wise",
                        "Foreign Hotel",
                        "card",
                    ]
                )
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Foreign Hotel",
                expense_type="hotel",
                amount=75,
                currency="USD",
            )
            with (
                patch(
                    "nlp_expenses.fx_rates.WeeklyCadFxResolver.resolve",
                    side_effect=FxRateUnavailable("No cached or published weekly rate."),
                ),
                patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense),
            ):
                initial = sync_reconciliation(trip, root, llm_mode="off")
            self.assertEqual(initial["summary"]["needs_review_count"], 1)
            self.assertEqual(initial["transactions"][0]["cad_completeness"], "partial")

            resolved = set_invoice_review(
                trip,
                "hotel.pdf",
                update_manual_cad=True,
                manual_cad={"amount": 105, "note": "Wise activity detail"},
            )
            self.assertEqual(resolved["summary"]["needs_review_count"], 0)
            self.assertEqual(resolved["transactions"][0]["cad_source"], "manual")
            self.assertAlmostEqual(resolved["transactions"][0]["fx_rate"], 1.4)

    def test_possible_duplicates_are_visible_and_resolvable_without_deletion(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_duplicates", mode="arvine")
            receipt = trip / "expenses_receipts" / "cafe.pdf"
            receipt.write_bytes(b"cafe")
            for filename in ("card-part-1.csv", "card-part-2.csv"):
                (trip / "card_statements" / filename).write_text(
                    "Date,Description,Amount,Currency\n2026-07-01,CLIENT CAFE,10,CAD\n",
                    encoding="utf-8",
                )
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Client Cafe",
                expense_type="meal",
                amount=10,
                currency="CAD",
            )
            warnings: list[str] = []
            with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
                initial = sync_reconciliation(
                    trip,
                    root,
                    llm_mode="off",
                    warning_callback=warnings.append,
                )
            duplicates = [item for item in initial["transactions"] if item["possible_duplicate"]]
            self.assertEqual(len(duplicates), 2)
            self.assertEqual(initial["summary"]["needs_review_count"], 2)
            self.assertTrue(
                any("Possible duplicate statement transactions" in warning for warning in warnings)
            )
            self.assertEqual(
                {source["file"] for item in duplicates for source in item["source_rows"]},
                {"card-part-1.csv", "card-part-2.csv"},
            )

            ignored = set_transaction_decision(
                trip,
                duplicates[0]["group_id"],
                "ignore",
                "Overlapping export; second file is authoritative",
            )
            self.assertEqual(ignored["summary"]["needs_review_count"], 1)
            self.assertEqual(ignored["expenses"][0]["cad_amount_used"], 10)
            ignored_transaction = next(
                item
                for item in ignored["transactions"]
                if item["group_id"] == duplicates[0]["group_id"]
            )
            self.assertTrue(ignored_transaction["ignored"])

            confirmed = set_transaction_decision(trip, duplicates[1]["group_id"], "keep")
            self.assertEqual(confirmed["summary"]["needs_review_count"], 0)
            self.assertEqual(confirmed["summary"]["transaction_count"], 1)
            self.assertEqual(confirmed["summary"]["audit_transaction_count"], 1)

            output = trip / "deduplicated.xlsx"
            with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
                generate_review(
                    trip,
                    root,
                    llm_mode="off",
                    statements_complete=True,
                    output_path=output,
                )
            sheet = load_workbook(output, data_only=False)["card_statements"]
            self.assertEqual(sheet.max_row, 3)
            ignored_rows = [
                row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 26).value == "ignored"
            ]
            self.assertEqual(len(ignored_rows), 1)
            self.assertFalse(sheet.cell(ignored_rows[0], 13).value)
            self.assertIn("Overlapping export", sheet.cell(ignored_rows[0], 27).value)

    def test_statement_coverage_lists_ranges_and_records_gap_acknowledgement(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_coverage", mode="arvine")
            receipt = trip / "expenses_receipts" / "hotel.pdf"
            receipt.write_bytes(b"hotel")
            statement = trip / "card_statements" / "card.csv"
            statement.write_text(
                "Date,Description,Amount,Currency,Account\n"
                "2026-07-01,HOTEL,130,CAD,1234\n"
                "2026-07-05,TAXI,25,CAD,1234\n",
                encoding="utf-8",
            )
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date="2026-07-01",
                supplier_name="Hotel",
                expense_type="hotel",
                amount=130,
                currency="CAD",
            )
            with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
                synced = sync_reconciliation(trip, root, llm_mode="off")
            coverage = synced["coverage"]
            self.assertEqual(len(coverage["accounts"]), 1)
            self.assertEqual(coverage["accounts"][0]["earliest_date"], "2026-07-01")
            self.assertEqual(coverage["accounts"][0]["latest_date"], "2026-07-05")
            self.assertEqual(coverage["accounts"][0]["source_files"], ["card.csv"])

            changed = set_coverage_settings(trip, ["GENERIC ••••1234", "AMEX ••••9999"])
            self.assertEqual(len(changed["gaps"]), 1)
            self.assertIn("AMEX", changed["gaps"][0]["message"])
            with self.assertRaisesRegex(ValueError, "Explain"):
                confirm_statement_coverage(trip)
            confirmation = confirm_statement_coverage(
                trip, "Corporate Amex was not used on this trip"
            )
            self.assertEqual(confirmation["gap_count"], 1)

            output = trip / "coverage.xlsx"
            with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
                generate_review(
                    trip,
                    root,
                    llm_mode="off",
                    statements_complete=True,
                    output_path=output,
                )
            summary = load_workbook(output, data_only=False)["expense_summary"]
            self.assertIn("Corporate Amex was not used", summary["E6"].value)
            self.assertEqual(summary["E7"].value, 1)

    def test_split_allocations_refund_and_personal_portion_flow_to_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_allocations", mode="arvine")
            first_file = trip / "expenses_receipts" / "first.pdf"
            second_file = trip / "expenses_receipts" / "second.pdf"
            first_file.write_bytes(b"first")
            second_file.write_bytes(b"second")
            statement = trip / "card_statements" / "card.csv"
            statement.write_text(
                "Date,Description,Amount,Currency\n"
                "2026-07-01,CONSOLIDATED CHARGE,300,CAD\n"
                "2026-07-03,MERCHANT REFUND,-20,CAD\n",
                encoding="utf-8",
            )

            def parsed(path: Path, **_kwargs) -> Expense:
                if path.name == "first.pdf":
                    return Expense(
                        source_file=path,
                        expense_id="",
                        date="2026-07-01",
                        supplier_name="First Vendor",
                        expense_type="hotel",
                        amount=100,
                        currency="CAD",
                    )
                return Expense(
                    source_file=path,
                    expense_id="",
                    date="2026-07-01",
                    supplier_name="Second Vendor",
                    expense_type="transport",
                    amount=150,
                    currency="CAD",
                )

            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                synced = sync_reconciliation(trip, root, llm_mode="off")
            purchase = next(
                item for item in synced["transactions"] if item["transaction_type"] == "purchase"
            )
            refund = next(
                item for item in synced["transactions"] if item["transaction_type"] == "refund"
            )

            partial = set_transaction_allocations(
                trip,
                purchase["group_id"],
                [
                    {
                        "type": "purchase",
                        "invoice_file": "first.pdf",
                        "cad_amount": 100,
                        "note": "Hotel share",
                    }
                ],
            )
            partial_group = next(
                item for item in partial["transactions"] if item["group_id"] == purchase["group_id"]
            )
            self.assertEqual(partial_group["allocation_status"], "unallocated")
            self.assertEqual(partial_group["allocation_balance"], 200)
            with self.assertRaisesRegex(ValueError, "Finish split allocations"):
                ensure_reconciliation_ready(trip)

            overallocated = set_transaction_allocations(
                trip,
                purchase["group_id"],
                [
                    {"type": "purchase", "invoice_file": "first.pdf", "cad_amount": 150},
                    {"type": "purchase", "invoice_file": "second.pdf", "cad_amount": 200},
                ],
            )
            over_group = next(
                item
                for item in overallocated["transactions"]
                if item["group_id"] == purchase["group_id"]
            )
            self.assertEqual(over_group["allocation_status"], "overallocated")
            self.assertEqual(over_group["allocation_balance"], -50)

            set_transaction_allocations(
                trip,
                purchase["group_id"],
                [
                    {
                        "type": "purchase",
                        "invoice_file": "first.pdf",
                        "cad_amount": 100,
                        "original_amount": 100,
                        "note": "Hotel portion",
                    },
                    {
                        "type": "purchase",
                        "invoice_file": "second.pdf",
                        "cad_amount": 150,
                        "original_amount": 150,
                        "note": "Transport portion",
                    },
                    {
                        "type": "personal",
                        "category": "Personal extension",
                        "cad_amount": 50,
                        "note": "Not reimbursable",
                    },
                ],
            )
            balanced = set_transaction_allocations(
                trip,
                refund["group_id"],
                [
                    {
                        "type": "refund",
                        "invoice_file": "first.pdf",
                        "cad_amount": -20,
                        "note": "Refund to hotel charge",
                    }
                ],
            )
            purchase_group = next(
                item
                for item in balanced["transactions"]
                if item["group_id"] == purchase["group_id"]
            )
            self.assertEqual(purchase_group["allocation_status"], "balanced")
            self.assertEqual(purchase_group["allocation_balance"], 0)
            expenses = {expense["source_file"]: expense for expense in balanced["expenses"]}
            self.assertEqual(expenses["first.pdf"]["cad_amount_used"], 80)
            self.assertEqual(expenses["second.pdf"]["cad_amount_used"], 150)
            self.assertEqual(balanced["summary"]["incomplete_allocation_count"], 0)

            output = trip / "allocations.xlsx"
            with patch("nlp_expenses.generator.parse_arvine_receipt", side_effect=parsed):
                generate_review(
                    trip,
                    root,
                    llm_mode="off",
                    statements_complete=True,
                    output_path=output,
                )
            workbook = load_workbook(output, data_only=False)
            cards = workbook["card_statements"]
            allocation_rows = [
                row
                for row in range(2, cards.max_row + 1)
                if cards.cell(row, 5).value == "allocation"
            ]
            self.assertEqual(len(allocation_rows), 4)
            personal_row = next(
                row for row in allocation_rows if cards.cell(row, 10).value == "personal"
            )
            self.assertFalse(cards.cell(personal_row, 13).value)
            self.assertIsNone(cards.cell(personal_row, 20).value)
            linked_cad = sum(
                cards.cell(row, 18).value
                for row in allocation_rows
                if cards.cell(row, 20).value == "20260701_#1"
            )
            self.assertEqual(linked_cad, 80)
            self.assertIn('"allocation"', workbook["expense_detail"]["AH2"].value)


if __name__ == "__main__":
    unittest.main()
