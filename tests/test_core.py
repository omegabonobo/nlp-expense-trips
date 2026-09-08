from __future__ import annotations

import csv
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from nlp_expenses.config import ask_openai_for_run
from nlp_expenses.extraction.alcohol import detect_alcohol, is_alcohol
from nlp_expenses.extraction.receipts import (
    MULTI_DOCUMENT_RECEIPT_PROMPT,
    apply_missing_date_fallback,
    find_date,
    find_filename_date,
    find_invoice_date,
    heuristic_parse_receipt,
    is_multi_page_pdf,
    line_item_from_llm,
    parse_receipt,
)
from nlp_expenses.extraction.statements import parse_csv_statement
from nlp_expenses.generator import (
    SUPPORTED_RECEIPTS,
    assign_simple_expense_ids,
    generate_review,
    resolve_run_settings,
)
from nlp_expenses.matching import (
    close_split_amount,
    enrich_expenses_from_statements,
    match_normalized_transactions,
    match_score,
)
from nlp_expenses.models import Expense, NormalizedTransaction, StatementTransaction
from nlp_expenses.trips import ensure_trip, validate_trip_name
from nlp_expenses.workbook import build_workbook


class CoreTests(unittest.TestCase):
    def test_unknown_llm_mode_is_rejected(self):
        with (
            tempfile.TemporaryDirectory() as tmp,
            self.assertRaisesRegex(ValueError, "Unknown LLM mode"),
        ):
            resolve_run_settings(Path(tmp), "sometimes", allow_openai_prompt=False)

    def test_trip_name_validation_and_creation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.assertTrue(validate_trip_name("202606_melbourne"))
            self.assertFalse(validate_trip_name("melbourne_202606"))
            trip = ensure_trip(root, "202606_melbourne")
            self.assertTrue((trip / "expenses_receipts").is_dir())
            self.assertTrue((trip / "card_statements").is_dir())

    def test_openai_prompt_can_decline_or_store_key(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with patch("builtins.input", return_value="n"):
                api_key, _model = ask_openai_for_run(root)
            self.assertIsNone(api_key)
            self.assertFalse((root / ".env").exists())

            with (
                patch("builtins.input", return_value="y"),
                patch("nlp_expenses.config.getpass", return_value="sk-test"),
            ):
                api_key, model = ask_openai_for_run(root)
            self.assertEqual(api_key, "sk-test")
            self.assertEqual(model, "gpt-5.2")
            self.assertIn("OPENAI_API_KEY=sk-test", (root / ".env").read_text(encoding="utf-8"))

    def test_filename_date_and_simple_expense_ids(self):
        self.assertEqual(find_date("Scanned_20260530 - Receipt.pdf"), "2026-05-30")
        self.assertEqual(find_filename_date("Scanned_20260530 - Receipt.pdf"), "2026-05-30")
        expenses = [
            Expense(source_file=Path("a.pdf"), expense_id="", date="2026-05-30"),
            Expense(source_file=Path("b.pdf"), expense_id="", date=None),
        ]
        assign_simple_expense_ids(expenses)
        self.assertEqual(expenses[0].expense_id, "20260530_#1")
        self.assertEqual(expenses[1].expense_id, "yyyymmdd_#2")

    def test_filename_date_fallback_supports_conservative_common_formats(self):
        cases = {
            "receipt_20260530.pdf": "2026-05-30",
            "receipt_2026-05-30.pdf": "2026-05-30",
            "receipt_2026_05_30.jpg": "2026-05-30",
            "receipt_2026.05.30.heic": "2026-05-30",
            "receipt_30-05-2026.png": "2026-05-30",
            "receipt_05-30-2026.png": "2026-05-30",
            "receipt_30052026.png": "2026-05-30",
            "receipt_05302026.png": "2026-05-30",
            "receipt_May_30_2026.pdf": "2026-05-30",
        }
        for filename, expected in cases.items():
            with self.subTest(filename=filename):
                self.assertEqual(find_filename_date(filename), expected)
        self.assertIsNone(find_filename_date("receipt_2026-13-40.pdf"))
        self.assertIsNone(find_filename_date("receipt_05-06-2026.pdf"))
        self.assertIsNone(find_filename_date("receipt_20260530_to_20260531.pdf"))

    def test_receipt_content_date_wins_and_filename_is_only_a_reviewable_fallback(self):
        with tempfile.TemporaryDirectory() as tmp:
            receipt_with_date = Path(tmp) / "receipt_20260530.pdf"
            parsed = heuristic_parse_receipt(
                receipt_with_date,
                "Cafe Montreal\nPurchase Date June 2, 2026\nTotal CAD 12.00",
            )
            self.assertEqual(parsed.date, "2026-06-02")
            self.assertIn("kept the receipt date", parsed.review_note)

            receipt_without_date = Path(tmp) / "receipt_2026_05_30.pdf"
            fallback = heuristic_parse_receipt(
                receipt_without_date,
                "Cafe Montreal\nCappuccino CAD 5.00\nTotal CAD 5.00",
            )
            self.assertEqual(fallback.date, "2026-05-30")
            self.assertIn("inferred from the source filename", fallback.review_note)
            self.assertLessEqual(fallback.confidence, 0.70)

    def test_missing_structured_date_falls_back_after_openai_extraction(self):
        expense = Expense(
            source_file=Path("receipt_20260530.pdf"),
            expense_id="",
            date=None,
            confidence=0.92,
            review_note="Structured with OpenAI receipt extraction.",
        )
        apply_missing_date_fallback(expense, expense.source_file, "Cafe Montreal\nTotal CAD 5.00")
        self.assertEqual(expense.date, "2026-05-30")
        self.assertIn("inferred from the source filename", expense.review_note)
        self.assertLessEqual(expense.confidence, 0.70)

    def test_asking_for_api_key_forces_llm_over_high_confidence_heuristics(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202606_test")
            receipt = trip / "expenses_receipts" / "Scanned_20260601-test.pdf"
            receipt.write_bytes(b"dummy")
            mocked_expense = Expense(
                source_file=receipt, expense_id="", date="2026-06-01", supplier_name="LLM Cafe"
            )
            with (
                patch(
                    "nlp_expenses.generator.ask_openai_for_run",
                    return_value=("sk-test", "gpt-test"),
                ),
                patch(
                    "nlp_expenses.generator.parse_receipt", return_value=mocked_expense
                ) as parse_receipt_mock,
            ):
                output = generate_review(trip, root, llm_mode="ask")
            self.assertTrue(output.exists())
            self.assertEqual(parse_receipt_mock.call_args.kwargs["use_llm"], True)
            self.assertEqual(parse_receipt_mock.call_args.kwargs["force_llm"], True)

    def test_alcohol_detection(self):
        self.assertTrue(is_alcohol("Glass of shiraz"))
        self.assertTrue(is_alcohol("Pint lager"))
        self.assertTrue(is_alcohol("Canta"))
        self.assertTrue(is_alcohol("Pisco Sour"))
        self.assertTrue(is_alcohol("GLS 2023 JC Own Lobethal Chardonnay"))
        self.assertTrue(is_alcohol("Stomping Ground IPA"))
        self.assertTrue(is_alcohol("2x Peroni Nastro Azzurro"))
        self.assertTrue(is_alcohol("Hefeweizen 500ml"))
        self.assertTrue(is_alcohol("French 75"))
        self.assertTrue(is_alcohol("Long Island Iced Tea"))
        self.assertTrue(is_alcohol("Don Julio Blanco"))
        self.assertTrue(is_alcohol("Corona 330ml"))
        self.assertTrue(is_alcohol("Aviation"))
        self.assertTrue(is_alcohol("2 x Penicillin"))
        self.assertTrue(is_alcohol("Hard seltzer"))
        self.assertTrue(is_alcohol("Rosé"))
        self.assertTrue(is_alcohol("Craft lager 5.2% ABV"))
        self.assertFalse(is_alcohol("Flat white coffee"))
        self.assertFalse(is_alcohol("Sparkling Water"))
        self.assertFalse(is_alcohol("Ginger Beer"))
        self.assertFalse(is_alcohol("Root beer float"))
        self.assertFalse(is_alcohol("Heineken 0.0% non-alcoholic"))
        self.assertFalse(is_alcohol("Heineken 0.0"))
        self.assertFalse(is_alcohol("Virgin Mojito"))
        self.assertFalse(is_alcohol("Beer battered fish"))
        self.assertFalse(is_alcohol("Red wine vinegar"))
        self.assertFalse(is_alcohol("Americano coffee"))
        self.assertFalse(is_alcohol("Galician Scotch Filet"))
        self.assertFalse(is_alcohol("Scotch fillet steak"))
        self.assertFalse(is_alcohol("GST (10% incl.)"))
        self.assertFalse(is_alcohol("Weekend surcharge 15%"))
        self.assertFalse(is_alcohol("10% discount"))
        self.assertTrue(is_alcohol("TIGER SCHe"))
        self.assertTrue(is_alcohol("ASAHTML"))
        self.assertTrue(is_alcohol("Strawberry Fielss"))
        self.assertTrue(is_alcohol("East Skipper"))

    def test_alcohol_detection_explains_its_decision(self):
        cocktail = detect_alcohol("French 75")
        self.assertTrue(cocktail.is_alcohol)
        self.assertEqual(cocktail.reason, "recognized cocktail")
        self.assertEqual(cocktail.matched_term, "french 75")
        self.assertGreaterEqual(cocktail.confidence, 0.95)

        excluded = detect_alcohol("Virgin Mojito")
        self.assertFalse(excluded.is_alcohol)
        self.assertEqual(excluded.reason, "non-alcoholic drink style")
        self.assertEqual(excluded.matched_term, "virgin")

    def test_openai_alcohol_flag_cannot_turn_tax_percentage_into_alcohol(self):
        item = line_item_from_llm(
            {"description": "GST (10% incl.)", "amount": 0.07, "is_alcohol": True},
            extraction_confidence=0.95,
            include_images=True,
        )
        self.assertFalse(item.is_alcohol)
        self.assertTrue(item.included)
        self.assertEqual(item.alcohol_reason, "receipt tax or charge")

    def test_receipt_heuristic_parses_total_and_corrected_amount(self):
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "receipt.pdf"
            file_path.write_bytes(b"dummy")
            text = "\n".join(
                [
                    "Fancy Hanks",
                    "1 Jun 2026",
                    "Burger 20.00",
                    "Beer 9.00",
                    "Total $29.00",
                ]
            )
            expense = heuristic_parse_receipt(file_path, text)
            self.assertEqual(expense.date, "2026-06-01")
            self.assertEqual(expense.currency, "AUD")
            self.assertAlmostEqual(expense.amount, 29.0)
            self.assertAlmostEqual(expense.corrected_amount_in_currency, 20.0)
            self.assertTrue(
                any(
                    item.description == "Alcohol adjustment - manual" and item.is_alcohol
                    for item in expense.line_items
                )
            )

    def test_invoice_date_is_preferred_for_flights(self):
        text = "\n".join(
            [
                "FlightTo Class Date Departure Arrival",
                "GA0718 D 29May 22:05 05:30",
                "Issuing Airline and date IATA: GARUDA INDONESIA 25May26 : 15394643",
                "Total Amount : IDR 20598982",
            ]
        )
        self.assertEqual(find_invoice_date(text), "2026-05-25")
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "Garuda - DPS to MEL.pdf"
            file_path.write_bytes(b"dummy")
            expense = heuristic_parse_receipt(file_path, text)
            self.assertEqual(expense.date, "2026-05-25")
            self.assertEqual(expense.expense_type, "flight")
            receipt_total_items = [
                item for item in expense.line_items if item.description == "Receipt total"
            ]
            self.assertEqual(len(receipt_total_items), 1)
            self.assertEqual(receipt_total_items[0].amount, 20598982.0)
            self.assertTrue(
                any(
                    item.description == "Alcohol adjustment - manual" and item.is_alcohol
                    for item in expense.line_items
                )
            )

    def test_juni_scanned_receipt_keeps_receipt_date_and_restaurant_supplier(self):
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "Scanned_20260530 - Receipt - May 30 2026 - 9-27 PM.pdf"
            file_path.write_bytes(b"dummy")
            text = "\n".join(
                [
                    "TAX INVOICE",
                    "TASLE ACCOUNT 14 = Id/Check 36143",
                    "136 Exhibition Street",
                    "Served by Beth Juni « Juni Restaurant 01",
                    "Covers : 3",
                    "30/8/2026 at 9:17 pm",
                    "Food Sales $ 222.00",
                    "1 x SALMON SASHIMI $ 76.00",
                    "2 « STOMPING GROUND IPA $ 28.00",
                    "2 « ASAHT 400ML $ 30.00",
                    "Total $ 314.00",
                ]
            )
            expense = heuristic_parse_receipt(file_path, text)
            self.assertEqual(expense.date, "2026-08-30")
            self.assertIn("kept the receipt date", expense.review_note)
            self.assertEqual(expense.supplier_name, "Juni Restaurant")
            self.assertEqual(expense.expense_type, "meal-dinner")
            self.assertTrue(any(item.is_alcohol for item in expense.line_items))

    def test_known_restaurant_cafe_names_classify_as_meals(self):
        cases = [
            ("Nigel\nCappuccino $5.50\nBanana Bread $6.00\nTotal $11.50", "meal-breakfast"),
            (
                "Reine & La Rue\nThank you for dining at Reine & La Rue\nTotal Inc Tax: $521.00",
                "meal-dinner",
            ),
            (
                "GABRIEL\nCAPPUCINO x 4\npork belly Benedict 1 $27.00\nSubtotal $43.50",
                "meal-breakfast",
            ),
            (
                "Tax ywvOl\nThame you for dining wan US Bt Farmers\nDaughte’s\nTotal $474.10",
                "meal-dinner",
            ),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            for idx, (text, expected_type) in enumerate(cases):
                file_path = Path(tmp) / f"receipt-{idx}.pdf"
                file_path.write_bytes(b"dummy")
                expense = heuristic_parse_receipt(file_path, text)
                self.assertEqual(expense.expense_type, expected_type)

    def test_restaurant_additive_charges_are_line_items_but_included_gst_is_not(self):
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "Gmail - Receipt from Gillott espresso #rFUK.pdf"
            file_path.write_bytes(b"dummy")
            text = "\n".join(
                [
                    "Gillott espresso",
                    "Purchase Date May 31, 2026",
                    "Eggs Benedict $23.00",
                    "Cappuccino $5.70 A classic espresso coffee topped with steamed milk and a thick, velvety foam. Mug (have here) ($1.00)",
                    "Dine In Small ($0.50)",
                    "A classic espresso coffee topped with steamed milk and a thick, velvety foam. Mug (have here) ($1.00)",
                    "weekend surcharge (15%) $4.30",
                    "Total includes GST $3.00",
                    "Total $33.00",
                ]
            )
            expense = heuristic_parse_receipt(file_path, text)
            items = {item.description: item.amount for item in expense.line_items}
            self.assertEqual(items["Cappuccino"], 5.70)
            self.assertEqual(items["weekend surcharge"], 4.30)
            self.assertNotIn("Total includes GST", items)
            self.assertNotIn("Dine In Small", items)
            self.assertFalse(
                any("classic espresso" in description.lower() for description in items)
            )

    def test_uber_eats_promotions_remain_negative_and_reconcile_the_receipt(self):
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "Receipt_09Jul2026_063952.pdf"
            file_path.write_bytes(b"dummy")
            text = "\n".join(
                [
                    "Here's your receipt for Restaurant Bombay Mahal (Mont-Royal).",
                    "Total CA$47.50",
                    "2 Poulet au Beurre / Butter Chicken CA$35.00",
                    "2 Naal a l'Ail / Garlic Naan CA$8.00",
                    "1 Riz Vapeur / Steamed Rice CA$4.50",
                    "Tax CA$5.47",
                    "Delivery Fee",
                    "CA$0.99",
                    "Service Fee",
                    "CA$6.50",
                    "Tip CA$5.49",
                    "Promotion -CA$17.50",
                    "Service Fee Discount -CA$0.95",
                    "American Express ••••1003 CA$47.50",
                    "Uber Delivery",
                ]
            )

            expense = heuristic_parse_receipt(file_path, text)
            items = {item.description: item.amount for item in expense.line_items}

            self.assertEqual(expense.expense_type, "meal-dinner")
            self.assertEqual(items["Delivery Fee"], 0.99)
            self.assertEqual(items["Service Fee"], 6.50)
            self.assertEqual(items["Tip"], 5.49)
            self.assertEqual(items["Promotion"], -17.50)
            self.assertEqual(items["Service Fee Discount"], -0.95)
            self.assertNotIn("American Express ••••", items)
            self.assertAlmostEqual(
                sum(item.amount or 0 for item in expense.line_items if not item.synthetic),
                47.50,
            )

    def test_ocr_decimal_variants_and_payment_total_are_parsed(self):
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "Scanned_20260602-1936.pdf"
            file_path.write_bytes(b"dummy")
            text = "\n".join(
                [
                    "Craft and Wheelbarrow",
                    "02/06/2026",
                    "BEVERAGE",
                    "Balter Mpe,Ping 1 17.90",
                    "BEVERAGE Total $17.90",
                    "FOOD",
                    "Chicken Peis 1 34,50",
                    "FOOD Total $34.50",
                    "AMERICAN EXPR. $53. 32",
                    "Natl. Handling Fee $0.92",
                    "GST Included $4.85",
                ]
            )
            expense = heuristic_parse_receipt(file_path, text)
            self.assertEqual(expense.amount, 53.32)
            items = {item.description: item.amount for item in expense.line_items}
            self.assertEqual(items["Chicken Peis"], 34.50)
            self.assertEqual(items["Natl. Handling Fee"], 0.92)
            self.assertFalse(any(amount == 4.85 for amount in items.values()))

    def test_tax_summary_rows_do_not_become_restaurant_line_items(self):
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "Scanned_20260604-2259.pdf"
            file_path.write_bytes(b"dummy")
            text = "\n".join(
                [
                    "Reine & La Rue",
                    "2026-06-04 22:58:15",
                    "Canta $28.00",
                    "Pisco Sour $25.00",
                    "Bread x 3 $27.00",
                    "GLS 2023 JC Own Lobethal Chardonnay $25 .00",
                    "Subtotal: $521.00",
                    "Total ex tax: $473.63",
                    "- Tax Free $0.00",
                    "- GST $47.37",
                    "Total Inc Tax: $521 00",
                ]
            )
            expense = heuristic_parse_receipt(file_path, text)
            self.assertEqual(expense.amount, 521.0)
            items = {item.description: item.amount for item in expense.line_items}
            alcohol_items = {item.description for item in expense.line_items if item.is_alcohol}
            self.assertNotIn("Total ex tax", items)
            self.assertNotIn("GST", items)
            self.assertNotIn("Total Inc Tax", items)
            self.assertTrue(any("Chardonnay" in description for description in items))
            self.assertIn("Canta", alcohol_items)
            self.assertIn("Pisco Sour", alcohol_items)
            self.assertTrue(any("Chardonnay" in description for description in alcohol_items))

    def test_positive_meal_shortfall_defaults_to_tip(self):
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "Scanned_20260605-1429.pdf"
            file_path.write_bytes(b"dummy")
            text = "\n".join(
                [
                    "Gabriel",
                    "5 Jun 2026",
                    "pork belly Benedict 1 $27.00",
                    "DOUBLE ESPRESSO 1 $5.00",
                    "Cinnamon Scroll 1 $6.50",
                    "Subtotal $43.50",
                    "Card surcharge (1.5%) $0.85",
                ]
            )
            expense = heuristic_parse_receipt(file_path, text)
            gap_items = [item for item in expense.line_items if item.description == "Tip"]
            self.assertEqual(len(gap_items), 1)
            self.assertAlmostEqual(gap_items[0].amount, 5.0)
            self.assertIn("Tip inferred from the positive gap", gap_items[0].review_note)

    def test_small_positive_meal_shortfall_also_defaults_to_tip(self):
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "meal.pdf"
            file_path.write_bytes(b"dummy")
            expense = heuristic_parse_receipt(
                file_path,
                "Bistro Restaurant\n7 Sep 2026\nDinner $98.00\nTotal $100.00",
            )
            tips = [item for item in expense.line_items if item.description == "Tip"]
            self.assertEqual(len(tips), 1)
            self.assertEqual(tips[0].amount, 2.0)

    def test_multi_document_prompt_covers_split_bill_tip_and_surcharge(self):
        self.assertIn(
            "itemized bill/invoice and a card/EFTPOS payment receipt", MULTI_DOCUMENT_RECEIPT_PROMPT
        )
        self.assertIn("'Meal share'", MULTI_DOCUMENT_RECEIPT_PROMPT)
        self.assertIn("remaining difference as a line named 'Tip'", MULTI_DOCUMENT_RECEIPT_PROMPT)
        self.assertIn("explicitly shown tip, gratuity, surcharge", MULTI_DOCUMENT_RECEIPT_PROMPT)

    def test_multi_page_pdf_is_detected(self):
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as tmp:
            one_page = Path(tmp) / "one.pdf"
            two_pages = Path(tmp) / "two.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            with one_page.open("wb") as handle:
                writer.write(handle)
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.add_blank_page(width=100, height=100)
            with two_pages.open("wb") as handle:
                writer.write(handle)
            self.assertFalse(is_multi_page_pdf(one_page))
            self.assertTrue(is_multi_page_pdf(two_pages))

    def test_multi_page_pdf_forces_vision_aware_llm_pass(self):
        heuristic = Expense(
            source_file=Path("bill.pdf"),
            expense_id="",
            date="2026-09-07",
            supplier_name="Bistro",
            expense_type="hotel",
            amount=100.0,
            currency="AUD",
            confidence=0.95,
        )
        with (
            patch(
                "nlp_expenses.extraction.receipts.extract_text", return_value=("text", "pdf_text")
            ),
            patch(
                "nlp_expenses.extraction.receipts.heuristic_parse_receipt", return_value=heuristic
            ),
            patch("nlp_expenses.extraction.receipts.is_multi_page_pdf", return_value=True),
            patch("nlp_expenses.extraction.receipts.llm_parse_receipt", return_value=None) as llm,
        ):
            parse_receipt(Path("bill.pdf"), use_llm=True)
        self.assertTrue(llm.call_args.kwargs["include_images"])

    def test_implausible_restaurant_amounts_are_not_accepted_as_menu_items(self):
        with tempfile.TemporaryDirectory() as tmp:
            file_path = Path(tmp) / "Scanned_20260605-1429.pdf"
            file_path.write_bytes(b"dummy")
            text = "\n".join(
                [
                    "GABRIEL",
                    "BIPINKUMAR $9 676 584 984",
                    "CAPPUCINO $5.50",
                    "pork belly Benedict $27.00",
                    "DOUBLE ESPRESSO ~« 1 $5.00",
                    "Cinnamon Scroll * 1 $6.50",
                    "Subtotal $38.50",
                    "Card surcharge (1.5%) $085",
                ]
            )
            expense = heuristic_parse_receipt(file_path, text)
            self.assertEqual(expense.amount, 39.35)
            self.assertNotIn(984.0, [item.amount for item in expense.line_items])
            self.assertNotIn(9.0, [item.amount for item in expense.line_items])
            self.assertIn(0.85, [item.amount for item in expense.line_items])

    def test_statement_merchant_can_rescue_poor_supplier_name(self):
        expense = Expense(
            source_file=Path("receipt.pdf"),
            expense_id="20260531_#6",
            date="2026-05-31",
            supplier_name="Tax ywvOl",
            expense_type="meal-dinner",
            amount=119.84,
            currency="AUD",
            raw_text="Farmers Daughte’s Total $119.84",
        )
        transaction = StatementTransaction(
            source_file=Path("statement.xls"),
            date="2026-05-31",
            description="ZLR*FARMERS DAUGHTER FA MELBOURNE",
            amount_cad=122.01,
            foreign_amount=119.84,
            foreign_currency="AUD",
            suggested_expense_id="20260531_#6",
            match_confidence=0.65,
        )
        enrich_expenses_from_statements([expense], [transaction])
        self.assertEqual(expense.supplier_name, "Farmers Daughters")

    def test_parse_csv_statement(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "statement.csv"
            with path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(
                    handle, fieldnames=["Date", "Description", "Amount", "Foreign Spend Amount"]
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "Date": "2026-06-01",
                        "Description": "FANCY HANKS MELBOURNE",
                        "Amount": "35.18",
                        "Foreign Spend Amount": "34.56 AUSTRALIAN DOLLAR",
                    }
                )
            rows = parse_csv_statement(path)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].date, "2026-06-01")
            self.assertEqual(rows[0].foreign_currency, "AUD")

    def test_match_score_uses_date_amount_and_supplier(self):
        expense = Expense(
            source_file=Path("receipt.pdf"),
            expense_id="EXP-1",
            date="2026-06-01",
            supplier_name="Fancy Hanks",
            amount=34.56,
            currency="AUD",
        )
        tx = StatementTransaction(
            source_file=Path("statement.csv"),
            date="2026-06-01",
            description="BBBQ - FANCY HANKS MELBOURNE",
            amount_cad=35.18,
            foreign_amount=34.56,
            foreign_currency="AUD",
        )
        self.assertGreaterEqual(match_score(expense, tx), 0.72)

    def test_same_date_and_amount_can_auto_match_with_a_weak_merchant_name(self):
        expense = Expense(
            source_file=Path("receipt.pdf"),
            expense_id="EXP-1",
            date="2026-07-10",
            supplier_name="Receipt merchant unreadable",
            amount=42.50,
            currency="CAD",
        )
        transaction = NormalizedTransaction(
            source_file=Path("card.csv"),
            source_row=2,
            provider="amex",
            transaction_group_id="TX-1",
            funding_leg_id="TX-1:1",
            transaction_date="2026-07-10",
            description="SQ *LOCAL PURCHASE",
            match_eligible=True,
            purchase_amount=42.50,
            purchase_currency="CAD",
            cad_amount=42.50,
            cad_completeness="complete",
        )

        match_normalized_transactions([expense], [transaction])

        self.assertEqual(transaction.expense_id, expense.expense_id)
        self.assertEqual(transaction.match_status, "auto")
        self.assertGreaterEqual(transaction.match_confidence, 0.72)

    def test_shared_receipt_auto_matches_the_per_employee_card_amount(self):
        expense = Expense(
            source_file=Path("receipt.pdf"),
            expense_id="EXP-SHARED",
            date="2026-07-10",
            supplier_name="Team Dinner",
            amount=314.0,
            currency="CAD",
            number_of_people=2,
        )
        transaction = NormalizedTransaction(
            source_file=Path("card.csv"),
            source_row=2,
            provider="amex",
            transaction_group_id="TX-SHARED",
            funding_leg_id="TX-SHARED:1",
            transaction_date="2026-07-10",
            description="TEAM DINNER",
            match_eligible=True,
            purchase_amount=157.0,
            purchase_currency="CAD",
            cad_amount=157.0,
            cad_completeness="complete",
        )

        match_normalized_transactions([expense], [transaction])

        self.assertEqual(transaction.expense_id, expense.expense_id)
        self.assertEqual(transaction.match_status, "auto")

    def test_same_restaurant_date_with_receipt_thirty_percent_lower_is_review_suggestion(self):
        expense = Expense(
            source_file=Path("receipt.pdf"),
            expense_id="EXP-1",
            date="2026-07-10",
            supplier_name="Bistro Montreal",
            amount=70.0,
            currency="CAD",
        )
        transaction = NormalizedTransaction(
            source_file=Path("card.csv"),
            source_row=2,
            provider="amex",
            transaction_group_id="TX-1",
            funding_leg_id="TX-1:1",
            transaction_date="2026-07-10",
            description="BISTRO MONTREAL",
            match_eligible=True,
            purchase_amount=100.0,
            purchase_currency="CAD",
            cad_amount=100.0,
            cad_completeness="complete",
        )

        match_normalized_transactions([expense], [transaction])

        self.assertEqual(transaction.suggested_expense_id, expense.expense_id)
        self.assertFalse(transaction.expense_id)
        self.assertEqual(transaction.match_status, "suggested")
        self.assertGreaterEqual(transaction.match_confidence, 0.72)
        self.assertIn("30% below the card total", transaction.match_review_reason)

    def test_split_amount_can_still_match_statement(self):
        self.assertTrue(close_split_amount(400.0, 100.0))
        expense = Expense(
            source_file=Path("receipt.pdf"),
            expense_id="EXP-1",
            date="2026-06-01",
            supplier_name="Fancy Hanks",
            amount=400.0,
            currency="AUD",
        )
        tx = StatementTransaction(
            source_file=Path("statement.csv"),
            date="2026-06-01",
            description="BBBQ - FANCY HANKS MELBOURNE",
            amount_cad=101.80,
            foreign_amount=100.0,
            foreign_currency="AUD",
        )
        self.assertGreaterEqual(match_score(expense, tx), 0.72)

    def test_workbook_formulas_and_empty_statement_tab(self):
        with tempfile.TemporaryDirectory() as tmp:
            trip = Path(tmp) / "202606_test"
            trip.mkdir()
            expense = Expense(
                source_file=Path("receipt.pdf"),
                expense_id="EXP-1",
                date="2026-06-01",
                supplier_name="Cafe",
                expense_type="meal",
                amount=10,
                currency="AUD",
                corrected_amount_in_currency=8,
                confidence=0.8,
            )
            transaction = StatementTransaction(
                source_file=Path("statement.csv"),
                date="2026-06-02",
                description="Unmatched",
                amount_cad=99,
            )
            output = build_workbook(trip, [expense], [transaction])
            wb = load_workbook(output, data_only=False)
            self.assertEqual(wb["expense_list"]["E1"].value, "amount_in_currency")
            self.assertEqual(wb["expense_list"]["F1"].value, "receipt_total_in_currency")
            self.assertEqual(wb["expense_list"]["G1"].value, "amount_check")
            self.assertEqual(wb["expense_list"]["I1"].value, "number_of_person")
            self.assertEqual(wb["expense_list"]["I2"].value, 1)
            self.assertEqual(
                wb["expense_list"]["E2"].value,
                '=IF(COUNTIFS(expense_line_items!$A:$A,$A2,expense_line_items!$F:$F,">0")=0,"",SUMIFS(expense_line_items!$F:$F,expense_line_items!$A:$A,$A2))',
            )
            self.assertEqual(
                wb["expense_list"]["G2"].value,
                '=IF($F2="","missing",IF($E2="","receipt total used",IF(ABS($E2-$F2)<=MAX(0.05,$F2*0.03),"ok","mismatch")))',
            )
            self.assertIn("COUNTIFS(expense_line_items!", wb["expense_list"]["J2"].value)
            self.assertIn("$F2/IF(OR($I2", wb["expense_list"]["J2"].value)
            self.assertEqual(
                wb["expense_list"]["K2"].value,
                '=IF($R2<>"",$R2,IF(COUNTIF(card_statements!$D:$D,$A2)=0,"",SUMIFS(card_statements!$C:$C,card_statements!$D:$D,$A2)))',
            )
            self.assertEqual(
                wb["expense_list"]["L2"].value,
                '=IF(OR($K2="",$K2=0),"",IF($H2="CAD",1,IF(OR($V2="",$V2=0),"",$K2/$V2)))',
            )
            self.assertIn('"statement_person_share"', wb["expense_list"]["M2"].value)
            self.assertIn('"statement_receipt_total"', wb["expense_list"]["M2"].value)
            self.assertEqual(wb["expense_list"]["Q1"].value, "include")
            self.assertEqual(wb["expense_list"]["R1"].value, "manual_CAD_override")
            self.assertEqual(wb["expense_list"]["T1"].value, "statement_purchase_amount")
            self.assertEqual(wb["expense_list"]["U1"].value, "statement_purchase_currency")
            self.assertEqual(wb["expense_list"]["V1"].value, "accounting_original_basis")
            self.assertEqual(wb["expense_list"]["W1"].value, "accounting_basis_status")
            self.assertIn('IF(OR($T2="",$T2=0),$F2', wb["expense_list"]["V2"].value)
            line_headers = [cell.value for cell in wb["expense_line_items"][1]]
            self.assertIn("alcohol_detection_confidence", line_headers)
            self.assertIn("alcohol_detection_reason", line_headers)
            self.assertIn("alcohol_matched_term", line_headers)
            cf_rules = list(wb["expense_list"].conditional_formatting["K2:K2"])
            self.assertEqual(len(cf_rules), 1)
            self.assertEqual(cf_rules[0].operator, "equal")
            self.assertEqual(cf_rules[0].formula, ["0"])
            self.assertEqual(wb["card_statements"]["A1"].value, "date")
            self.assertEqual(wb["expense_list"]["I2"].fill.fgColor.rgb, "00FFF2CC")
            self.assertEqual(wb["card_statements"]["D2"].fill.fgColor.rgb, "00C00000")

    def test_reviewed_receipt_date_updates_generated_expense_id(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202606_test")
            receipt = trip / "expenses_receipts" / "receipt.pdf"
            receipt.write_bytes(b"dummy")
            expense = Expense(
                source_file=receipt,
                expense_id="",
                date=None,
                supplier_name="Reviewed Cafe",
                amount=10,
                currency="AUD",
            )

            def apply_review(_trip, expenses, **_kwargs):
                expenses[0].date = "2026-06-03"
                return True

            captured = {}

            def build(trip_dir, expenses, transactions, **_kwargs):
                captured["expense_id"] = expenses[0].expense_id
                output = trip_dir / "review.xlsx"
                output.touch()
                return output

            with (
                patch("nlp_expenses.generator.parse_receipt", return_value=expense),
                patch("nlp_expenses.generator.apply_line_item_review", side_effect=apply_review),
                patch("nlp_expenses.generator.build_workbook", side_effect=build),
            ):
                output = generate_review(trip, root, llm_mode="off", mode="ivado")

            self.assertEqual(captured["expense_id"], "20260603_#1")
            self.assertEqual(output, (trip / "review.xlsx").resolve())

    @pytest.mark.integration
    def test_integration_existing_melbourne_receipt_count(self):
        root = Path(__file__).resolve().parents[1]
        trip = root / "trips" / "202606_melbourne"
        if not trip.exists():
            self.skipTest("fixture trip not present")
        with tempfile.TemporaryDirectory() as tmp:
            test_root = Path(tmp)
            test_trip = test_root / "trips" / trip.name
            test_trip.parent.mkdir()
            shutil.copytree(
                trip,
                test_trip,
                ignore=shutil.ignore_patterns("*.xlsx", "*.zip", "benchmark_*.json"),
            )
            output = generate_review(test_trip, test_root, llm_mode="off")
            wb = load_workbook(output, read_only=True)
            receipt_count = sum(
                1
                for path in (test_trip / "expenses_receipts").iterdir()
                if path.is_file() and path.suffix.lower() in SUPPORTED_RECEIPTS
            )
            self.assertEqual(wb["expense_list"].max_row - 1, receipt_count)


if __name__ == "__main__":
    unittest.main()
