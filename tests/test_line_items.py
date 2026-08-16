from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from nlp_expenses.line_items import (
    add_line_item,
    apply_line_item_review,
    ensure_line_item_review_ready,
    line_item_review_view,
    load_line_item_review_state,
    receipt_scan_status,
    remove_line_item,
    reset_receipt_review,
    save_line_item_review,
    save_line_item_review_state,
    set_expense_review,
    set_line_item_review,
    sync_line_item_review,
)
from nlp_expenses.models import Expense, LineItem, NormalizedTransaction
from nlp_expenses.trip_metadata import save_trip_metadata
from nlp_expenses.trips import ensure_trip
from nlp_expenses.workbook import build_arvine_workbook, build_workbook


def meal_expense(path: Path) -> Expense:
    return Expense(
        source_file=path,
        expense_id="MEAL-1",
        date="2026-07-01",
        supplier_name="Bistro",
        expense_type="meal",
        amount=115.0,
        currency="CAD",
        country="Canada",
        province="QC",
        gst_hst=5.0,
        qst=10.0,
        line_items=[
            LineItem(description="Dinner", amount=80.0, confidence=0.9),
            LineItem(
                description="French 75",
                amount=20.0,
                is_alcohol=True,
                included=False,
                confidence=0.9,
                alcohol_confidence=0.97,
                alcohol_reason="recognized cocktail",
                alcohol_matched_term="french 75",
            ),
            LineItem(
                description="GST/HST", amount=5.0, included=True, confidence=1.0, synthetic=True
            ),
            LineItem(description="QST", amount=10.0, included=True, confidence=1.0, synthetic=True),
        ],
    )


class LineItemReviewTests(unittest.TestCase):
    def test_incremental_scan_extracts_only_new_receipts_and_preserves_existing_review(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_incremental-scan", mode="arvine")
            first_path = trip / "expenses_receipts" / "first.pdf"
            first_path.write_bytes(b"first")
            with patch(
                "nlp_expenses.generator.parse_arvine_receipt",
                return_value=meal_expense(first_path),
            ):
                sync_line_item_review(trip, root, llm_mode="off", only_unscanned=True)

            set_expense_review(trip, first_path.name, {"currency": "QAR"})
            second_path = trip / "expenses_receipts" / "second.pdf"
            second_path.write_bytes(b"second")
            scan = receipt_scan_status(trip)
            self.assertEqual(scan["scanned_count"], 1)
            self.assertEqual(scan["unscanned_count"], 1)
            self.assertEqual(
                {item["source_file"]: item["status"] for item in scan["receipts"]},
                {"first.pdf": "scanned", "second.pdf": "not_scanned"},
            )

            second_expense = meal_expense(second_path)
            second_expense.supplier_name = "Second Bistro"
            with patch(
                "nlp_expenses.generator.parse_arvine_receipt",
                return_value=second_expense,
            ) as parser:
                sync_line_item_review(trip, root, llm_mode="off", only_unscanned=True)
            parser.assert_called_once()
            self.assertEqual(parser.call_args.args[0], second_path)

            review = line_item_review_view(trip)
            self.assertFalse(review["stale"])
            self.assertEqual(len(review["receipts"]), 2)
            first = next(item for item in review["receipts"] if item["source_file"] == "first.pdf")
            self.assertEqual(first["currency"], "QAR")
            self.assertIn("currency", first["overridden_fields"])

    def test_explicit_line_and_receipt_review_moves_expense_to_ready(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_ready-state", mode="arvine")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            save_line_item_review(trip, [meal_expense(receipt)])

            initial = line_item_review_view(trip)["receipts"][0]
            self.assertEqual(initial["status"], "ok")
            self.assertFalse(initial["reviewed"])
            self.assertEqual(line_item_review_view(trip)["summary"]["ok_count"], 1)
            first_line = initial["line_items"][0]

            one_line = set_line_item_review(
                trip,
                receipt.name,
                first_line["line_id"],
                {"reviewed": True},
            )["receipts"][0]
            self.assertEqual(one_line["line_items"][0]["status"], "ready")
            self.assertEqual(one_line["status"], "ok")

            ready = set_expense_review(trip, receipt.name, {"reviewed": True})["receipts"][0]
            self.assertEqual(ready["status"], "ready")
            self.assertTrue(ready["reviewed"])
            self.assertTrue(all(item["reviewed"] for item in ready["line_items"]))

            corrected = set_expense_review(trip, receipt.name, {"currency": "qar"})["receipts"][0]
            self.assertEqual(corrected["currency"], "QAR")
            self.assertEqual(corrected["status"], "ok")
            self.assertFalse(corrected["reviewed"])

            set_expense_review(trip, receipt.name, {"reviewed": True})
            changed = set_line_item_review(
                trip,
                receipt.name,
                first_line["line_id"],
                {"description": "Dinner corrected"},
            )["receipts"][0]
            self.assertEqual(changed["status"], "ok")
            self.assertFalse(changed["reviewed"])
            changed_line = next(
                item for item in changed["line_items"] if item["line_id"] == first_line["line_id"]
            )
            self.assertFalse(changed_line["reviewed"])

    def test_receipt_meal_classification_is_derived_from_editable_expense_type(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_meal-toggle", mode="arvine")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            save_line_item_review(trip, [meal_expense(receipt)])

            detected = line_item_review_view(trip)["receipts"][0]
            self.assertTrue(detected["is_meal"])
            self.assertEqual(detected["non_meal_expense_type"], "other")

            non_meal = set_expense_review(trip, receipt.name, {"expense_type": "transport"})[
                "receipts"
            ][0]
            self.assertFalse(non_meal["is_meal"])
            self.assertEqual(non_meal["non_meal_expense_type"], "transport")

            meal = set_expense_review(trip, receipt.name, {"expense_type": "meal"})["receipts"][0]
            self.assertTrue(meal["is_meal"])
            self.assertEqual(meal["non_meal_expense_type"], "transport")

    def test_expense_currency_must_come_from_iso_currency_list(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_currency-list", mode="arvine")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            save_line_item_review(trip, [meal_expense(receipt)])

            with self.assertRaisesRegex(ValueError, "highlighted expense fields"):
                set_expense_review(trip, receipt.name, {"currency": "ZZZ"})

            qar = set_expense_review(trip, receipt.name, {"currency": "qar"})["receipts"][0]
            self.assertEqual(qar["currency"], "QAR")

    def test_default_payer_override_survives_rescan_and_reset_restores_default(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_payer", mode="arvine")
            save_trip_metadata(
                trip,
                {
                    "claim_program": "arvine_only",
                    "default_paid_by": "arvine_corporate_bmo",
                    "payer_confirmed": True,
                },
            )
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            expense = meal_expense(receipt)
            save_line_item_review(trip, [expense])
            self.assertEqual(
                line_item_review_view(trip)["receipts"][0]["paid_by"],
                "company_card",
            )

            set_expense_review(trip, receipt.name, {"paid_by": "employee_personal"})
            save_line_item_review(trip, [meal_expense(receipt)])
            rescanned = line_item_review_view(trip)["receipts"][0]
            self.assertEqual(rescanned["paid_by"], "traveller_personal")
            self.assertTrue(rescanned["paid_by_overridden"])

            reset = reset_receipt_review(trip, receipt.name)["receipts"][0]
            self.assertEqual(reset["paid_by"], "company_card")
            self.assertFalse(reset["paid_by_overridden"])

            save_trip_metadata(
                trip,
                {
                    "claim_program": "arvine_only",
                    "default_paid_by": "employee_personal",
                    "payer_confirmed": True,
                },
            )
            updated_default = line_item_review_view(trip)["receipts"][0]
            self.assertEqual(updated_default["paid_by"], "traveller_personal")
            self.assertEqual(updated_default["auto_paid_by"], "traveller_personal")

    def test_whole_receipt_ivado_exclusion_keeps_a_reason(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_receipt_exclusion", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            save_line_item_review(trip, [meal_expense(receipt)])

            excluded = set_expense_review(
                trip,
                receipt.name,
                {
                    "included_in_arvine": True,
                    "included_in_ivado": False,
                    "ivado_exclusion_reason": "non_business",
                },
            )["receipts"][0]
            self.assertFalse(excluded["included_in_ivado"])
            self.assertEqual(excluded["ivado_exclusion_reason"], "non_business")

            restored = set_expense_review(
                trip,
                receipt.name,
                {
                    "included_in_ivado": True,
                    "ivado_exclusion_reason": "non_business",
                },
            )["receipts"][0]
            self.assertTrue(restored["included_in_ivado"])
            self.assertIsNone(restored["ivado_exclusion_reason"])

    def test_summary_excludes_zero_value_manual_alcohol_placeholder(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_summary", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            expense = meal_expense(receipt)
            expense.line_items.append(
                LineItem(
                    description="Alcohol adjustment - manual",
                    amount=0.0,
                    is_alcohol=True,
                    included=False,
                    alcohol_confidence=1.0,
                )
            )
            save_line_item_review(trip, [expense])

            summary = line_item_review_view(trip)["summary"]
            self.assertEqual(summary["alcohol_count"], 1)
            self.assertEqual(summary["excluded_count"], 1)
            self.assertEqual(summary["line_count"], 5)

    def test_ivado_deactivates_alcohol_until_user_reclassifies_it(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_ivado", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            save_line_item_review(trip, [meal_expense(receipt)])

            view = line_item_review_view(trip)
            cocktail = next(
                item
                for item in view["receipts"][0]["line_items"]
                if item["description"] == "French 75"
            )
            self.assertTrue(cocktail["is_alcohol"])
            self.assertFalse(cocktail["included"])
            self.assertEqual(view["summary"]["excluded_count"], 1)

            with self.assertRaisesRegex(ValueError, "Mark the line as not alcohol"):
                set_line_item_review(
                    trip,
                    receipt.name,
                    cocktail["line_id"],
                    {"included": True},
                )

            updated = set_line_item_review(
                trip,
                receipt.name,
                cocktail["line_id"],
                {"is_alcohol": False},
            )
            cocktail = next(
                item
                for item in updated["receipts"][0]["line_items"]
                if item["description"] == "French 75"
            )
            self.assertFalse(cocktail["is_alcohol"])
            self.assertTrue(cocktail["included"])
            self.assertTrue(cocktail["alcohol_overridden"])

            fresh = meal_expense(receipt)
            self.assertTrue(apply_line_item_review(trip, [fresh], require_fresh=True))
            applied = next(item for item in fresh.line_items if item.description == "French 75")
            self.assertFalse(applied.is_alcohol)
            self.assertTrue(applied.included)
            self.assertEqual(fresh.corrected_amount_in_currency, 115.0)

            reset = reset_receipt_review(trip, receipt.name)
            cocktail = next(
                item
                for item in reset["receipts"][0]["line_items"]
                if item["description"] == "French 75"
            )
            self.assertFalse(cocktail["included"])
            self.assertFalse(cocktail["inclusion_overridden"])

    def test_ivado_excludes_every_alcohol_line_regardless_of_detection_confidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_ivado-low-confidence", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            expense = meal_expense(receipt)
            expense.line_items = [
                LineItem(description="Dinner", amount=95.0),
                LineItem(
                    description="St-Ambroise",
                    amount=20.0,
                    is_alcohol=True,
                    included=True,
                    alcohol_confidence=0.74,
                    alcohol_reason="OpenAI classified the line as alcohol",
                ),
            ]
            expense.gst_hst = 0.0
            expense.qst = 0.0

            save_line_item_review(trip, [expense])
            reviewed = line_item_review_view(trip)["receipts"][0]
            beer = next(item for item in reviewed["line_items"] if item["is_alcohol"])
            self.assertTrue(beer["included_in_arvine"])
            self.assertFalse(beer["included_in_ivado"])
            self.assertEqual(beer["ivado_exclusion_reason"], "alcohol")
            self.assertEqual(reviewed["arvine_included_total"], 115.0)
            self.assertEqual(reviewed["ivado_included_total"], 95.0)

    def test_shared_receipt_exposes_full_and_per_employee_adjusted_totals(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_shared-totals", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            expense = Expense(
                source_file=receipt,
                expense_id="",
                expense_type="meal",
                amount=314.0,
                currency="CAD",
                number_of_people=2,
                line_items=[
                    LineItem(description="Food", amount=222.0),
                    LineItem(description="Alcohol", amount=92.0, is_alcohol=True),
                ],
            )

            save_line_item_review(trip, [expense])
            reviewed = line_item_review_view(trip)["receipts"][0]

            self.assertEqual(reviewed["per_person_receipt_total"], 157.0)
            self.assertEqual(reviewed["per_person_line_total"], 157.0)
            self.assertEqual(reviewed["per_person_arvine_included_total"], 157.0)
            self.assertEqual(reviewed["per_person_ivado_included_total"], 111.0)
            self.assertEqual(reviewed["per_person_ivado_excluded_total"], 46.0)

    def test_tax_fields_and_protected_tax_lines_stay_synchronized(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_tax-lines", mode="arvine")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            expense = Expense(
                source_file=receipt,
                expense_id="",
                expense_type="meal",
                amount=18.29,
                subtotal=15.90,
                line_items=[
                    LineItem(description="Pizza", amount=15.90),
                    LineItem(description="GST/HST", amount=0.80),
                    LineItem(description="QST", amount=1.59),
                ],
            )
            save_line_item_review(trip, [expense])

            reviewed = line_item_review_view(trip)["receipts"][0]
            tax_lines = {
                item["system_type"]: item for item in reviewed["line_items"] if item["system_type"]
            }
            self.assertEqual(set(tax_lines), {"gst_hst", "qst"})
            self.assertEqual(tax_lines["gst_hst"]["description"], "GST/HST")
            self.assertEqual(tax_lines["gst_hst"]["amount"], 0.80)
            self.assertEqual(tax_lines["qst"]["amount"], 1.59)

            updated = set_expense_review(trip, receipt.name, {"gst_hst": 1.00, "qst": ""})[
                "receipts"
            ][0]
            tax_lines = {
                item["system_type"]: item for item in updated["line_items"] if item["system_type"]
            }
            self.assertEqual(updated["gst_hst"], 1.00)
            self.assertEqual(updated["qst"], 0.0)
            self.assertEqual(tax_lines["gst_hst"]["amount"], 1.00)
            self.assertEqual(tax_lines["qst"]["amount"], 0.0)

            updated = set_line_item_review(
                trip,
                receipt.name,
                tax_lines["qst"]["line_id"],
                {"amount": 1.75},
            )["receipts"][0]
            self.assertEqual(updated["qst"], 1.75)
            with self.assertRaisesRegex(ValueError, "managed automatically"):
                set_line_item_review(
                    trip,
                    receipt.name,
                    tax_lines["gst_hst"]["line_id"],
                    {"description": "Sales tax"},
                )
            with self.assertRaisesRegex(ValueError, "cannot be removed"):
                remove_line_item(trip, receipt.name, tax_lines["gst_hst"]["line_id"])

            with self.assertRaisesRegex(ValueError, "Tax amount cannot be negative"):
                set_line_item_review(
                    trip,
                    receipt.name,
                    tax_lines["gst_hst"]["line_id"],
                    {"amount": -1.0},
                )

    def test_promotions_can_be_added_and_edited_as_negative_line_items(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_promotion", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            expense = Expense(
                source_file=receipt,
                expense_id="",
                expense_type="meal",
                amount=10.0,
                currency="CAD",
                line_items=[LineItem(description="Food and fees", amount=15.0)],
            )
            save_line_item_review(trip, [expense])

            added = add_line_item(
                trip,
                receipt.name,
                "Promotion",
                -5.0,
                included=True,
                is_alcohol=True,
            )["receipts"][0]
            promotion = next(
                item for item in added["line_items"] if item["description"] == "Promotion"
            )

            self.assertEqual(promotion["amount"], -5.0)
            self.assertFalse(promotion["is_alcohol"])
            self.assertTrue(promotion["included_in_ivado"])
            self.assertEqual(added["line_total"], 10.0)
            self.assertTrue(added["reconciled"])

            edited = set_line_item_review(
                trip,
                receipt.name,
                promotion["line_id"],
                {"amount": -4.5},
            )["receipts"][0]
            promotion = next(
                item for item in edited["line_items"] if item["description"] == "Promotion"
            )
            self.assertEqual(promotion["amount"], -4.5)

            with self.assertRaisesRegex(ValueError, "valid line-item amount"):
                add_line_item(trip, receipt.name, "Invalid", "nan")

    def test_corrected_subtotal_updates_unchanged_ocr_receipt_total(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_corrected-total", mode="arvine")
            receipt = trip / "expenses_receipts" / "cafe.pdf"
            receipt.write_bytes(b"cafe")
            expense = Expense(
                source_file=receipt,
                expense_id="",
                expense_type="meal",
                amount=67.0,
                subtotal=67.0,
                currency="QAR",
                line_items=[LineItem(description="Cafe purchase", amount=67.0)],
            )
            save_line_item_review(trip, [expense])

            corrected = set_expense_review(
                trip,
                receipt.name,
                {"amount": 67.0, "subtotal": 57.0, "gst_hst": 0, "qst": 0},
            )["receipts"][0]

            self.assertEqual(corrected["subtotal"], 57.0)
            self.assertEqual(corrected["amount"], 57.0)
            self.assertEqual(corrected["receipt_total"], 57.0)
            self.assertEqual(corrected["field_overrides"]["amount"], 57.0)

            explicitly_distinct = set_expense_review(
                trip,
                receipt.name,
                {"amount": 60.0, "subtotal": 55.0},
            )["receipts"][0]
            self.assertEqual(explicitly_distinct["amount"], 60.0)

    def test_every_expense_has_zero_default_tax_lines(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_zero-tax-lines", mode="arvine")
            receipt = trip / "expenses_receipts" / "hotel.pdf"
            receipt.write_bytes(b"hotel")
            save_line_item_review(
                trip,
                [
                    Expense(
                        source_file=receipt,
                        expense_id="",
                        date="2026-07-01",
                        supplier_name="Hotel",
                        expense_type="hotel",
                        amount=200,
                        currency="CAD",
                        line_items=[LineItem(description="Hotel stay", amount=200)],
                    )
                ],
            )

            reviewed = line_item_review_view(trip)["receipts"][0]
            tax_lines = [item for item in reviewed["line_items"] if item["system_type"]]
            self.assertEqual(reviewed["status"], "ok")
            self.assertEqual(
                [(item["description"], item["amount"]) for item in tax_lines],
                [("GST/HST", 0.0), ("QST", 0.0)],
            )

    def test_non_meal_receipt_without_saved_purchase_lines_gets_a_balancing_subtotal(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_non-meal-gap", mode="arvine")
            receipt = trip / "expenses_receipts" / "other.pdf"
            receipt.write_bytes(b"other")
            save_line_item_review(
                trip,
                [
                    Expense(
                        source_file=receipt,
                        expense_id="",
                        date="2026-07-24",
                        supplier_name="DAOUST DE 210",
                        expense_type="other",
                        amount=30.45,
                        currency="CAD",
                    )
                ],
            )

            reviewed = line_item_review_view(trip)["receipts"][0]
            self.assertEqual(reviewed["line_total"], 30.45)
            self.assertEqual(reviewed["difference"], 0.0)
            self.assertTrue(reviewed["reconciled"])
            self.assertEqual(reviewed["automatic_status"], "not_applicable")
            self.assertEqual(reviewed["status"], "ok")
            self.assertEqual(
                [item["description"] for item in reviewed["line_items"]],
                ["Receipt subtotal", "GST/HST", "QST"],
            )

    def test_legacy_uber_tax_only_scan_uses_saved_subtotal_without_rescanning(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_uber", mode="company")
            receipt = trip / "expenses_receipts" / "uber.pdf"
            receipt.write_bytes(b"uber")
            save_line_item_review(
                trip,
                [
                    Expense(
                        source_file=receipt,
                        expense_id="",
                        date="2026-07-01",
                        supplier_name="Uber",
                        expense_type="transport",
                        amount=16.56,
                        subtotal=14.40,
                        gst_hst=0.72,
                        qst=1.44,
                        currency="CAD",
                    )
                ],
            )

            reviewed = line_item_review_view(trip)["receipts"][0]

            self.assertEqual(reviewed["line_total"], 16.56)
            self.assertEqual(reviewed["difference"], 0.0)
            self.assertEqual(
                [(item["description"], item["amount"]) for item in reviewed["line_items"]],
                [("Receipt subtotal", 14.40), ("GST/HST", 0.72), ("QST", 1.44)],
            )

    def test_removed_receipt_does_not_lock_edits_for_remaining_receipts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_removed-receipt", mode="arvine")
            first = trip / "expenses_receipts" / "first.pdf"
            second = trip / "expenses_receipts" / "second.pdf"
            first.write_bytes(b"first")
            second.write_bytes(b"second")
            first_expense = meal_expense(first)
            first_expense.expense_id = "MEAL-FIRST"
            second_expense = meal_expense(second)
            second_expense.expense_id = "MEAL-SECOND"
            save_line_item_review(trip, [first_expense, second_expense])

            second.unlink()
            view = line_item_review_view(trip)
            self.assertFalse(view["stale"])
            self.assertEqual([receipt["source_file"] for receipt in view["receipts"]], [first.name])

            dinner = next(
                item
                for item in view["receipts"][0]["line_items"]
                if item["description"] == "Dinner"
            )
            updated = set_line_item_review(
                trip,
                first.name,
                dinner["line_id"],
                {"amount": 79.0},
            )
            self.assertEqual(
                next(
                    item
                    for item in updated["receipts"][0]["line_items"]
                    if item["line_id"] == dinner["line_id"]
                )["amount"],
                79.0,
            )
            removed = remove_line_item(trip, first.name, dinner["line_id"])
            self.assertNotIn(
                dinner["line_id"],
                [item["line_id"] for item in removed["receipts"][0]["line_items"]],
            )

    def test_legacy_review_repairs_arvine_and_low_confidence_alcohol_flags(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_legacy-program-flags", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            expense = meal_expense(receipt)
            expense.line_items[1].alcohol_confidence = 0.74
            save_line_item_review(trip, [expense])
            state = load_line_item_review_state(trip)
            state["version"] = 2
            stored_receipt = state["receipts"][0]
            stored_receipt["included_in_arvine"] = False
            stored_alcohol = next(
                item for item in stored_receipt["line_items"] if item["is_alcohol"]
            )
            stored_alcohol["included"] = True
            stored_alcohol["included_in_arvine"] = False
            stored_alcohol["included_in_ivado"] = True
            stored_alcohol["ivado_exclusion_reason"] = None
            save_line_item_review_state(trip, state)

            reviewed = line_item_review_view(trip)["receipts"][0]
            alcohol = next(item for item in reviewed["line_items"] if item["is_alcohol"])
            self.assertTrue(reviewed["included_in_arvine"])
            self.assertTrue(alcohol["included_in_arvine"])
            self.assertFalse(alcohol["included_in_ivado"])
            self.assertEqual(alcohol["ivado_exclusion_reason"], "alcohol")

    def test_arvine_clears_alcohol_metadata_and_rejects_line_exclusion(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_arvine", mode="arvine")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            save_line_item_review(trip, [meal_expense(receipt)])
            view = line_item_review_view(trip)
            cocktail = next(
                item
                for item in view["receipts"][0]["line_items"]
                if item["description"] == "French 75"
            )
            self.assertTrue(cocktail["included"])
            self.assertFalse(cocktail["is_alcohol"])
            self.assertEqual(cocktail["alcohol_reason"], "")

            with self.assertRaisesRegex(ValueError, "company report includes every receipt line"):
                set_line_item_review(
                    trip,
                    receipt.name,
                    cocktail["line_id"],
                    {"included": False, "note": "Not reimbursable"},
                )
            receipt_view = line_item_review_view(trip)["receipts"][0]
            self.assertEqual(receipt_view["claimable_ratio"], 1.0)
            self.assertEqual(receipt_view["included_total"], 115.0)
            ensure_line_item_review_ready(trip)

    def test_user_line_choice_survives_small_openai_wording_change(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_ivado-rescan", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            first = meal_expense(receipt)
            first.line_items = [
                LineItem(
                    description="Balter Mpe,Ping",
                    amount=17.90,
                    is_alcohol=True,
                    alcohol_confidence=0.95,
                ),
                LineItem(description="Chicken Peis", amount=34.50),
            ]
            save_line_item_review(trip, [first], llm_mode="required")
            reviewed = line_item_review_view(trip)
            alcohol = next(
                item for item in reviewed["receipts"][0]["line_items"] if item["is_alcohol"]
            )
            set_line_item_review(
                trip,
                receipt.name,
                alcohol["line_id"],
                {"is_alcohol": False},
            )

            second = meal_expense(receipt)
            second.line_items = [
                LineItem(
                    description="Balter Mpe, Ping 1",
                    amount=17.90,
                    is_alcohol=True,
                    alcohol_confidence=0.95,
                ),
                LineItem(description="Chicken Peis", amount=34.50),
            ]
            save_line_item_review(trip, [second], llm_mode="required")
            rescanned = line_item_review_view(trip)
            corrected = next(
                item
                for item in rescanned["receipts"][0]["line_items"]
                if item["description"].startswith("Balter")
            )
            self.assertFalse(corrected["is_alcohol"])
            self.assertTrue(corrected["included_in_ivado"])
            self.assertTrue(corrected["alcohol_overridden"])

    def test_manually_added_line_survives_rescan(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_manual-line-rescan", mode="arvine")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            expense = meal_expense(receipt)
            save_line_item_review(trip, [expense], llm_mode="required")
            add_line_item(trip, receipt.name, "Tip confirmed manually", 5.0)

            save_line_item_review(trip, [meal_expense(receipt)], llm_mode="required")
            descriptions = {
                item["description"]
                for item in line_item_review_view(trip)["receipts"][0]["line_items"]
            }
            self.assertIn("Tip confirmed manually", descriptions)

    def test_receipt_source_change_makes_review_stale(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_stale", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"first")
            save_line_item_review(trip, [meal_expense(receipt)])
            self.assertFalse(line_item_review_view(trip)["stale"])
            receipt.write_bytes(b"changed")
            self.assertTrue(line_item_review_view(trip)["stale"])
            with self.assertRaisesRegex(ValueError, "current receipt line items"):
                ensure_line_item_review_ready(trip)

    def test_unreconciled_receipt_blocks_generation_only_when_a_line_is_excluded(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            trip = ensure_trip(root, "202607_gap", mode="ivado")
            receipt = trip / "expenses_receipts" / "meal.pdf"
            receipt.write_bytes(b"meal")
            expense = meal_expense(receipt)
            expense.line_items[-1] = LineItem(
                description="Unreconciled meal item - review",
                amount=10.0,
                included=True,
                synthetic=True,
            )
            save_line_item_review(trip, [expense])
            view = line_item_review_view(trip)
            self.assertEqual(view["receipts"][0]["status"], "review")
            self.assertTrue(view["receipts"][0]["blocking"])
            with self.assertRaisesRegex(ValueError, "Resolve line-item totals"):
                ensure_line_item_review_ready(trip)

    def test_workbooks_use_ivado_inclusion_and_arvine_full_claim(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            ivado = ensure_trip(root, "202607_ivado-book", mode="ivado")
            ivado_receipt = ivado / "expenses_receipts" / "meal.pdf"
            ivado_receipt.write_bytes(b"meal")
            ivado_expense = meal_expense(ivado_receipt)
            ivado_expense.expense_id = "MEAL-1"
            ivado_output = build_workbook(ivado, [ivado_expense], [])
            ivado_book = load_workbook(ivado_output, data_only=False)
            line_headers = [cell.value for cell in ivado_book["expense_line_items"][1]]
            self.assertEqual(line_headers[7:9], ["is_alcohol", "included"])
            self.assertIn("expense_line_items!$I:$I,TRUE", ivado_book["expense_list"]["J2"].value)

            arvine = ensure_trip(root, "202607_arvine-book", mode="arvine")
            arvine_receipt = arvine / "expenses_receipts" / "meal.pdf"
            arvine_receipt.write_bytes(b"meal")
            expense = meal_expense(arvine_receipt)
            save_line_item_review(arvine, [expense])
            view = line_item_review_view(arvine)
            cocktail = next(
                item
                for item in view["receipts"][0]["line_items"]
                if item["description"] == "French 75"
            )
            with self.assertRaisesRegex(ValueError, "company report includes every receipt line"):
                set_line_item_review(
                    arvine,
                    arvine_receipt.name,
                    cocktail["line_id"],
                    {"included": False},
                )
            apply_line_item_review(arvine, [expense], require_fresh=True)
            statement = arvine / "card_statements" / "card.csv"
            statement.write_text("", encoding="utf-8")
            transaction = NormalizedTransaction(
                source_file=statement,
                source_row=2,
                provider="amex",
                transaction_group_id="G1",
                funding_leg_id="G1:1",
                transaction_date="2026-07-01",
                description="BISTRO",
                match_eligible=True,
                purchase_amount=115.0,
                purchase_currency="CAD",
                settlement_amount=115.0,
                settlement_currency="CAD",
                cad_amount=115.0,
                cad_completeness="complete",
            )
            output = build_arvine_workbook(arvine, [expense], [transaction])
            book = load_workbook(output, data_only=False)
            detail = book["expense_detail"]
            self.assertIn("expense_line_items!$I:$I,TRUE", detail["AM2"].value)
            self.assertEqual(detail["AN2"].value, '=IF($AL2="","",$AL2-$AM2)')
            self.assertIn("$AM2/$AL2", detail["AO2"].value)
            self.assertIn("$AO2", detail["U2"].value)
            self.assertIn("$AO2", detail["W2"].value)
            self.assertEqual(book.sheetnames[-1], "expense_line_items")


if __name__ == "__main__":
    unittest.main()
