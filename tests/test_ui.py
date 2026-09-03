from __future__ import annotations

import json
import os
import re
import stat
import tempfile
import threading
import time
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from nlp_expenses.generator import generate_review
from nlp_expenses.jobs import JobManager
from nlp_expenses.lifecycle import record_generated_workbook
from nlp_expenses.line_items import save_line_item_review
from nlp_expenses.models import Expense, LineItem
from nlp_expenses.reconciliation_state import (
    load_reconciliation_state,
    save_reconciliation_state,
)
from nlp_expenses.trip_metadata import save_trip_metadata
from nlp_expenses.trips import ensure_trip, trip_mode
from nlp_expenses.ui import create_app, open_local_url
from nlp_expenses.workbook import build_workbook


class UITests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.app = create_app(self.root, access_token="test-token")
        self.app.testing = True
        self.client = self.app.test_client()
        response = self.client.get("/?token=test-token")
        self.assertEqual(response.status_code, 302)
        with self.client.session_transaction() as session:
            self.csrf = session["csrf_token"]
        self.headers = {"X-CSRF-Token": self.csrf}

    def tearDown(self):
        self.temp.cleanup()

    def complete_trip_metadata(self, trip: Path) -> None:
        save_trip_metadata(
            trip,
            {
                "claim_program": "ivado_sponsored" if trip_mode(trip) == "ivado" else "arvine_only",
                "traveller": "Florent",
                "company": "Example Corp.",
                "sponsor": "IVADO Labs" if trip_mode(trip) == "ivado" else "",
                "start_date": "2026-07-01",
                "end_date": "2026-07-03",
                "business_purpose": "Client workshop",
                "approver": "Manager",
                "payment_method": "Personal card reimbursement",
                "default_paid_by": "employee_personal",
                "payer_confirmed": True,
            },
        )

    def test_launcher_token_and_csrf_are_required(self):
        other = create_app(self.root, access_token="secret").test_client()
        self.assertEqual(other.get("/").status_code, 403)
        other.get("/?token=secret")
        self.assertEqual(
            other.post(
                "/api/trips", json={"month": "2026-07", "description": "Montreal"}
            ).status_code,
            403,
        )

    def test_mac_browser_launch_uses_native_open(self):
        with (
            patch("nlp_expenses.ui.sys.platform", "darwin"),
            patch("nlp_expenses.ui.shutil.which", return_value="/usr/bin/open"),
            patch("nlp_expenses.ui.subprocess.run") as native_open,
            patch("nlp_expenses.ui.webbrowser.open") as generic_open,
        ):
            open_local_url("http://127.0.0.1:8765/?token=test")
        native_open.assert_called_once_with(
            ["open", "http://127.0.0.1:8765/?token=test"],
            check=False,
        )
        generic_open.assert_not_called()

    def test_raw_template_has_direct_open_guard(self):
        template = (
            Path(__file__).parents[1] / "nlp_expenses" / "templates" / "index.html"
        ).read_text(encoding="utf-8")
        self.assertIn('window.location.protocol === "file:"', template)
        self.assertIn("This is the application template, not the running app", template)

    def test_legacy_company_values_render_with_neutral_product_language(self):
        trip = ensure_trip(self.root, "202607_company-trip", mode="arvine")
        save_trip_metadata(
            trip,
            {
                "claim_program": "arvine_only",
                "traveller": "Colleague",
                "default_paid_by": "arvine_corporate_bmo",
            },
        )

        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)

        self.assertNotIn("Arvine", html)
        self.assertIn("Own-company reimbursement", html)
        self.assertIn("Default payment source", html)
        self.assertIn("Company funds/card", html)

    def test_upload_cards_reveal_folders_and_best_quality_is_disabled_without_key(self):
        trip = ensure_trip(self.root, "202607_montreal", mode="arvine")
        with patch.dict(os.environ, {"OPENAI_API_KEY": ""}, clear=False):
            response = self.client.get(f"/?trip={trip.name}")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertEqual(html.count("Reveal folder in Finder"), 2)
        self.assertIn('data-target="receipts"', html)
        self.assertIn('data-target="statements"', html)
        self.assertEqual(html.count('value="best"'), 1)
        self.assertIn("Receipt extraction method", html)
        self.assertIn("API charges may apply", html)
        self.assertIn(
            "Public weekly FX data is fetched only when no exact CAD amount exists",
            html,
        )
        self.assertIn("Download standard CSV", html)
        self.assertIn("/static/standard-statement-template.csv", html)
        self.assertNotIn('name="quality"', html)
        self.assertRegex(
            html, r'<input type="radio" name="receipt-quality" value="best"[^>]*disabled'
        )
        self.assertIn(str((self.root / ".env").resolve()), html)

    def test_standard_statement_template_is_downloadable(self):
        response = self.client.get("/static/standard-statement-template.csv")
        try:
            self.assertEqual(response.status_code, 200)
            self.assertEqual(
                response.get_data(as_text=True).strip(),
                (
                    "transaction_date,description,purchase_amount,purchase_currency,"
                    "cad_amount,transaction_type,posted_date,account,cardholder,category"
                ),
            )
        finally:
            response.close()

    def test_best_quality_controls_are_enabled_when_key_is_configured(self):
        trip = ensure_trip(self.root, "202607_montreal", mode="arvine")
        (self.root / ".env").write_text("OPENAI_API_KEY=sk-test-ui\n", encoding="utf-8")
        with patch.dict(os.environ, {"OPENAI_API_KEY": "sk-test-ui"}, clear=False):
            response = self.client.get(f"/?trip={trip.name}")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn("OpenAI key saved", html)
        self.assertRegex(
            html, r'<input type="radio" name="receipt-quality" value="best"[^>]*checked'
        )

    def test_file_state_endpoint_detects_files_added_outside_the_ui(self):
        trip = ensure_trip(self.root, "202607_montreal", mode="arvine")
        initial = self.client.get(f"/api/trips/{trip.name}/file-state")
        self.assertEqual(initial.status_code, 200)
        initial_payload = initial.get_json()
        self.assertFalse(initial_payload["busy"])
        self.assertEqual(initial_payload["file_state"]["receipts"], 0)

        (trip / "expenses_receipts" / "finder-added.pdf").write_bytes(b"%PDF-finder")
        updated = self.client.get(f"/api/trips/{trip.name}/file-state")
        self.assertEqual(updated.status_code, 200)
        updated_state = updated.get_json()["file_state"]
        self.assertEqual(updated_state["receipts"], 1)
        self.assertNotEqual(updated_state["signature"], initial_payload["file_state"]["signature"])

        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Refresh files", html)
        self.assertIn("finder-added.pdf", html)

    def test_create_upload_validate_and_remove(self):
        created = self.client.post(
            "/api/trips",
            json={
                "month": "2026-07",
                "description": "Montreal",
                "claim_program": "arvine_only",
            },
            headers=self.headers,
        )
        self.assertEqual(created.status_code, 201)
        trip_name = created.get_json()["trip"]["name"]
        statement = b"Date,Date Processed,Description,Card Member,Account,Amount\n2026-07-01,2026-07-02,Cafe,A Person,1234,12.50\n"
        uploaded = self.client.post(
            f"/api/trips/{trip_name}/upload/statements",
            data={"files": (BytesIO(statement), "amex.csv")},
            content_type="multipart/form-data",
            headers=self.headers,
        )
        self.assertEqual(uploaded.status_code, 200)
        validation = uploaded.get_json()["trip"]["statements"][0]["validation"]
        self.assertEqual(validation["provider"], "amex")
        self.assertEqual(validation["rows_normalized"], 1)

        removed = self.client.post(
            f"/api/trips/{trip_name}/remove-file",
            json={"kind": "statements", "filename": "amex.csv"},
            headers=self.headers,
        )
        self.assertEqual(removed.status_code, 200)
        self.assertEqual(removed.get_json()["trip"]["statements"], [])

        sponsored = self.client.post(
            "/api/trips",
            json={
                "month": "2026-08",
                "description": "Sponsored Montreal",
                "claim_program": "ivado_sponsored",
                "metadata": {"claim_program": "ivado_sponsored"},
            },
            headers=self.headers,
        )
        self.assertEqual(sponsored.status_code, 201)
        sponsored_trip = sponsored.get_json()["trip"]
        self.assertEqual(sponsored_trip["mode"], "ivado")
        self.assertEqual(sponsored_trip["metadata"]["sponsor"], "IVADO Labs")

    def test_removing_a_receipt_keeps_remaining_line_controls_editable(self):
        trip = ensure_trip(self.root, "202607_remove-receipt-review", mode="arvine")
        first = trip / "expenses_receipts" / "first.pdf"
        second = trip / "expenses_receipts" / "second.pdf"
        first.write_bytes(b"first")
        second.write_bytes(b"second")
        expenses = [
            Expense(
                source_file=first,
                expense_id="FIRST",
                date="2026-07-01",
                supplier_name="First Cafe",
                expense_type="meal",
                amount=20,
                currency="CAD",
                line_items=[LineItem(description="Lunch", amount=20)],
            ),
            Expense(
                source_file=second,
                expense_id="SECOND",
                date="2026-07-02",
                supplier_name="Second Cafe",
                expense_type="meal",
                amount=25,
                currency="CAD",
                line_items=[LineItem(description="Dinner", amount=25)],
            ),
        ]
        save_line_item_review(trip, expenses)

        removed = self.client.post(
            f"/api/trips/{trip.name}/remove-file",
            json={"kind": "receipts", "filename": second.name},
            headers=self.headers,
        )
        self.assertEqual(removed.status_code, 200)
        review = removed.get_json()["trip"]["line_item_review"]
        self.assertFalse(review["stale"])
        self.assertEqual([receipt["source_file"] for receipt in review["receipts"]], [first.name])

        first_line = review["receipts"][0]["line_items"][0]
        edited = self.client.post(
            f"/api/trips/{trip.name}/line-items/item",
            json={
                "source_file": first.name,
                "line_id": first_line["line_id"],
                "fields": {"amount": 19.0},
            },
            headers=self.headers,
        )
        self.assertEqual(edited.status_code, 200)
        deleted = self.client.post(
            f"/api/trips/{trip.name}/line-items/remove",
            json={"source_file": first.name, "line_id": first_line["line_id"]},
            headers=self.headers,
        )
        self.assertEqual(deleted.status_code, 200)

    def test_delete_trip_requires_exact_confirmation_and_clears_all_trip_data(self):
        trip = ensure_trip(self.root, "202607_delete-from-ui", mode="arvine")
        receipt = trip / "expenses_receipts" / "nested" / "receipt.pdf"
        receipt.parent.mkdir(parents=True)
        receipt.write_bytes(b"receipt")
        (trip / "card_statements" / "statement.csv").write_text(
            "date,description,amount,currency\n2026-07-01,Cafe,10,CAD\n",
            encoding="utf-8",
        )
        (trip / "expense_review.xlsx").write_bytes(b"workbook")
        (trip / "trip_package.zip").write_bytes(b"package")

        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Delete trip and all data", html)
        self.assertIn("Permanently delete trip", html)
        self.assertIn(f"Type <code>{trip.name}</code> to confirm", html)

        rejected = self.client.delete(
            f"/api/trips/{trip.name}",
            json={"confirmation": "wrong-trip"},
            headers=self.headers,
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertTrue(trip.is_dir())

        deleted = self.client.delete(
            f"/api/trips/{trip.name}",
            json={"confirmation": trip.name},
            headers=self.headers,
        )
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(deleted.get_json()["deleted"], trip.name)
        self.assertFalse(trip.exists())
        self.assertEqual(
            self.client.get(f"/api/trips/{trip.name}").status_code,
            404,
        )

    def test_trip_metadata_policy_and_lifecycle_endpoints(self):
        trip = ensure_trip(self.root, "202607_lifecycle", mode="ivado")
        metadata = {
            "claim_program": "ivado_sponsored",
            "traveller": "Florent",
            "company": "Example Inc.",
            "sponsor": "IVADO Labs",
            "start_date": "2026-07-01",
            "end_date": "2026-07-02",
            "business_purpose": "Client workshop",
            "approver": "Reviewer",
            "payment_method": "Employee reimbursement",
            "default_paid_by": "employee_personal",
            "payer_confirmed": True,
            "policy": {
                "allowed_categories": ["flight", "hotel", "transport", "meal", "other"],
                "meal_limit_cad": 75,
            },
        }
        updated = self.client.post(
            f"/api/trips/{trip.name}/metadata",
            json={"metadata": metadata},
            headers=self.headers,
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(
            updated.get_json()["trip"]["metadata"]["business_purpose"], "Client workshop"
        )

        exception = self.client.post(
            f"/api/trips/{trip.name}/policy-exception",
            json={"warning_id": "meal_limit:meal.pdf", "note": "Approved client dinner."},
            headers=self.headers,
        )
        self.assertEqual(exception.status_code, 200)
        self.assertEqual(
            exception.get_json()["trip"]["metadata"]["policy_exceptions"]["meal_limit:meal.pdf"],
            "Approved client dinner.",
        )

        receipt = trip / "expenses_receipts" / "receipt.pdf"
        receipt.write_bytes(b"fixture")
        expense = Expense(
            source_file=receipt,
            expense_id="EXP-1",
            date="2026-07-01",
            supplier_name="Cafe",
            expense_type="meal",
            amount=20,
            currency="CAD",
        )
        workbook = trip / "expense_review_202607_lifecycle_ivado_20260725-120000.xlsx"
        build_workbook(trip, [expense], [], output_path=workbook)
        record_generated_workbook(self.root, trip, workbook)
        approved = self.client.post(
            f"/api/trips/{trip.name}/approve",
            json={
                "workbook": workbook.name,
                "reviewer": "Reviewer",
                "note": "All documents reviewed.",
            },
            headers=self.headers,
        )
        self.assertEqual(approved.status_code, 200)
        self.assertTrue(approved.get_json()["approval"]["current"])

        exported = self.client.post(
            f"/api/trips/{trip.name}/export-package",
            json={},
            headers=self.headers,
        )
        self.assertEqual(exported.status_code, 200)
        package_name = exported.get_json()["package"]["name"]
        downloaded = self.client.get(
            f"/api/trips/{trip.name}/download-package?filename={package_name}"
        )
        self.assertEqual(downloaded.status_code, 200)
        self.assertTrue(downloaded.data.startswith(b"PK"))
        downloaded.close()

        archived = self.client.post(
            f"/api/trips/{trip.name}/archive",
            json={"archived": True},
            headers=self.headers,
        )
        self.assertEqual(archived.status_code, 200)
        self.assertEqual(archived.get_json()["trip"]["lifecycle"]["status"], "archived")

    def test_ambiguous_statement_date_convention_can_be_selected(self):
        trip = ensure_trip(self.root, "202607_ambiguous-date", mode="arvine")
        statement = b"Date,Description,Amount,Currency\n07/01/2026,Ambiguous purchase,10,CAD\n"
        uploaded = self.client.post(
            f"/api/trips/{trip.name}/upload/statements",
            data={"files": (BytesIO(statement), "ambiguous.csv")},
            content_type="multipart/form-data",
            headers=self.headers,
        )
        self.assertEqual(uploaded.status_code, 200)
        validation = uploaded.get_json()["trip"]["statements"][0]["validation"]
        self.assertTrue(validation["date_convention_required"])
        self.assertTrue(validation["errors"])

        selected = self.client.post(
            f"/api/trips/{trip.name}/statement-date-convention",
            json={"filename": "ambiguous.csv", "convention": "month_first"},
            headers=self.headers,
        )
        self.assertEqual(selected.status_code, 200)
        validation = selected.get_json()["trip"]["statements"][0]["validation"]
        self.assertFalse(validation["date_convention_required"])
        self.assertEqual(validation["date_convention"], "month_first")
        self.assertIn("2026-07-01", validation["date_samples"][0])

    def test_ambiguous_statement_columns_can_be_mapped_and_previewed(self):
        trip = ensure_trip(self.root, "202607_statement-mapping", mode="arvine")
        statement = (
            b"Example Card export\n"
            b"When,Merchant Label,Txn Value,ISO Code\n"
            b"13/07/2026,TRAIN,-45.20,CAD\n"
        )
        uploaded = self.client.post(
            f"/api/trips/{trip.name}/upload/statements",
            data={"files": (BytesIO(statement), "custom.csv")},
            content_type="multipart/form-data",
            headers=self.headers,
        )
        self.assertEqual(uploaded.status_code, 200)
        validation = uploaded.get_json()["trip"]["statements"][0]["validation"]
        self.assertTrue(validation["mapping_required"])
        self.assertTrue(validation["errors"])
        self.assertEqual(validation["header_row"], 2)

        mapped = self.client.post(
            f"/api/trips/{trip.name}/statement-import-profile",
            json={
                "filename": "custom.csv",
                "header_row": 2,
                "mapping": {
                    "transaction_date": 0,
                    "description": 1,
                    "amount": 2,
                    "settlement_currency": 3,
                },
                "sign_convention": "negative_purchase",
                "date_convention": "day_first",
            },
            headers=self.headers,
        )
        self.assertEqual(mapped.status_code, 200)
        validation = mapped.get_json()["trip"]["statements"][0]["validation"]
        self.assertFalse(validation["errors"])
        self.assertTrue(validation["profile_reused"])
        self.assertEqual(validation["preview"][0]["description"], "TRAIN")
        self.assertEqual(validation["preview"][0]["settlement_amount"], 45.20)

        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Review statement mapping", html)
        self.assertIn("saved mapping reused", html)
        self.assertIn("Row 3", html)

    def test_generation_gates_and_progress_polling(self):
        trip = ensure_trip(self.root, "202607_montreal", mode="arvine")
        no_receipt = self.client.post(
            f"/api/trips/{trip.name}/generate",
            json={"quality": "basic", "statements_complete": True},
            headers=self.headers,
        )
        self.assertEqual(no_receipt.status_code, 400)

        receipt = trip / "expenses_receipts" / "receipt.pdf"
        receipt.write_bytes(b"fixture")
        missing_confirmation = self.client.post(
            f"/api/trips/{trip.name}/generate",
            json={"quality": "basic", "statements_complete": False},
            headers=self.headers,
        )
        self.assertEqual(missing_confirmation.status_code, 400)

        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Cafe",
            expense_type="meal",
            amount=12.5,
            currency="CAD",
        )
        save_line_item_review(trip, [expense])
        self.complete_trip_metadata(trip)
        finalized = self.client.post(
            f"/api/trips/{trip.name}/finalize",
            json={"statements_complete": True},
            headers=self.headers,
        )
        self.assertEqual(finalized.status_code, 200)
        with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/generate",
                json={},
                headers=self.headers,
            )
            self.assertEqual(started.status_code, 202)
            job_id = started.get_json()["job"]["id"]
            job = self.wait_for_job(job_id)

        self.assertEqual(job["status"], "succeeded_warnings")
        self.assertIn("No statement files", job["warnings"][0])
        self.assertTrue(job["output_name"].startswith(f"expense_review_{trip.name}_company_"))
        workbook_path = trip / job["output_name"]
        self.assertTrue(workbook_path.exists())
        self.assertEqual(
            load_workbook(workbook_path, read_only=True).sheetnames,
            ["Expense Report", "Accounting Rows"],
        )
        accounting = load_workbook(workbook_path, data_only=False)["Accounting Rows"]
        self.assertEqual(
            [accounting.cell(row, 4).value for row in range(2, 8)],
            [
                "Travel – Non-meal",
                "Meals – Deductible (50%)",
                "Meals – Non-deductible (50%)",
                "GST Receivable",
                "QST Receivable",
                "Shareholder Current Account",
            ],
        )
        self.assertEqual(accounting["B2"].value, "Expense report – Client workshop")
        self.assertEqual(accounting["B7"].value, "Florent - Client workshop")
        self.assertEqual(accounting["C7"].value, "Payment of business trip")
        self.assertEqual(accounting["E7"].value, "Bank – Checking")
        manifest_path = trip / "trip-reimbursement-manifest.v3.ndjson"
        records = [
            json.loads(line)
            for line in manifest_path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
        self.assertEqual([record["kind"] for record in records], ["receipt", "trip_report"])
        self.assertEqual(records[-1]["claim_program"], "arvine_only")
        details = self.client.get(f"/api/trips/{trip.name}").get_json()["trip"]
        self.assertEqual(
            {entry["kind"] for entry in details["lifecycle"]["current_generation"]["artifacts"]},
            {"company_report", "reimbursement_manifest"},
        )

        approved = self.client.post(
            f"/api/trips/{trip.name}/approve",
            json={
                "workbook": job["output_name"],
                "reviewer": "Manager",
                "note": "Reviewed the canonical report and contract controls.",
            },
            headers=self.headers,
        )
        self.assertEqual(approved.status_code, 200)
        exported = self.client.post(
            f"/api/trips/{trip.name}/export-package",
            json={},
            headers=self.headers,
        )
        self.assertEqual(exported.status_code, 200)
        package_path = trip / exported.get_json()["package"]["name"]
        with zipfile.ZipFile(package_path) as archive:
            names = set(archive.namelist())
            self.assertIn(job["output_name"], names)
            self.assertIn("trip-reimbursement-manifest.v3.ndjson", names)
            self.assertIn("approval-manifest.json", names)

    def test_sponsored_bundle_adds_ivado_adapter_from_the_same_contract(self):
        trip = ensure_trip(self.root, "202607_montreal-sponsored", mode="ivado")
        self.complete_trip_metadata(trip)
        receipt = trip / "expenses_receipts" / "dinner.pdf"
        receipt.write_bytes(b"fixture")
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Bistro",
            expense_type="meal",
            amount=115,
            currency="CAD",
            line_items=[
                LineItem(description="Dinner and tip", amount=95),
                LineItem(
                    description="Wine",
                    amount=20,
                    is_alcohol=True,
                    alcohol_confidence=0.99,
                ),
            ],
        )
        save_line_item_review(trip, [expense])
        finalized = self.client.post(
            f"/api/trips/{trip.name}/finalize",
            json={},
            headers=self.headers,
        )
        self.assertEqual(finalized.status_code, 200)

        with patch("nlp_expenses.generator.parse_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/generate",
                json={},
                headers=self.headers,
            )
            terminal = self.wait_for_job(started.get_json()["job"]["id"])
        self.assertEqual(terminal["status"], "succeeded")
        primary = trip / terminal["output_name"]
        ivado = trip / terminal["output_name"].replace("_company_", "_ivado_", 1)
        self.assertTrue(primary.is_file())
        self.assertTrue(ivado.is_file())
        self.assertEqual(
            load_workbook(ivado, read_only=True).sheetnames,
            [
                "modèle - Template FR EN",
                "Card Statements",
                "Reconciliation",
                "Directives & instructions - FR",
                "Guidelines & Instructions - EN",
            ],
        )
        ivado_workbook = load_workbook(ivado, data_only=False)
        expense_report = ivado_workbook["modèle - Template FR EN"]
        self.assertEqual(expense_report["H15"].value, 95)
        self.assertEqual(expense_report["N15"].value, 95)
        statement_headers = [cell.value for cell in ivado_workbook["Card Statements"][5]]
        self.assertEqual(statement_headers[6], "Statement CAD")
        self.assertEqual(statement_headers[13], "Allocated IVADO Claim CAD")
        reconciliation = ivado_workbook["Reconciliation"]
        rows = list(reconciliation.iter_rows(values_only=True))
        receipt_header = next(index for index, row in enumerate(rows) if "Vendor" in row)
        bistro_row = next(row for row in rows[receipt_header + 1 :] if row[2] == "Bistro")
        self.assertEqual(bistro_row[14], 20)
        self.assertEqual(bistro_row[15], 95)
        self.assertEqual(bistro_row[16], 20)
        self.assertEqual(bistro_row[17], 115)
        self.assertNotIn("Receipt Line", {value for row in rows for value in row})
        records = [
            json.loads(line)
            for line in (trip / "trip-reimbursement-manifest.v3.ndjson")
            .read_text(encoding="utf-8")
            .splitlines()
            if line.strip()
        ]
        receipt_record = records[0]
        report_record = records[-1]
        self.assertEqual(receipt_record["arvine_reimbursable_cad"], 115)
        self.assertEqual(receipt_record["ivado_claimable_cad"], 95)
        self.assertEqual(receipt_record["ivado_excluded_cad"], 20)
        self.assertEqual(report_record["employee_reimbursement_total_cad"], 115)
        self.assertEqual(report_record["ivado_claim_total_cad"], 95)

        accounting = load_workbook(primary, data_only=False)["Accounting Rows"]
        posting_rows = [
            tuple(accounting.cell(row, column).value for column in range(2, 7))
            for row in range(2, 7)
        ]
        self.assertEqual(
            posting_rows,
            [
                (
                    "Florent",
                    "Reimbursable travel expense (passthrough IVADO)",
                    "Expenses Recoverable from Clients",
                    "Shareholder Current Account",
                    95,
                ),
                (
                    "Florent",
                    "Arvine-borne alcohol – deductible 50%",
                    "Meals – Deductible (50%)",
                    "Shareholder Current Account",
                    10,
                ),
                (
                    "Florent",
                    "Arvine-borne alcohol – non-deductible 50%",
                    "Meals – Non-deductible (50%)",
                    "Shareholder Current Account",
                    10,
                ),
                (
                    "IVADO Labs",
                    "Client workshop - Invoice sent to IL by Arvine",
                    "Accounts Receivable",
                    "Expenses Recoverable from Clients",
                    95,
                ),
                (
                    "IVADO Labs",
                    "Client workshop - Reimbursement from IL",
                    "Bank – Checking",
                    "Accounts Receivable",
                    95,
                ),
            ],
        )

    def test_workbook_generation_inherits_the_receipt_scan_quality(self):
        trip = ensure_trip(self.root, "202607_quality-source", mode="ivado")
        receipt = trip / "expenses_receipts" / "receipt.pdf"
        receipt.write_bytes(b"fixture")
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Cafe",
            expense_type="meal",
            amount=12.5,
            currency="CAD",
        )
        save_line_item_review(trip, [expense], llm_mode="required")
        self.complete_trip_metadata(trip)
        finalized = self.client.post(
            f"/api/trips/{trip.name}/finalize",
            json={},
            headers=self.headers,
        )
        self.assertEqual(finalized.status_code, 200)
        (self.root / ".env").write_text("OPENAI_API_KEY=sk-test-ui\n", encoding="utf-8")

        def create_workbook(*_args, **kwargs):
            kwargs["output_path"].write_bytes(b"workbook")
            return kwargs["output_path"]

        with patch("nlp_expenses.jobs.generate_review", side_effect=create_workbook):
            started = self.client.post(
                f"/api/trips/{trip.name}/generate",
                json={"quality": "basic"},
                headers=self.headers,
            )
            self.assertEqual(started.status_code, 202)
            self.assertEqual(started.get_json()["job"]["quality"], "best")
            terminal = self.wait_for_job(started.get_json()["job"]["id"])
        self.assertEqual(terminal["status"], "succeeded")

    def test_line_item_review_excludes_alcohol_until_reclassified(self):
        trip = ensure_trip(self.root, "202607_line-review", mode="ivado")
        receipt = trip / "expenses_receipts" / "meal.pdf"
        receipt.write_bytes(b"fixture")
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Bistro",
            expense_type="meal",
            amount=30,
            currency="CAD",
            line_items=[
                LineItem(description="Dinner", amount=20, included=True),
                LineItem(
                    description="French 75",
                    amount=10,
                    is_alcohol=True,
                    included=False,
                    alcohol_confidence=0.97,
                    alcohol_reason="recognized cocktail",
                    alcohol_matched_term="french 75",
                ),
            ],
        )
        save_line_item_review(trip, [expense])
        view = self.client.get(f"/?trip={trip.name}")
        self.assertEqual(view.status_code, 200)
        html = view.get_data(as_text=True)
        self.assertIn("Extract and review receipts", html)
        self.assertIn("French 75", html)
        self.assertIn("recognized cocktail", html)
        self.assertIn("Employees sharing bill", html)
        self.assertIn('class="people-review-form"', html)
        self.assertIn('class="button secondary small receipt-items-toggle"', html)
        self.assertIn('id="collapse-all-receipts"', html)
        self.assertIn('id="expand-review-receipts"', html)
        self.assertIn('id="expand-all-receipts"', html)
        self.assertNotIn('name="approver"', html)
        self.assertNotIn("Policy controls", html)
        self.assertNotIn("Settlement and IVADO handoff", html)

        current = self.client.get(f"/api/trips/{trip.name}").get_json()["trip"]["line_item_review"]
        cocktail = next(
            item
            for item in current["receipts"][0]["line_items"]
            if item["description"] == "French 75"
        )
        updated = self.client.post(
            f"/api/trips/{trip.name}/line-items/item",
            json={
                "source_file": receipt.name,
                "line_id": cocktail["line_id"],
                "fields": {"is_alcohol": False},
            },
            headers=self.headers,
        )
        self.assertEqual(updated.status_code, 200)
        cocktail = next(
            item
            for item in updated.get_json()["line_item_review"]["receipts"][0]["line_items"]
            if item["description"] == "French 75"
        )
        self.assertTrue(cocktail["included"])
        self.assertFalse(cocktail["is_alcohol"])
        self.assertTrue(cocktail["alcohol_overridden"])

        reset = self.client.post(
            f"/api/trips/{trip.name}/line-items/reset",
            json={"source_file": receipt.name},
            headers=self.headers,
        )
        self.assertEqual(reset.status_code, 200)
        cocktail = next(
            item
            for item in reset.get_json()["line_item_review"]["receipts"][0]["line_items"]
            if item["description"] == "French 75"
        )
        self.assertFalse(cocktail["included"])

    def test_receipt_review_controls_currency_and_secure_inline_preview(self):
        trip = ensure_trip(self.root, "202607_receipt-controls", mode="arvine")
        receipt = trip / "expenses_receipts" / "meals" / "dinner.pdf"
        receipt.parent.mkdir(parents=True)
        receipt.write_bytes(b"%PDF-1.4\nfixture")
        expense = Expense(
            source_file=Path("meals/dinner.pdf"),
            expense_id="",
            date="2026-07-05",
            supplier_name="Dinner",
            expense_type="meal",
            amount=40,
            currency="EUR",
            number_of_people=2,
            line_items=[LineItem(description="Dinner", amount=40)],
        )
        save_line_item_review(trip, [expense])

        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn('class="receipt-file-link receipt-preview"', html)
        self.assertIn('class="currency-review-form"', html)
        self.assertIn("Per employee (÷ 2)", html)
        self.assertIn("Receipt 20.00", html)
        self.assertNotIn("Saved automatically", html)
        self.assertNotIn(">Save share</button>", html)
        self.assertIn('class="receipt-meal-toggle"', html)
        self.assertIn("Meal expense · 50% deductible", html)
        self.assertIn('list="currency-options"', html)
        self.assertIn('<option value="QAR">Qatari Rial</option>', html)
        self.assertIn('class="receipt-reviewed"', html)
        self.assertIn('class="line-item-reviewed"', html)
        self.assertIn('data-autosave-field="description"', html)
        self.assertIn('data-autosave-field="amount"', html)
        amount_inputs = re.findall(
            r'<input class="input compact-input line-item-amount"[^>]+>', html
        )
        self.assertTrue(
            any(
                'value="40.00"' in input_html and 'min="0"' not in input_html
                for input_html in amount_inputs
            )
        )
        self.assertIn('class="line-save-status"', html)
        self.assertNotIn('class="button secondary small save-line-item"', html)
        self.assertIn("Click outside this window or press Esc to close", html)

        preview = self.client.get(
            f"/api/trips/{trip.name}/receipt",
            query_string={"filename": "meals/dinner.pdf"},
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data, b"%PDF-1.4\nfixture")
        self.assertIn("inline", preview.headers["Content-Disposition"])
        preview.close()

        escaped = self.client.get(
            f"/api/trips/{trip.name}/receipt",
            query_string={"filename": "../outside.pdf"},
        )
        self.assertEqual(escaped.status_code, 400)

        corrected = self.client.post(
            f"/api/trips/{trip.name}/line-items/expense",
            json={"source_file": "meals/dinner.pdf", "fields": {"currency": "cad"}},
            headers=self.headers,
        )
        self.assertEqual(corrected.status_code, 200)
        reviewed = corrected.get_json()["line_item_review"]["receipts"][0]
        self.assertEqual(reviewed["currency"], "CAD")
        self.assertEqual(reviewed["status"], "ok")

        ready = self.client.post(
            f"/api/trips/{trip.name}/line-items/expense",
            json={"source_file": "meals/dinner.pdf", "fields": {"reviewed": True}},
            headers=self.headers,
        )
        self.assertEqual(ready.status_code, 200)
        reviewed = ready.get_json()["line_item_review"]["receipts"][0]
        self.assertEqual(reviewed["status"], "ready")
        self.assertTrue(reviewed["line_items"][0]["reviewed"])

    def test_arvine_line_review_hides_alcohol_and_ivado_controls(self):
        trip = ensure_trip(self.root, "202607_arvine-lines", mode="arvine")
        receipt = trip / "expenses_receipts" / "meal.pdf"
        receipt.write_bytes(b"fixture")
        save_line_item_review(
            trip,
            [
                Expense(
                    source_file=receipt,
                    expense_id="",
                    date="2026-07-01",
                    supplier_name="Bistro",
                    expense_type="meal",
                    amount=30,
                    currency="CAD",
                    line_items=[
                        LineItem(
                            description="Wine",
                            amount=30,
                            is_alcohol=True,
                            alcohol_confidence=0.99,
                        )
                    ],
                )
            ],
        )

        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Alcohol classification and exclusions do not apply.", html)
        self.assertNotIn('class="line-item-alcohol"', html)
        self.assertNotIn('class="line-item-ivado"', html)
        self.assertNotIn('name="included_in_ivado"', html)

    def test_app_first_expense_line_and_finalization_endpoints(self):
        trip = ensure_trip(self.root, "202607_app-first-api", mode="ivado")
        self.complete_trip_metadata(trip)
        receipt = trip / "expenses_receipts" / "meal.pdf"
        receipt.write_bytes(b"fixture")
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Bistro",
            expense_type="meal",
            amount=40,
            currency="CAD",
            line_items=[LineItem(description="Dinner", amount=40)],
        )
        save_line_item_review(trip, [expense])

        updated = self.client.post(
            f"/api/trips/{trip.name}/line-items/expense",
            json={
                "source_file": receipt.name,
                "fields": {
                    "number_of_people": 2,
                    "included": True,
                    "business_purpose": "Client dinner",
                },
            },
            headers=self.headers,
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(
            updated.get_json()["line_item_review"]["receipts"][0]["number_of_people"],
            2,
        )
        self.assertEqual(
            updated.get_json()["consolidation"]["summary"]["claimable_cad"],
            20,
        )

        added = self.client.post(
            f"/api/trips/{trip.name}/line-items/add",
            json={
                "source_file": receipt.name,
                "description": "Tip",
                "amount": 5,
                "included": True,
            },
            headers=self.headers,
        )
        self.assertEqual(added.status_code, 200)
        manual = next(
            item
            for item in added.get_json()["line_item_review"]["receipts"][0]["line_items"]
            if item["description"] == "Tip"
        )
        removed = self.client.post(
            f"/api/trips/{trip.name}/line-items/remove",
            json={"source_file": receipt.name, "line_id": manual["line_id"]},
            headers=self.headers,
        )
        self.assertEqual(removed.status_code, 200)

        finalized = self.client.post(
            f"/api/trips/{trip.name}/finalize",
            json={},
            headers=self.headers,
        )
        self.assertEqual(finalized.status_code, 200)
        self.assertTrue(finalized.get_json()["consolidation"]["is_finalized"])
        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Review and finalize the claim", html)
        self.assertIn("Export the finalized claim to Excel", html)
        self.assertIn("20.00 CAD", html)

    def test_ivado_reconciliation_is_available_in_the_frontend(self):
        trip = ensure_trip(self.root, "202607_ivado-reconcile", mode="ivado")
        receipt = trip / "expenses_receipts" / "cafe.pdf"
        receipt.write_bytes(b"fixture")
        (trip / "card_statements" / "card.csv").write_text(
            "Date,Description,Amount,Foreign Spend Amount\n2026-07-01,CAFE,75,50 AUD\n",
            encoding="utf-8",
        )
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Cafe",
            expense_type="meal",
            amount=50,
            currency="AUD",
            line_items=[LineItem(description="Lunch", amount=50)],
        )
        save_line_item_review(trip, [expense])
        with patch("nlp_expenses.generator.parse_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/reconcile",
                json={},
                headers=self.headers,
            )
            self.assertEqual(started.status_code, 202)
            self.wait_for_job(started.get_json()["job"]["id"])

        reconciliation = self.client.get(f"/api/trips/{trip.name}/reconciliation").get_json()[
            "reconciliation"
        ]
        self.assertTrue(reconciliation["available"])
        self.assertEqual(reconciliation["summary"]["matched_invoice_count"], 1)
        self.assertEqual(reconciliation["expenses"][0]["cad_amount_used"], 75)
        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Reconcile invoices and card statements", html)
        self.assertIn("Receipt and card matches", html)
        self.assertIn("Statement transactions without a receipt", html)
        self.assertNotIn("<h4>Statement mappings</h4>", html)
        self.assertIn("1.500000", html)

        people = self.client.post(
            f"/api/trips/{trip.name}/line-items/expense",
            json={"source_file": receipt.name, "fields": {"number_of_people": 2}},
            headers=self.headers,
        )
        self.assertEqual(people.status_code, 200)
        basis = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/statement-basis",
            json={
                "source_file": receipt.name,
                "basis": "personal_share",
            },
            headers=self.headers,
        )
        self.assertEqual(basis.status_code, 200)
        changed = basis.get_json()["reconciliation"]["expenses"][0]
        self.assertEqual(changed["statement_amount_basis"], "personal_share")
        changed_html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Card amount represents", changed_html)
        self.assertIn("Traveller personal share", changed_html)
        self.assertIn("reviewed alcohol-free share", changed_html)

    def test_line_item_scan_job_populates_ivado_review(self):
        trip = ensure_trip(self.root, "202607_line-scan", mode="ivado")
        receipt = trip / "expenses_receipts" / "meal.pdf"
        receipt.write_bytes(b"fixture")
        before_scan = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("<h3>Receipts</h3>", before_scan)
        self.assertIn("Not scanned", before_scan)
        self.assertIn("Scan 1 new/changed receipt", before_scan)
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Bistro",
            expense_type="meal",
            amount=30,
            currency="CAD",
            line_items=[
                LineItem(description="Dinner", amount=20),
                LineItem(
                    description="French 75",
                    amount=10,
                    is_alcohol=True,
                    alcohol_confidence=0.97,
                    alcohol_reason="recognized cocktail",
                ),
            ],
        )
        with patch("nlp_expenses.generator.parse_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/line-items/sync",
                json={"quality": "basic"},
                headers=self.headers,
            )
            self.assertEqual(started.status_code, 202)
            terminal = self.wait_for_job(started.get_json()["job"]["id"])
        self.assertEqual(terminal["kind"], "line_items")
        self.assertEqual(terminal["status"], "succeeded")
        review = self.client.get(f"/api/trips/{trip.name}").get_json()["trip"]["line_item_review"]
        self.assertTrue(review["available"])
        self.assertEqual(review["summary"]["line_count"], 4)
        self.assertEqual(review["summary"]["excluded_count"], 1)
        after_scan = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Scanned", after_scan)
        self.assertIn("All receipts scanned", after_scan)
        self.assertIn("Rescan all receipts", after_scan)

    def test_reconciliation_sync_and_manual_mapping_api(self):
        trip = ensure_trip(self.root, "202607_reconcile", mode="arvine")
        receipt = trip / "expenses_receipts" / "hotel.pdf"
        receipt.write_bytes(b"fixture")
        statement = trip / "card_statements" / "card.csv"
        statement.write_text(
            "Date,Description,Amount,Currency,Foreign Amount\n"
            "2026-07-01,FOREIGN HOTEL,130,CAD,100 USD\n",
            encoding="utf-8",
        )
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Foreign Hotel",
            expense_type="hotel",
            amount=100,
            currency="USD",
        )
        stale_generation = self.client.post(
            f"/api/trips/{trip.name}/generate",
            json={"quality": "basic", "statements_complete": True},
            headers=self.headers,
        )
        self.assertEqual(stale_generation.status_code, 400)
        self.assertIn("Sync and auto-match", stale_generation.get_json()["error"])

        unscanned_reconciliation = self.client.post(
            f"/api/trips/{trip.name}/reconcile",
            json={},
            headers=self.headers,
        )
        self.assertEqual(unscanned_reconciliation.status_code, 400)
        self.assertIn(
            "Scan the current receipts in Step 3", unscanned_reconciliation.get_json()["error"]
        )

        save_line_item_review(trip, [expense])
        with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/reconcile",
                json={"quality": "best"},
                headers=self.headers,
            )
            self.assertEqual(started.status_code, 202)
            self.assertEqual(started.get_json()["job"]["quality"], "basic")
            self.assertTrue(started.get_json()["job"]["only_unmatched"])
            job = self.wait_for_job(started.get_json()["job"]["id"])
        self.assertEqual(job["kind"], "reconciliation")
        self.assertEqual(job["status"], "succeeded")

        fetched = self.client.get(f"/api/trips/{trip.name}/reconciliation")
        reconciliation = fetched.get_json()["reconciliation"]
        self.assertEqual(reconciliation["summary"]["matched_invoice_count"], 1)
        transaction = reconciliation["transactions"][0]
        self.assertAlmostEqual(transaction["fx_rate"], 1.3)
        self.assertEqual(
            reconciliation["expenses"][0]["statement_matches"][0]["group_id"],
            transaction["group_id"],
        )
        self.assertTrue(reconciliation["expenses"][0]["match_suggestions"][0]["current"])
        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn('class="button secondary small open-receipt-matcher"', html)
        self.assertIn("All cards and accounts", html)
        self.assertIn('id="card-match-date"', html)
        self.assertIn('id="card-match-date-window"', html)
        self.assertIn("±3 days", html)
        self.assertIn("every eligible uploaded-card transaction remains searchable", html)
        self.assertIn("assumed ok", html)
        self.assertIn('class="current-match-card"', html)
        self.assertIn(transaction["description"], html)
        self.assertIn(f"{transaction['cad_amount']:.2f} CAD", html)
        self.assertIn("Sync unmatched receipts", html)
        self.assertIn("Rebuild automatic matches", html)

        with patch(
            "nlp_expenses.generator.parse_arvine_receipt",
            side_effect=AssertionError("reconciliation should reuse Step 3"),
        ) as parser:
            rebuild = self.client.post(
                f"/api/trips/{trip.name}/reconcile",
                json={"only_unmatched": False},
                headers=self.headers,
            )
            self.assertEqual(rebuild.status_code, 202)
            self.assertFalse(rebuild.get_json()["job"]["only_unmatched"])
            rebuilt_job = self.wait_for_job(rebuild.get_json()["job"]["id"])
        parser.assert_not_called()
        self.assertEqual(rebuilt_job["status"], "succeeded")

        cleared = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/mapping",
            json={"group_id": transaction["group_id"], "expense_file": None, "use_auto": False},
            headers=self.headers,
        )
        self.assertEqual(cleared.status_code, 200)
        changed = cleared.get_json()["reconciliation"]["transactions"][0]
        self.assertIsNone(changed["expense_file"])
        self.assertEqual(changed["match_status"], "unmatched")
        unmatched_html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Statement transactions without a receipt", unmatched_html)
        self.assertIn("FOREIGN HOTEL", unmatched_html)
        self.assertIn('class="icon-button exclude-transaction"', unmatched_html)
        self.assertNotIn("transaction-disposition-dialog", unmatched_html)

        edited = self.client.post(
            f"/api/trips/{trip.name}/line-items/expense",
            json={"source_file": receipt.name, "fields": {"currency": "EUR"}},
            headers=self.headers,
        )
        self.assertEqual(edited.status_code, 200)
        self.assertEqual(edited.get_json()["line_item_review"]["receipts"][0]["currency"], "EUR")
        refreshed = self.client.get(f"/api/trips/{trip.name}/reconciliation")
        updated_reconciliation = refreshed.get_json()["reconciliation"]
        self.assertFalse(updated_reconciliation["stale"])
        self.assertEqual(updated_reconciliation["expenses"][0]["currency"], "EUR")
        self.assertIsNone(updated_reconciliation["transactions"][0]["expense_file"])
        self.assertEqual(updated_reconciliation["transactions"][0]["match_status"], "unmatched")
        refreshed_html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn('class="button secondary small open-receipt-matcher"', refreshed_html)
        self.assertNotIn('class="button secondary small resync-card-matcher"', refreshed_html)

        legacy_state = load_reconciliation_state(trip)
        legacy_state["expenses"][0]["currency"] = "USD"
        legacy_state["invoice_overrides"].pop(receipt.name, None)
        legacy_state["requires_resync"] = True
        legacy_state.pop("requires_resync_reason", None)
        save_reconciliation_state(trip, legacy_state)
        migrated = self.client.get(f"/api/trips/{trip.name}/reconciliation").get_json()[
            "reconciliation"
        ]
        self.assertFalse(migrated["stale"])
        self.assertEqual(migrated["expenses"][0]["currency"], "EUR")

    def test_active_job_locks_trip_source_mutations(self):
        trip = ensure_trip(self.root, "202607_locked", mode="ivado")
        (trip / "expenses_receipts" / "receipt.pdf").write_bytes(b"fixture")
        manager = JobManager(self.root)
        app = create_app(self.root, access_token="locked-token", job_manager=manager)
        app.testing = True
        client = app.test_client()
        client.get("/?token=locked-token")
        with client.session_transaction() as session:
            headers = {"X-CSRF-Token": session["csrf_token"]}

        release = threading.Event()

        def wait_for_release(*_args, **_kwargs):
            release.wait(timeout=5)
            return trip / "expense_review_locked.xlsx"

        try:
            with patch("nlp_expenses.jobs.generate_review", side_effect=wait_for_release):
                job = manager.start(trip.name, "basic", statements_complete=False)
                for _ in range(100):
                    if manager.get(job.id).status == "running":
                        break
                    time.sleep(0.01)

                uploaded = client.post(
                    f"/api/trips/{trip.name}/upload/receipts",
                    data={"files": (BytesIO(b"new"), "new.pdf")},
                    content_type="multipart/form-data",
                    headers=headers,
                )
                self.assertEqual(uploaded.status_code, 409)
                removed = client.post(
                    f"/api/trips/{trip.name}/remove-file",
                    json={"kind": "receipts", "filename": "receipt.pdf"},
                    headers=headers,
                )
                self.assertEqual(removed.status_code, 409)
                switched = client.post(
                    f"/api/trips/{trip.name}/mode",
                    json={"mode": "arvine"},
                    headers=headers,
                )
                self.assertEqual(switched.status_code, 409)
                deleted = client.delete(
                    f"/api/trips/{trip.name}",
                    json={"confirmation": trip.name},
                    headers=headers,
                )
                self.assertEqual(deleted.status_code, 409)
                self.assertTrue(trip.is_dir())
        finally:
            release.set()
            self.wait_for_manager(manager, job.id)

    def test_invoice_review_api_reports_validation_and_persists_override(self):
        trip = ensure_trip(self.root, "202607_invoice-review", mode="arvine")
        receipt = trip / "expenses_receipts" / "meal.pdf"
        receipt.write_bytes(b"fixture")
        (trip / "card_statements" / "card.csv").write_text(
            "Date,Description,Amount,Currency,Foreign Amount\n"
            "2026-07-01,CLIENT DINNER,130,CAD,100 USD\n",
            encoding="utf-8",
        )
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Client Dinner",
            expense_type="meal",
            amount=100,
            currency="USD",
        )
        save_line_item_review(trip, [expense])
        self.complete_trip_metadata(trip)
        with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/reconcile",
                json={"quality": "basic"},
                headers=self.headers,
            )
            self.wait_for_job(started.get_json()["job"]["id"])

        invalid = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/invoice",
            json={
                "source_file": "meal.pdf",
                "fields": {"date": "not-a-date", "amount": "-12", "currency": "US dollars"},
            },
            headers=self.headers,
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual({"date", "amount", "currency"}, set(invalid.get_json()["fields"]))

        saved = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/invoice",
            json={
                "source_file": "meal.pdf",
                "fields": {
                    "date": "2026-07-02",
                    "vendor": "Corrected Restaurant",
                    "amount": "95",
                    "currency": "EUR",
                },
                "update_manual_cad": True,
                "manual_cad": {"amount": "142.25", "note": "Card portal confirmation"},
            },
            headers=self.headers,
        )
        self.assertEqual(saved.status_code, 200)
        view = saved.get_json()["reconciliation"]
        self.assertFalse(view["stale"])
        self.assertEqual(view["expenses"][0]["vendor"], "Corrected Restaurant")
        self.assertEqual(view["expenses"][0]["manual_cad_override"]["amount"], 142.25)

    def test_duplicate_transaction_decision_api(self):
        trip = ensure_trip(self.root, "202607_duplicate-api", mode="arvine")
        receipt = trip / "expenses_receipts" / "cafe.pdf"
        receipt.write_bytes(b"fixture")
        for filename in ("one.csv", "two.csv"):
            (trip / "card_statements" / filename).write_text(
                "Date,Description,Amount,Currency\n2026-07-01,CLIENT CAFE,10,CAD\n",
                encoding="utf-8",
            )
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Client Cafe",
            expense_type="meal",
            amount=10,
            currency="CAD",
        )
        save_line_item_review(trip, [expense])
        self.complete_trip_metadata(trip)
        with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/reconcile",
                json={"quality": "basic"},
                headers=self.headers,
            )
            self.wait_for_job(started.get_json()["job"]["id"])
        view = self.client.get(f"/api/trips/{trip.name}/reconciliation").get_json()[
            "reconciliation"
        ]
        duplicate = next(item for item in view["transactions"] if item["possible_duplicate"])

        saved = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/transaction-decision",
            json={"group_id": duplicate["group_id"], "action": "ignore", "note": ""},
            headers=self.headers,
        )
        self.assertEqual(saved.status_code, 200)
        changed = next(
            item
            for item in saved.get_json()["reconciliation"]["transactions"]
            if item["group_id"] == duplicate["group_id"]
        )
        self.assertEqual(changed["duplicate_resolution"], "ignore")
        self.assertTrue(changed["ignored"])

    def test_ivado_unrelated_transaction_can_be_excluded_from_the_trip(self):
        trip = ensure_trip(self.root, "202607_ivado-disposition", mode="ivado")
        receipt = trip / "expenses_receipts" / "taxi.pdf"
        receipt.write_bytes(b"fixture")
        (trip / "card_statements" / "card.csv").write_text(
            "Date,Description,Amount,Foreign Spend Amount\n"
            "2026-07-01,TAXI,75,50 AUD\n"
            "2026-07-01,SPOTIFY,20,\n",
            encoding="utf-8",
        )
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Taxi",
            expense_type="transport",
            amount=50,
            currency="AUD",
        )
        save_line_item_review(trip, [expense])
        self.complete_trip_metadata(trip)
        with patch("nlp_expenses.generator.parse_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/reconcile",
                json={"quality": "basic"},
                headers=self.headers,
            )
            self.wait_for_job(started.get_json()["job"]["id"])
        view = self.client.get(f"/api/trips/{trip.name}/reconciliation").get_json()[
            "reconciliation"
        ]
        unrelated = next(item for item in view["transactions"] if "SPOTIFY" in item["description"])

        page = self.client.get(f"/?trip={trip.name}")
        self.assertIn(b'class="icon-button exclude-transaction"', page.data)
        self.assertNotIn(b"transaction-disposition-dialog", page.data)
        saved = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/transaction-decision",
            json={
                "group_id": unrelated["group_id"],
                "action": "ignore",
            },
            headers=self.headers,
        )
        self.assertEqual(saved.status_code, 200)
        changed = next(
            item
            for item in saved.get_json()["reconciliation"]["transactions"]
            if item["group_id"] == unrelated["group_id"]
        )
        self.assertTrue(changed["ignored"])
        self.assertEqual(changed["decision_note"], "")

        excluded_page = self.client.get(f"/?trip={trip.name}")
        self.assertIn(
            b"excluded statement transaction(s) \xc2\xb7 restore if needed", excluded_page.data
        )
        self.assertIn(b'class="button secondary small restore-transaction"', excluded_page.data)

        restored = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/transaction-decision",
            json={"group_id": unrelated["group_id"], "action": "keep", "note": ""},
            headers=self.headers,
        )
        self.assertEqual(restored.status_code, 200)
        restored_item = next(
            item
            for item in restored.get_json()["reconciliation"]["transactions"]
            if item["group_id"] == unrelated["group_id"]
        )
        self.assertFalse(restored_item["ignored"])

    def test_statement_coverage_gap_requires_generation_explanation(self):
        trip = ensure_trip(self.root, "202607_coverage-api", mode="arvine")
        self.complete_trip_metadata(trip)
        receipt = trip / "expenses_receipts" / "hotel.pdf"
        receipt.write_bytes(b"fixture")
        (trip / "card_statements" / "card.csv").write_text(
            "Date,Description,Amount,Currency,Account\n2026-07-01,HOTEL,100,CAD,1234\n",
            encoding="utf-8",
        )
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Hotel",
            expense_type="hotel",
            amount=100,
            currency="CAD",
        )
        save_line_item_review(trip, [expense])
        with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/reconcile",
                json={"quality": "basic"},
                headers=self.headers,
            )
            self.wait_for_job(started.get_json()["job"]["id"])

        settings = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/coverage-settings",
            json={"expected_accounts": ["GENERIC ••••1234", "AMEX ••••9999"]},
            headers=self.headers,
        )
        self.assertEqual(settings.status_code, 200)
        self.assertGreaterEqual(len(settings.get_json()["coverage"]["gaps"]), 1)

        html = self.client.get(f"/?trip={trip.name}").get_data(as_text=True)
        self.assertIn("Accounts expected for this trip (optional)", html)
        self.assertIn(
            "Completeness checklist only; this does not affect transaction matching.", html
        )
        self.assertIn("Save completeness checklist", html)

        rejected = self.client.post(
            f"/api/trips/{trip.name}/finalize",
            json={"statements_complete": True},
            headers=self.headers,
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertIn("coverage gaps", rejected.get_json()["error"])

        finalized = self.client.post(
            f"/api/trips/{trip.name}/finalize",
            json={
                "statements_complete": True,
                "coverage_acknowledgement": "Amex was not used",
            },
            headers=self.headers,
        )
        self.assertEqual(finalized.status_code, 200)
        with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
            accepted = self.client.post(
                f"/api/trips/{trip.name}/generate",
                json={},
                headers=self.headers,
            )
            self.assertEqual(accepted.status_code, 202)
            terminal = self.wait_for_job(accepted.get_json()["job"]["id"])
        self.assertIn(terminal["status"], {"succeeded", "succeeded_warnings"})

    def test_accounting_profile_api_updates_trip_and_can_become_default(self):
        trip = ensure_trip(self.root, "202607_profile-api", mode="arvine")
        profile = self.client.get(f"/api/trips/{trip.name}").get_json()["trip"][
            "accounting_profile"
        ]
        profile.update(
            {
                "company_legal_name": "Example Corp.",
                "traveller_reimbursement_type": "corporate_card",
                "counter_account": "Corporate Card Payable",
                "gst_hst_registrant": False,
                "qst_registrant": False,
                "meal_deduction_pct": 0.65,
            }
        )
        saved = self.client.post(
            f"/api/trips/{trip.name}/accounting-profile",
            json={"profile": profile, "make_default": True},
            headers=self.headers,
        )
        self.assertEqual(saved.status_code, 200)
        applied = saved.get_json()["trip"]["accounting_profile"]
        self.assertEqual(applied["counter_account"], "Corporate Card Payable")
        self.assertFalse(applied["gst_hst_registrant"])
        self.assertEqual(applied["meal_deduction_pct"], 0.65)

        later = ensure_trip(self.root, "202608_later-profile", mode="arvine")
        inherited = self.client.get(f"/api/trips/{later.name}").get_json()["trip"][
            "accounting_profile"
        ]
        self.assertEqual(inherited["company_legal_name"], "Example Corp.")
        self.assertEqual(inherited["counter_account"], "Corporate Card Payable")

    def test_split_allocation_api_reports_balance_and_blocks_generation(self):
        trip = ensure_trip(self.root, "202607_allocation-api", mode="arvine")
        receipt = trip / "expenses_receipts" / "hotel.pdf"
        receipt.write_bytes(b"fixture")
        (trip / "card_statements" / "card.csv").write_text(
            "Date,Description,Amount,Currency\n2026-07-01,HOTEL PACKAGE,300,CAD\n",
            encoding="utf-8",
        )
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Hotel",
            expense_type="hotel",
            amount=250,
            currency="CAD",
        )
        save_line_item_review(trip, [expense])
        with patch("nlp_expenses.generator.parse_arvine_receipt", return_value=expense):
            started = self.client.post(
                f"/api/trips/{trip.name}/reconcile",
                json={"quality": "basic"},
                headers=self.headers,
            )
            self.wait_for_job(started.get_json()["job"]["id"])
        view = self.client.get(f"/api/trips/{trip.name}/reconciliation").get_json()[
            "reconciliation"
        ]
        group_id = view["transactions"][0]["group_id"]

        draft = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/allocations",
            json={
                "group_id": group_id,
                "allocations": [
                    {"type": "purchase", "invoice_file": "hotel.pdf", "cad_amount": 250},
                    {
                        "type": "personal",
                        "category": "Personal",
                        "cad_amount": 25,
                        "note": "Personal",
                    },
                ],
            },
            headers=self.headers,
        )
        self.assertEqual(draft.status_code, 200)
        transaction = draft.get_json()["reconciliation"]["transactions"][0]
        self.assertEqual(transaction["allocation_status"], "unallocated")
        self.assertEqual(transaction["allocation_balance"], 25)
        rejected = self.client.post(
            f"/api/trips/{trip.name}/generate",
            json={"quality": "basic", "statements_complete": True},
            headers=self.headers,
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertIn("split allocations", rejected.get_json()["error"])

        balanced = self.client.post(
            f"/api/trips/{trip.name}/reconciliation/allocations",
            json={
                "group_id": group_id,
                "allocations": [
                    {"type": "purchase", "invoice_file": "hotel.pdf", "cad_amount": 250},
                    {
                        "type": "personal",
                        "category": "Personal",
                        "cad_amount": 50,
                        "note": "Personal extension",
                    },
                ],
            },
            headers=self.headers,
        )
        self.assertEqual(balanced.status_code, 200)
        self.assertEqual(
            balanced.get_json()["reconciliation"]["transactions"][0]["allocation_status"],
            "balanced",
        )

    def test_best_quality_requires_key_and_key_is_never_returned(self):
        trip = ensure_trip(self.root, "202607_montreal", mode="arvine")
        (trip / "expenses_receipts" / "receipt.pdf").write_bytes(b"fixture")
        with patch.dict(os.environ, {"OPENAI_API_KEY": ""}, clear=False):
            rejected = self.client.post(
                f"/api/trips/{trip.name}/generate",
                json={"quality": "best", "statements_complete": True},
                headers=self.headers,
            )
        self.assertEqual(rejected.status_code, 400)

        saved = self.client.post(
            "/api/settings/openai",
            json={"api_key": "sk-local-test"},
            headers=self.headers,
        )
        self.assertEqual(saved.status_code, 200)
        self.assertNotIn("sk-local-test", saved.get_data(as_text=True))
        env_path = self.root / ".env"
        self.assertIn("OPENAI_API_KEY=sk-local-test", env_path.read_text(encoding="utf-8"))
        self.assertEqual(stat.S_IMODE(env_path.stat().st_mode), 0o600)

    def test_download_is_limited_to_trip_workbooks(self):
        trip = ensure_trip(self.root, "202607_montreal", mode="arvine")
        workbook = trip / "expense_review_202607_montreal_arvine_20260714-153012.xlsx"
        workbook.write_bytes(b"xlsx")
        downloaded = self.client.get(
            f"/api/trips/{trip.name}/download-workbook",
            query_string={"filename": workbook.name},
        )
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded.data, b"xlsx")
        downloaded.close()
        escaped = self.client.get(
            f"/api/trips/{trip.name}/download-workbook",
            query_string={"filename": "../outside.xlsx"},
        )
        self.assertEqual(escaped.status_code, 400)

    def test_failed_job_removes_partial_output(self):
        trip = ensure_trip(self.root, "202607_failure", mode="ivado")
        manager = JobManager(self.root)

        def fail_after_partial(*_args, **kwargs):
            kwargs["output_path"].write_bytes(b"partial")
            raise RuntimeError("simulated failure")

        with patch("nlp_expenses.jobs.generate_review", side_effect=fail_after_partial):
            job = manager.start(trip.name, "basic", statements_complete=False)
            terminal = self.wait_for_manager(manager, job.id)
        self.assertEqual(terminal.status, "failed")
        self.assertEqual(list(trip.glob("expense_review_*.xlsx")), [])

    def test_best_quality_fallback_is_reported_as_a_warning(self):
        trip = ensure_trip(self.root, "202607_fallback", mode="ivado")
        receipt = trip / "expenses_receipts" / "receipt.pdf"
        receipt.write_bytes(b"fixture")
        expense = Expense(
            source_file=receipt,
            expense_id="",
            date="2026-07-01",
            supplier_name="Cafe",
            expense_type="meal",
            amount=12.5,
            currency="CAD",
            review_note="Local extraction only.",
        )
        warnings: list[str] = []
        output = trip / "expense_review_fallback.xlsx"
        with (
            patch(
                "nlp_expenses.generator.get_openai_settings", return_value=("sk-test", "gpt-test")
            ),
            patch("nlp_expenses.generator.parse_receipt", return_value=expense),
        ):
            result = generate_review(
                trip,
                self.root,
                llm_mode="required",
                output_path=output,
                warning_callback=warnings.append,
                allow_openai_prompt=False,
            )
        self.assertEqual(result, output.resolve())
        self.assertEqual(len(warnings), 1)
        self.assertIn("local extraction was used", warnings[0])
        self.assertTrue(output.exists())

    def wait_for_job(self, job_id: str) -> dict:
        for _ in range(100):
            response = self.client.get(f"/api/jobs/{job_id}")
            self.assertEqual(response.status_code, 200)
            job = response.get_json()["job"]
            if job["status"] not in {"queued", "running"}:
                return job
            time.sleep(0.02)
        self.fail("Generation job did not finish")

    def wait_for_manager(self, manager: JobManager, job_id: str):
        for _ in range(100):
            job = manager.get(job_id)
            if job.status not in {"queued", "running"}:
                return job
            time.sleep(0.01)
        self.fail("Generation job did not finish")


if __name__ == "__main__":
    unittest.main()
