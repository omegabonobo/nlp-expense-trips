from __future__ import annotations

import csv
import re
import subprocess
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from nlp_expenses.models import StatementTransaction

from .text import extract_text


DATE_RE = re.compile(r"\b(\d{1,2})\s+([A-Za-z]{3,9})\.?\s+(20\d{2})\b")
MONEY_RE = re.compile(r"^-?\$?\s*\d{1,3}(?:,\d{3})*(?:\.\d{2})$")
FOREIGN_RE = re.compile(r"(?P<amount>-?\d{1,3}(?:,\d{3})*(?:\.\d{2})|-?\d+)\s+(?P<currency>[A-Z ]+)")


def parse_statement_file(path: Path) -> list[StatementTransaction]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".csv":
            return parse_csv_statement(path)
        if suffix == ".xlsx":
            return parse_xlsx_statement(path)
        if suffix == ".xls":
            rows = parse_xls_statement(path)
            if rows:
                return rows
        if suffix == ".pdf":
            text, _method = extract_text(path)
            return parse_statement_text(path, text)
    except Exception:
        return []
    return []


def parse_all_statements(statements_dir: Path) -> list[StatementTransaction]:
    transactions: list[StatementTransaction] = []
    for path in sorted(p for p in statements_dir.iterdir() if p.is_file()):
        transactions.extend(parse_statement_file(path))
    return transactions


def parse_csv_statement(path: Path) -> list[StatementTransaction]:
    text = path.read_text(encoding="utf-8-sig", errors="ignore")
    sample = text[:2048]
    dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
    reader = csv.DictReader(text.splitlines(), dialect=dialect)
    return rows_to_transactions(path, list(reader))


def parse_xlsx_statement(path: Path) -> list[StatementTransaction]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header_index = find_header_index(rows)
    headers = [str(v or "").strip() for v in rows[header_index]]
    dict_rows = []
    for row in rows[header_index + 1 :]:
        dict_rows.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return rows_to_transactions(path, dict_rows)


def parse_xls_statement(path: Path) -> list[StatementTransaction]:
    try:
        import xlrd
    except Exception:
        return parse_statement_text(path, strings_text(path))
    workbook = xlrd.open_workbook(str(path))
    sheet = workbook.sheet_by_index(0)
    rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    header_index = find_header_index(rows)
    headers = [str(v or "").strip() for v in rows[header_index]]
    dict_rows = []
    for row in rows[header_index + 1 :]:
        dict_rows.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return rows_to_transactions(path, dict_rows)


def rows_to_transactions(path: Path, rows: list[dict]) -> list[StatementTransaction]:
    transactions: list[StatementTransaction] = []
    for row in rows:
        normalized = {normalize_key(k): v for k, v in row.items()}
        date = first_value(normalized, ["date", "transaction_date", "posted_date", "date_processed"])
        description = first_value(normalized, ["description", "merchant", "details", "transaction", "name"])
        amount = first_value(normalized, ["amount", "cad_amount", "debit", "charge", "charges_adjustments"])
        foreign = first_value(normalized, ["foreign_spend_amount", "foreign_amount", "original_amount"])
        parsed_amount = parse_money(amount)
        if parsed_amount is None:
            continue
        foreign_amount, foreign_currency = parse_foreign(foreign)
        transactions.append(
            StatementTransaction(
                source_file=path,
                date=parse_date(date),
                description=str(description or "").strip(),
                amount_cad=parsed_amount,
                foreign_amount=foreign_amount,
                foreign_currency=foreign_currency,
            )
        )
    return transactions


def parse_statement_text(path: Path, text: str) -> list[StatementTransaction]:
    transactions: list[StatementTransaction] = []
    current_date: str | None = None
    pending_desc = ""
    pending: StatementTransaction | None = None
    in_transaction_table = False
    for raw_line in text.splitlines():
        line = raw_line.strip().strip("\x0c%#!")
        if not line:
            continue
        if re.search(r"additional information|foreign spend amount", line, re.I):
            in_transaction_table = True
            continue
        if not in_transaction_table:
            continue
        date = parse_date(line)
        if date:
            current_date = date
            pending_desc = ""
            pending = None
            continue
        if not current_date:
            continue
        foreign_amount, foreign_currency = parse_foreign(line)
        if pending and foreign_amount is not None:
            pending.foreign_amount = foreign_amount
            pending.foreign_currency = foreign_currency
            pending = None
            continue
        amount = parse_text_cad_amount(line)
        if amount is not None:
            pending = StatementTransaction(
                source_file=path,
                date=current_date,
                description=pending_desc,
                amount_cad=amount,
            )
            transactions.append(pending)
            pending_desc = ""
            continue
        if looks_like_description(line):
            pending_desc = line
    return transactions


def parse_text_cad_amount(line: str) -> float | None:
    if "$" not in line:
        return None
    if re.search(r'"\$"|#|_', line):
        return None
    return parse_money(line)


def strings_text(path: Path) -> str:
    try:
        result = subprocess.run(["strings", "-n", "3", str(path)], check=True, capture_output=True, text=True)
        return result.stdout
    except Exception:
        return ""


def find_header_index(rows: list[list]) -> int:
    for idx, row in enumerate(rows[:30]):
        joined = " ".join(str(v or "").lower() for v in row)
        if "date" in joined and ("amount" in joined or "description" in joined):
            return idx
    return 0


def normalize_key(key: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(key or "").strip().lower()).strip("_")


def first_value(row: dict, keys: list[str]):
    for key in keys:
        if key in row and row[key] not in (None, ""):
            return row[key]
    return None


def parse_money(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace("$", "").replace(",", "").strip()
    try:
        parsed = float(text)
        return -parsed if negative else parsed
    except ValueError:
        return None


def parse_foreign(value) -> tuple[float | None, str | None]:
    if value is None:
        return None, None
    text = str(value).upper().replace(",", "")
    match = FOREIGN_RE.search(text)
    if not match:
        return None, None
    try:
        amount = float(match.group("amount"))
    except ValueError:
        return None, None
    currency_words = match.group("currency").strip()
    currency = {
        "AUSTRALIAN DOLLAR": "AUD",
        "CANADIAN DOLLAR": "CAD",
        "US DOLLAR": "USD",
        "INDONESIAN RUPIAH": "IDR",
    }.get(currency_words, currency_words[:3])
    return amount, currency


def parse_date(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and value > 20000:
        try:
            base = datetime(1899, 12, 30)
            return datetime.fromordinal(base.toordinal() + int(value)).strftime("%Y-%m-%d")
        except Exception:
            return None
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d %b %Y", "%d %B %Y", "%d %b. %Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    match = DATE_RE.search(text)
    if match:
        day, month, year = match.groups()
        for fmt in ("%d %b %Y", "%d %B %Y"):
            try:
                return datetime.strptime(f"{day} {month[:3]} {year}", "%d %b %Y").strftime("%Y-%m-%d")
            except ValueError:
                continue
    return None


def looks_like_description(line: str) -> bool:
    if MONEY_RE.match(line):
        return False
    if re.match(r"^\d+(?:\.\d+)?$", line):
        return False
    return bool(re.search(r"[A-Za-z]{3}", line)) and not re.search(r"summary|total|account|cardmember", line, re.I)
