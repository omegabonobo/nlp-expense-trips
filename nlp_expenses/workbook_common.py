from __future__ import annotations

import os
import subprocess
import uuid
from contextlib import suppress
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from nlp_expenses.models import Expense


def save_workbook_atomic(workbook: Workbook, out_path: Path) -> Path:
    out_path = out_path.resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = out_path.parent / f".{out_path.stem}.{uuid.uuid4().hex}.tmp.xlsx"
    try:
        workbook.save(temporary)
        os.replace(temporary, out_path)
    finally:
        temporary.unlink(missing_ok=True)
    remove_macos_metadata(out_path)
    return out_path


def style_table_header(ws, row: int, width: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col in range(1, width + 1):
        cell = ws.cell(row, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def set_widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def excel_date(value: str | None):
    if not value:
        return None
    try:
        return datetime.fromisoformat(value).date()
    except ValueError:
        return value


def arvine_extraction_status(expense: Expense) -> str:
    if expense.amount is None or not expense.currency:
        return "manual"
    if not expense.date or not expense.supplier_name or expense.supplier_name == "Unknown supplier":
        return "review"
    return "ok"


def configure_arvine_calculation(workbook: Workbook) -> None:
    calculation = getattr(workbook, "calculation", None)
    if calculation is None:
        return
    calculation.calcMode = "auto"
    calculation.fullCalcOnLoad = True
    calculation.forceFullCalc = True


def append_clean(ws, values: list) -> None:
    ws.append([None if value == "" else value for value in values])


def set_filter_range(ws, width: int, height: int) -> None:
    end_col = get_column_letter(width)
    ws.auto_filter.ref = f"A1:{end_col}{height}"


def remove_macos_metadata(path: Path) -> None:
    for attr in ("com.apple.quarantine", "com.apple.provenance", "com.apple.lastuseddate#PS"):
        with suppress(AttributeError, OSError):
            os.removexattr(path, attr)
        with suppress(OSError):
            subprocess.run(
                ["xattr", "-d", attr, str(path)],
                check=False,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
