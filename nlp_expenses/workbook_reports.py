from __future__ import annotations

from importlib.resources import as_file, files
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from nlp_expenses.workbook_common import (
    excel_date,
    save_workbook_atomic,
    set_filter_range,
    style_table_header,
)

IVADO_TEMPLATE_SHEET = "modèle - Template FR EN"
IVADO_CARD_SHEET = "Card Statements"
IVADO_RECONCILIATION_SHEET = "Reconciliation"
IVADO_INCLUDE_DETAILED_RECONCILIATION = False
IVADO_TEMPLATE_RESOURCE = files("nlp_expenses").joinpath(
    "static/ivado_expense_report_template.xlsx"
)


def build_reimbursement_report_workbook(
    trip_dir: Path,
    records: list[dict],
    output_path: Path,
) -> Path:
    """Build a compact, human-readable company expense report from contract v3."""

    receipts, report = split_manifest_records(records)
    workbook = Workbook()
    report_ws = workbook.active
    report_ws.title = "Expense Report"
    accounting_ws = workbook.create_sheet("Accounting Rows")

    write_minimal_arvine_report_sheet(report_ws, trip_dir, report, receipts)
    write_minimal_accounting_sheet(accounting_ws, report, receipts)
    style_minimal_report_workbook(workbook)
    return save_workbook_atomic(workbook, output_path)


def build_ivado_claim_workbook(
    trip_dir: Path,
    records: list[dict],
    output_path: Path,
    *,
    include_detailed_reconciliation: bool | None = None,
) -> Path:
    """Populate IVADO's official template and append auditable support schedules."""

    receipts, report = split_manifest_records(records)
    if report.get("claim_program") != "ivado_sponsored":
        raise ValueError("An IVADO workbook is only produced for IVADO-reimbursed trips.")

    with as_file(IVADO_TEMPLATE_RESOURCE) as template_path:
        workbook = load_workbook(template_path)

    report_ws = workbook[IVADO_TEMPLATE_SHEET]
    write_ivado_template_sheet(report_ws, report, receipts)
    statements_ws = replace_support_sheet(
        workbook,
        IVADO_CARD_SHEET,
        1,
        "card statement",
        "card_statement",
        "card_statements",
    )
    reconciliation_ws = replace_support_sheet(
        workbook,
        IVADO_RECONCILIATION_SHEET,
        2,
        "reconciliation_fv",
    )
    reconciliation = load_reconciliation_for_report(trip_dir)
    write_ivado_card_statement_sheet(
        statements_ws,
        report,
        receipts,
        reconciliation,
    )
    write_ivado_reconciliation_sheet(
        reconciliation_ws,
        report,
        receipts,
        reconciliation,
        include_line_items=(
            IVADO_INCLUDE_DETAILED_RECONCILIATION
            if include_detailed_reconciliation is None
            else include_detailed_reconciliation
        ),
    )
    configure_ivado_template_workbook(workbook)
    return save_workbook_atomic(workbook, output_path)


MINIMAL_ARVINE_HEADERS = [
    "Receipt #",
    "Date",
    "Vendor",
    "Description",
    "Category",
    "Receipt File",
    "Employees Sharing",
    "Currency",
    "Full Receipt",
    "Traveller Share",
    "Reimbursement Basis CAD",
    "Paid By",
    "Employee Reimbursement CAD",
    "Corporate Paid CAD",
    "IVADO Claim CAD",
    "IVADO Removed CAD",
]

IVADO_CARD_AUDIT_HEADERS = [
    "Receipt #",
    "Transaction Date",
    "Statement Account",
    "Statement Merchant",
    "Original Amount",
    "Original Currency",
    "Statement CAD",
    "Statement Basis",
    "Match Status",
    "Receipt File",
    "Receipt Date",
    "Receipt Vendor",
    "People",
    "Allocated IVADO Claim CAD",
    "Allocated IVADO Removed CAD",
    "Source Statement(s)",
]

IVADO_RECONCILIATION_HEADERS = [
    "Receipt\n#",
    "Receipt\nDate",
    "Vendor",
    "Description",
    "Expense\nType",
    "Receipt\nFile",
    "Local\nCurr",
    "Full Receipt\n(Local Curr)",
    "People",
    "Statement\nBasis",
    "Traveller Share\n(Local Curr)",
    "Statement CAD\n(Gross Share)",
    "Effective\nCAD Rate",
    "Eligible After Excl.\n(Local Curr)",
    "Alcohol Removed\n(Local Curr)",
    "IVADO Claim\n(CAD)",
    "IVADO Removed\n(CAD)",
    "Traveller Reimb.\n(CAD)",
]

IVADO_LINE_AUDIT_HEADERS = [
    "Receipt #",
    "Receipt Date",
    "Vendor",
    "Receipt File",
    "People",
    "Line ID",
    "Receipt Line",
    "Full Line Amount",
    "Traveller Share",
    "Currency",
    "Alcohol",
    "Included in IVADO",
    "Removed Original Share",
    "Removed CAD",
    "Exclusion Reason",
    "Review Note",
]


def write_minimal_arvine_report_sheet(
    ws,
    trip_dir: Path,
    report: dict,
    receipts: list[dict],
) -> None:
    ws["A1"] = "Company Trip Expense Report"
    ws["A3"] = "Traveller"
    ws["B3"] = report.get("traveller")
    ws["D3"] = "Report date"
    ws["E3"] = report.get("report_date")
    ws["A4"] = "Purpose"
    ws["B4"] = report.get("description")
    ws["D4"] = "Reimbursement program"
    ws["E4"] = readable_claim_program(report.get("claim_program"))
    ws["A5"] = "Trip"
    ws["B5"] = trip_dir.name
    ws["D5"] = "Contract"
    ws["E5"] = "trip-reimbursement-manifest.v3.ndjson"
    ws.append([])
    ws.append(MINIMAL_ARVINE_HEADERS)
    ws["A6"] = (
        "When a receipt is matched, the reimbursement basis is the full card charge, "
        "including any tip or adjustment; receipt tax amounts remain unchanged. "
        "For IVADO trips, Arvine reimburses the full reviewed business share while the "
        "IVADO claim excludes alcohol and other sponsor-only removals."
    )
    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=len(MINIMAL_ARVINE_HEADERS))
    header_row = 7
    for sequence, receipt in enumerate(receipts, start=1):
        people = max(1, int(receipt.get("number_of_people") or 1))
        traveller_share = round(float(receipt.get("total") or 0) / people, 2)
        ws.append(
            [
                sequence,
                receipt.get("document_date"),
                receipt.get("vendor"),
                receipt.get("description"),
                receipt.get("expense_type"),
                receipt.get("source_file"),
                people,
                receipt.get("currency"),
                receipt.get("total"),
                traveller_share,
                receipt.get("total_cad"),
                readable_paid_by(receipt.get("paid_by")),
                receipt.get("arvine_reimbursable_cad"),
                (
                    receipt.get("total_cad")
                    if receipt.get("paid_by") == "arvine_corporate_bmo"
                    else 0.0
                ),
                receipt.get("ivado_claimable_cad"),
                receipt.get("ivado_excluded_cad"),
            ]
        )
        source_cell = ws.cell(ws.max_row, 6)
        source_cell.hyperlink = f"expenses_receipts/{receipt.get('source_file')}"
        source_cell.style = "Hyperlink"
    total_row = ws.max_row + 1
    ws.cell(total_row, 1, "TOTAL")
    for column in (9, 10, 11, 13, 14, 15, 16):
        letter = get_column_letter(column)
        ws.cell(total_row, column, f"=SUM({letter}{header_row + 1}:{letter}{total_row - 1})")
    ws.freeze_panes = f"A{header_row + 1}"
    set_filter_range(ws, len(MINIMAL_ARVINE_HEADERS), max(total_row - 1, header_row + 1))


def write_minimal_accounting_sheet(ws, report: dict, receipts: list[dict]) -> None:
    """Write posting-ready journal rows for ordinary and IVADO-sponsored trips."""

    headers = [
        "Invoice Date",
        "Vendor / Customer",
        "Description",
        "Account",
        "Counter-Account",
        "Amount CAD",
        "Posting Step",
    ]
    ws.append(headers)
    report_date = excel_date(str(report.get("report_date") or ""))
    trip_label = str(report.get("description") or report.get("trip_id") or "Business trip")
    traveller = str(report.get("traveller") or "Traveller")
    summary = report.get("accounting_summary", {})
    employee_total = round(float(report.get("employee_reimbursement_total_cad") or 0), 2)

    rows: list[list[object]] = []
    if report.get("claim_program") == "ivado_sponsored":
        employee_ivado_claim = round(
            sum(
                float(receipt.get("ivado_claimable_cad") or 0)
                for receipt in receipts
                if receipt.get("paid_by") == "employee_personal"
            ),
            2,
        )
        alcohol_borne = round(
            sum(employee_ivado_alcohol_cad(receipt) for receipt in receipts),
            2,
        )
        alcohol_deductible = round(alcohol_borne * 0.5, 2)
        alcohol_nondeductible = round(alcohol_borne - alcohol_deductible, 2)
        meal_deductible = round(float(summary.get("meal_deductible_cad") or 0), 2)
        meal_nondeductible = round(float(summary.get("meal_non_deductible_cad") or 0), 2)
        origin_rows = [
            (
                "Reimbursable travel expense (passthrough IVADO)",
                "Expenses Recoverable from Clients",
                employee_ivado_claim,
            ),
            (
                "Arvine-borne alcohol – deductible 50%",
                "Meals – Deductible (50%)",
                alcohol_deductible,
            ),
            (
                "Arvine-borne alcohol – non-deductible 50%",
                "Meals – Non-deductible (50%)",
                alcohol_nondeductible,
            ),
        ]
        for description, account, amount in origin_rows:
            if amount or "passthrough" in description:
                rows.append(
                    [
                        report_date,
                        traveller,
                        description,
                        account,
                        "Shareholder Current Account",
                        round(amount, 2),
                        "Record traveller expenses",
                    ]
                )

        other_meal_rows = [
            (
                "Arvine-borne IVADO exclusion – meal deductible",
                "Meals – Deductible (50%)",
                round(meal_deductible - alcohol_deductible, 2),
            ),
            (
                "Arvine-borne IVADO exclusion – meal non-deductible",
                "Meals – Non-deductible (50%)",
                round(meal_nondeductible - alcohol_nondeductible, 2),
            ),
        ]
        for description, account, amount in other_meal_rows:
            if amount > 0.005:
                rows.append(
                    [
                        report_date,
                        traveller,
                        description,
                        account,
                        "Shareholder Current Account",
                        amount,
                        "Record traveller expenses",
                    ]
                )

        # Preserve a posting path for an unusual IVADO-only exclusion that is
        # not a meal/alcohol item instead of silently forcing it into alcohol.
        non_meal = round(float(summary.get("travel_non_meal_cad") or 0), 2)
        if non_meal:
            rows.append(
                [
                    report_date,
                    traveller,
                    "Arvine-borne IVADO exclusion – non-meal",
                    "Travel – Non-meal",
                    "Shareholder Current Account",
                    non_meal,
                    "Record traveller expenses",
                ]
            )

        sponsor_claim = round(float(report.get("ivado_claim_total_cad") or 0), 2)
        rows.extend(
            [
                [
                    report_date,
                    "IVADO Labs",
                    f"{trip_label} - Invoice sent to IL by Arvine",
                    "Accounts Receivable",
                    "Expenses Recoverable from Clients",
                    sponsor_claim,
                    "Invoice IVADO",
                ],
                [
                    report_date,
                    "IVADO Labs",
                    f"{trip_label} - Reimbursement from IL",
                    "Bank – Checking",
                    "Accounts Receivable",
                    sponsor_claim,
                    "Receive IVADO reimbursement",
                ],
            ]
        )
        control_label = "Traveller expense booking"
        control_expected = employee_total
        origin_amounts = [row[5] for row in rows if row[6] == "Record traveller expenses"]
    else:
        vendor = f"Expense report – {trip_label}"
        components = [
            (
                "Airfare + taxi, excluding GST/QST",
                "Travel – Non-meal",
                "travel_non_meal_cad",
            ),
            (
                "Meal deductible 50%, excluding GST/QST",
                "Meals – Deductible (50%)",
                "meal_deductible_cad",
            ),
            (
                "Meal non-deductible 50%, excluding GST/QST",
                "Meals – Non-deductible (50%)",
                "meal_non_deductible_cad",
            ),
            ("GST paid on travel + meal", "GST Receivable", "gst_receivable_cad"),
            ("QST paid on travel + meal", "QST Receivable", "qst_receivable_cad"),
        ]
        for description, account, key in components:
            rows.append(
                [
                    report_date,
                    vendor,
                    description,
                    account,
                    "Shareholder Current Account",
                    round(float(summary.get(key) or 0), 2),
                    "Record expense report",
                ]
            )
        rows.append(
            [
                report_date,
                f"{traveller} - {trip_label}",
                "Payment of business trip",
                "Shareholder Current Account",
                "Bank – Checking",
                employee_total,
                "Pay traveller",
            ]
        )
        control_label = "Expense report booking"
        control_expected = employee_total
        origin_amounts = [row[5] for row in rows if row[6] == "Record expense report"]

    for row in rows:
        ws.append(row)
    control_row = ws.max_row + 2
    ws.cell(control_row, 1, control_label)
    ws.cell(control_row, 6, round(sum(float(value) for value in origin_amounts), 2))
    ws.cell(control_row + 1, 1, "Expected traveller reimbursement")
    ws.cell(control_row + 1, 6, control_expected)
    ws.cell(control_row + 2, 1, "Difference")
    ws.cell(control_row + 2, 6, f"=ROUND(F{control_row}-F{control_row + 1},2)")
    ws.cell(control_row + 3, 1, "Status")
    ws.cell(control_row + 3, 6, f'=IF(ABS(F{control_row + 2})<=0.02,"PASS","REVIEW")')
    ws.freeze_panes = "A2"
    set_filter_range(ws, len(headers), 1 + len(rows))


def employee_ivado_alcohol_cad(receipt: dict) -> float:
    """Allocate a receipt's exact CAD removal to its excluded alcohol lines."""

    if receipt.get("paid_by") != "employee_personal":
        return 0.0
    removed_cad = round(
        max(
            0.0,
            float(receipt.get("total_cad") or 0) - float(receipt.get("ivado_claimable_cad") or 0),
        ),
        2,
    )
    if removed_cad <= 0:
        return 0.0
    alcohol_original = max(0.0, ivado_removed_original(receipt, alcohol=True))
    other_original = max(0.0, ivado_removed_original(receipt, alcohol=False))
    removed_original = alcohol_original + other_original
    if removed_original <= 0:
        return 0.0
    return round(removed_cad * alcohol_original / removed_original, 2)


def replace_support_sheet(
    workbook: Workbook,
    name: str,
    index: int,
    *legacy_names: str,
):
    for candidate in dict.fromkeys((name, *legacy_names)):
        if candidate in workbook.sheetnames:
            workbook.remove(workbook[candidate])
    return workbook.create_sheet(name, index)


def load_reconciliation_for_report(trip_dir: Path) -> dict:
    try:
        from nlp_expenses.reconciliation import reconciliation_view

        return reconciliation_view(trip_dir)
    except Exception:
        return {"expenses": [], "transactions": []}


def write_ivado_template_sheet(ws, report: dict, receipts: list[dict]) -> None:
    claim_receipts = [
        receipt for receipt in receipts if float(receipt.get("ivado_claimable_cad") or 0) > 0
    ]
    if len(claim_receipts) > 50:
        raise ValueError(
            "The supplied IVADO form supports 50 claimed receipts. "
            "Split the trip into more than one IVADO form."
        )

    ws["D6"] = report.get("traveller")
    ws["D7"] = excel_date(str(report.get("report_date") or ""))
    ws["D7"].number_format = "yyyy-mm-dd"
    reset_ivado_template_calculations(ws)

    for sequence, receipt in enumerate(claim_receipts, start=1):
        row = 14 + sequence
        claim_cad = round(float(receipt.get("ivado_claimable_cad") or 0), 2)
        tax_values = ivado_tax_values(receipt, claim_cad)
        category_amount = round(max(0.0, claim_cad - sum(tax_values.values())), 2)
        people = max(1, int(receipt.get("number_of_people") or 1))
        description = str(receipt.get("description") or receipt.get("expense_type") or "Expense")
        annotations = []
        if people > 1:
            annotations.append(f"traveller share (1 of {people})")
        if ivado_removed_original(receipt, alcohol=True) > 0.005:
            annotations.append("alcohol removed")
        if ivado_removed_original(receipt, alcohol=False) > 0.005:
            annotations.append("other item removed")
        if annotations:
            description = f"{description} — {', '.join(annotations)}"

        ws.cell(row, 1, sequence)
        ws.cell(row, 2, excel_date(receipt.get("document_date")))
        ws.cell(row, 2).number_format = "yyyy-mm-dd"
        ws.cell(row, 3, receipt.get("vendor"))
        ws.cell(row, 4, description)
        ws.cell(row, 5, ivado_location(receipt))
        ws.cell(row, 6, "Yes/Oui")
        ws.cell(
            row,
            7,
            (
                ivado_original_claim_amount(receipt)
                if str(receipt.get("currency") or "").upper() != "CAD"
                else None
            ),
        )
        ws.cell(row, 8, claim_cad)
        ws.cell(row, ivado_category_column(receipt.get("expense_type")), category_amount)
        ws.cell(row, 18, tax_values["gst"])
        ws.cell(row, 19, tax_values["qst"])
        ws.cell(row, 20, tax_values["hst"])
        ws.cell(row, 21, tax_values["other"])

    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 75


def reset_ivado_template_calculations(ws) -> None:
    for row in range(15, 65):
        for column in range(2, 22):
            ws.cell(row, column).value = None
        ws.cell(row, 1, row - 14)
        ws.cell(row, 22, f"=SUM(I{row}:U{row})")
        ws.cell(row, 23, f"=ROUND(V{row}-H{row},2)")
        ws.cell(row, 24).value = None
        ws.cell(row, 25, f'=IF(X{row}="ME",R{row}*0.5,IF(X{row}="NR",R{row},0))')
        ws.cell(row, 26, f'=IF(X{row}="ME",S{row}*0.5,IF(X{row}="NR",S{row},0))')
        ws.cell(row, 27, f'=IF(X{row}="ME",T{row}*0.5,IF(X{row}="NR",T{row},0))')
        ws.cell(row, 28, f'=IF(X{row}="ME",U{row}*0.5,IF(X{row}="NR",U{row},0))')
        ws.cell(row, 29, f"=SUM(Y{row}:AB{row})")
    for column in range(9, 24):
        letter = get_column_letter(column)
        ws.cell(65, column, f"=SUM({letter}15:{letter}64)")
    for column in range(25, 30):
        letter = get_column_letter(column)
        ws.cell(65, column, f"=SUM({letter}15:{letter}64)")
    ws["W66"] = '=IF(ABS(W65)<=0.02,"OK","ERROR")'
    ws["I68"] = "=IF(I65>0,I65+$AC$65,0)"
    for column in range(10, 18):
        letter = get_column_letter(column)
        ws.cell(68, column, f"={letter}65")
    for column, adjustment_column in zip(range(18, 22), range(25, 29), strict=True):
        letter = get_column_letter(column)
        adjustment = get_column_letter(adjustment_column)
        ws.cell(68, column, f"={letter}65-{adjustment}65")
    ws["V68"] = "=SUM(I68:U68)"
    ws["V70"] = "=V65-V68"
    ws["W70"] = '=IF(ABS(V70)<=0.02,"Aucune différence / No difference","Difference")'


def write_ivado_card_statement_sheet(
    ws,
    report: dict,
    receipts: list[dict],
    reconciliation: dict,
) -> None:
    receipt_by_file = {str(item.get("source_file") or ""): item for item in receipts}
    sequence_by_file = {
        str(item.get("source_file") or ""): sequence
        for sequence, item in enumerate(receipts, start=1)
    }
    expense_by_file = {
        str(item.get("source_file") or ""): item
        for item in reconciliation.get("expenses", [])
        if isinstance(item, dict)
    }
    transactions = [
        item
        for item in reconciliation.get("transactions", [])
        if isinstance(item, dict) and item.get("expense_file") in receipt_by_file
    ]
    transaction_groups: dict[str, list[dict]] = {}
    for transaction in transactions:
        transaction_groups.setdefault(str(transaction.get("expense_file")), []).append(transaction)

    ws["A1"] = "Card statement transactions matched to the IVADO claim"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(IVADO_CARD_AUDIT_HEADERS))
    ws["A2"] = "Traveller"
    ws["B2"] = report.get("traveller")
    ws["D2"] = "IVADO claim CAD"
    ws["E2"] = report.get("ivado_claim_total_cad")
    ws["G2"] = "Gross matched statement CAD"
    ws["H2"] = 0.0
    ws["A3"] = (
        "Only transactions mapped to trip receipts are shown. Allocated claim and removal "
        "amounts reconcile exactly when multiple statement transactions fund one receipt."
    )
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(IVADO_CARD_AUDIT_HEADERS))
    header_row = 5
    for column, header in enumerate(IVADO_CARD_AUDIT_HEADERS, start=1):
        ws.cell(header_row, column, header)

    for source_file, group in transaction_groups.items():
        receipt = receipt_by_file[source_file]
        weights = [abs(float(item.get("cad_amount") or 0)) for item in group]
        allocated_claim = allocate_money(float(receipt.get("ivado_claimable_cad") or 0), weights)
        allocated_removed = allocate_money(float(receipt.get("ivado_excluded_cad") or 0), weights)
        for transaction, claim_part, removed_part in zip(
            group,
            allocated_claim,
            allocated_removed,
            strict=True,
        ):
            expense = expense_by_file.get(source_file, {})
            ws.append(
                [
                    sequence_by_file[source_file],
                    excel_date(transaction.get("transaction_date")),
                    transaction.get("account_label"),
                    transaction.get("description"),
                    transaction.get("purchase_amount"),
                    transaction.get("purchase_currency"),
                    transaction.get("cad_amount"),
                    readable_statement_basis(expense.get("statement_amount_basis")),
                    transaction.get("match_status"),
                    source_file,
                    excel_date(receipt.get("document_date")),
                    receipt.get("vendor"),
                    max(1, int(receipt.get("number_of_people") or 1)),
                    claim_part,
                    removed_part,
                    ", ".join(str(value) for value in transaction.get("source_files", [])),
                ]
            )
            ws.cell(ws.max_row, 2).number_format = "yyyy-mm-dd"
            ws.cell(ws.max_row, 11).number_format = "yyyy-mm-dd"
            add_relative_hyperlink(ws.cell(ws.max_row, 10), "expenses_receipts", source_file)
            source_files = transaction.get("source_files") or []
            if source_files:
                add_relative_hyperlink(
                    ws.cell(ws.max_row, 16),
                    "card_statements",
                    str(source_files[0]),
                )

    ws["H2"] = f"=SUM(G6:G{max(6, ws.max_row)})"

    style_support_sheet(ws, header_row, len(IVADO_CARD_AUDIT_HEADERS))
    for column in (5, 7, 14, 15):
        for cell in ws[get_column_letter(column)][header_row:]:
            cell.number_format = "#,##0.00;[Red](#,##0.00)"
    ws.freeze_panes = "A6"
    add_table_if_populated(
        ws,
        header_row,
        len(IVADO_CARD_AUDIT_HEADERS),
        "IVADOCardStatementAudit",
    )


def write_ivado_reconciliation_sheet(
    ws,
    report: dict,
    receipts: list[dict],
    reconciliation: dict,
    *,
    include_line_items: bool = False,
) -> None:
    expense_by_file = {
        str(item.get("source_file") or ""): item
        for item in reconciliation.get("expenses", [])
        if isinstance(item, dict)
    }
    ws["A1"] = "IVADO Receipt Reconciliation"
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(IVADO_RECONCILIATION_HEADERS),
    )
    ws["A2"] = "IVADO claim CAD"
    ws["B2"] = report.get("ivado_claim_total_cad")
    ws["D2"] = "Traveller reimbursement CAD"
    ws["E2"] = report.get("employee_reimbursement_total_cad")
    ws["G2"] = "IVADO removed CAD"
    ws["H2"] = report.get("ivado_excluded_total_cad")
    ws["A3"] = (
        "Gross statement CAD establishes the FX evidence. The IVADO claim and traveller "
        "reimbursement are independent: Arvine reimburses the full reviewed business share, "
        "while IVADO excludes alcohol and any manually reviewed sponsor removals."
    )
    ws.merge_cells(
        start_row=3,
        start_column=1,
        end_row=3,
        end_column=len(IVADO_RECONCILIATION_HEADERS),
    )
    header_row = 5
    for column, header in enumerate(IVADO_RECONCILIATION_HEADERS, start=1):
        ws.cell(header_row, column, header)

    for sequence, receipt in enumerate(receipts, start=1):
        source_file = str(receipt.get("source_file") or "")
        expense = expense_by_file.get(source_file, {})
        people = max(1, int(receipt.get("number_of_people") or 1))
        row = ws.max_row + 1
        ws.append(
            [
                sequence,
                excel_date(receipt.get("document_date")),
                receipt.get("vendor"),
                receipt.get("description"),
                receipt.get("expense_type"),
                source_file,
                receipt.get("currency"),
                receipt.get("total"),
                people,
                readable_statement_basis(expense.get("statement_amount_basis")),
                None,
                receipt.get("total_cad"),
                None,
                ivado_original_claim_amount(receipt),
                ivado_removed_original(receipt, alcohol=True),
                receipt.get("ivado_claimable_cad"),
                receipt.get("ivado_excluded_cad"),
                receipt.get("arvine_reimbursable_cad"),
            ]
        )
        ws.cell(row, 2).number_format = "yyyy-mm-dd"
        ws.cell(row, 11, f"=IFERROR(H{row}/MAX(1,I{row}),0)")
        ws.cell(row, 13, f'=IF(OR(K{row}=0,L{row}=0),"",L{row}/K{row})')
        add_relative_hyperlink(ws.cell(row, 6), "expenses_receipts", source_file)

    last_receipt_row = ws.max_row
    style_support_sheet(ws, header_row, len(IVADO_RECONCILIATION_HEADERS))
    for column in (8, 11, 14, 15, 16, 17, 18):
        for cell in ws[get_column_letter(column)][header_row:]:
            cell.number_format = "#,##0.00;[Red](#,##0.00)"
    for cell in ws["M"][header_row:]:
        cell.number_format = "0.000000"
    ws.row_dimensions[header_row].height = 48
    for cell in ws[header_row]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
    if last_receipt_row > header_row:
        table = Table(
            displayName="IVADOReceiptReconciliation",
            ref=(
                f"A{header_row}:{get_column_letter(len(IVADO_RECONCILIATION_HEADERS))}"
                f"{last_receipt_row}"
            ),
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    if include_line_items:
        line_title_row = last_receipt_row + 3
        line_header_row = line_title_row + 1
        ws.cell(line_title_row, 1, "Receipt line items and IVADO removals")
        ws.merge_cells(
            start_row=line_title_row,
            start_column=1,
            end_row=line_title_row,
            end_column=len(IVADO_LINE_AUDIT_HEADERS),
        )
        for column, header in enumerate(IVADO_LINE_AUDIT_HEADERS, start=1):
            ws.cell(line_header_row, column, header)
        for sequence, receipt in enumerate(receipts, start=1):
            append_ivado_line_audit_rows(ws, sequence, receipt)
        style_table_header(ws, line_header_row, len(IVADO_LINE_AUDIT_HEADERS))
        for column in (8, 9, 13, 14):
            for cell in ws[get_column_letter(column)][line_header_row:]:
                cell.number_format = "#,##0.00;[Red](#,##0.00)"
        if ws.max_row > line_header_row:
            table = Table(
                displayName="IVADOReceiptLineAudit",
                ref=f"A{line_header_row}:P{ws.max_row}",
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            ws.add_table(table)
    ws.freeze_panes = "A6"
    set_reconciliation_widths(ws)


def append_ivado_line_audit_rows(ws, sequence: int, receipt: dict) -> None:
    people = max(1, int(receipt.get("number_of_people") or 1))
    items = [
        item
        for item in receipt.get("line_items", [])
        if abs(float(item.get("amount") or 0)) >= 0.005
    ] or [
        {
            "line_id": f"{receipt.get('id')}-TOTAL",
            "description": receipt.get("description") or "Receipt total",
            "amount": receipt.get("total") or 0.0,
            "is_alcohol": False,
            "included_in_arvine": receipt.get("included_in_arvine", True),
            "included_in_ivado": receipt.get("included_in_ivado", True),
            "ivado_exclusion_reason": receipt.get("ivado_exclusion_reason"),
            "review_note": "No separate receipt lines were extracted.",
        }
    ]
    removed_shares = [
        (
            round(float(item.get("amount") or 0) / people, 2)
            if item.get("included_in_arvine", True) and not item.get("included_in_ivado", True)
            else 0.0
        )
        for item in items
    ]
    removed_allocations = allocate_money(
        float(receipt.get("ivado_excluded_cad") or 0),
        [max(0.0, value) for value in removed_shares],
    )
    for item, removed_share, removed_cad in zip(
        items,
        removed_shares,
        removed_allocations,
        strict=True,
    ):
        amount = float(item.get("amount") or 0)
        removed = item.get("included_in_arvine", True) and not item.get("included_in_ivado", True)
        row = ws.max_row + 1
        ws.append(
            [
                sequence,
                excel_date(receipt.get("document_date")),
                receipt.get("vendor"),
                receipt.get("source_file"),
                people,
                item.get("line_id"),
                item.get("description"),
                amount,
                round(amount / people, 2),
                receipt.get("currency"),
                "Yes" if item.get("is_alcohol") else "No",
                "Yes" if item.get("included_in_ivado", True) else "No",
                removed_share if removed else 0.0,
                removed_cad if removed else 0.0,
                item.get("ivado_exclusion_reason"),
                item.get("review_note"),
            ]
        )
        ws.cell(row, 2).number_format = "yyyy-mm-dd"
        add_relative_hyperlink(
            ws.cell(row, 4),
            "expenses_receipts",
            str(receipt.get("source_file") or ""),
        )
        if removed:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
        if item.get("is_alcohol"):
            ws.cell(row, 11).fill = PatternFill("solid", fgColor="F4CCCC")
            ws.cell(row, 11).font = Font(color="9C0006", bold=True)


def ivado_removed_original(receipt: dict, alcohol: bool) -> float:
    people = max(1, int(receipt.get("number_of_people") or 1))
    if receipt.get("included_in_arvine", True) and not receipt.get("included_in_ivado", True):
        return round(float(receipt.get("total") or 0) / people, 2) if not alcohol else 0.0
    return round(
        sum(
            float(item.get("amount") or 0)
            for item in receipt.get("line_items", [])
            if item.get("included_in_arvine", True)
            and not item.get("included_in_ivado", True)
            and bool(item.get("is_alcohol")) is alcohol
        )
        / people,
        2,
    )


def allocate_money(total: float, weights: list[float]) -> list[float]:
    if not weights:
        return []
    positive_total = sum(max(0.0, weight) for weight in weights)
    if positive_total <= 0:
        return [round(total, 2)] + [0.0] * (len(weights) - 1)
    allocations = [round(total * max(0.0, weight) / positive_total, 2) for weight in weights]
    allocations[-1] = round(allocations[-1] + round(total - sum(allocations), 2), 2)
    return allocations


def readable_statement_basis(value: object) -> str:
    return "Traveller share" if value == "personal_share" else "Full receipt"


def add_relative_hyperlink(cell, folder: str, filename: str) -> None:
    if not filename:
        return
    cell.hyperlink = f"{folder}/{filename}"
    cell.style = "Hyperlink"


def add_table_if_populated(ws, header_row: int, width: int, name: str) -> None:
    if ws.max_row <= header_row:
        return
    table = Table(
        displayName=name,
        ref=f"A{header_row}:{get_column_letter(width)}{ws.max_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def style_support_sheet(ws, header_row: int, width: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws["A1"].font = Font(bold=True, size=16, color="1F4E78")
    for cell in ws[2]:
        if cell.value not in (None, ""):
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
    ws["A3"].font = Font(italic=True, color="666666")
    style_table_header(ws, header_row, width)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=width):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    set_support_widths(ws, width)


def set_support_widths(ws, width: int) -> None:
    for column in range(1, width + 1):
        letter = get_column_letter(column)
        max_length = max(
            (len(str(cell.value)) for cell in ws[letter] if cell.value not in (None, "")),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(max_length + 2, 11), 34)


def set_reconciliation_widths(ws) -> None:
    widths = {
        "A": 9,
        "B": 12,
        "C": 21,
        "D": 24,
        "E": 12,
        "F": 25,
        "G": 9,
        "H": 15,
        "I": 8,
        "J": 16,
        "K": 16,
        "L": 17,
        "M": 13,
        "N": 17,
        "O": 16,
        "P": 15,
        "Q": 15,
        "R": 16,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.sheet_view.zoomScale = 75


def configure_ivado_template_workbook(workbook: Workbook) -> None:
    workbook.active = workbook.sheetnames.index(IVADO_TEMPLATE_SHEET)
    workbook.defined_names["Liste"] = DefinedName(
        "Liste",
        attr_text=f"'{IVADO_TEMPLATE_SHEET}'!$A$136:$A$137",
    )
    for validation in workbook[IVADO_TEMPLATE_SHEET].data_validations.dataValidation:
        if validation.formula1 == "#REF!":
            validation.formula1 = "Liste"
    if workbook.calculation is None:
        workbook.calculation = CalcProperties()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"


def ivado_original_claim_amount(receipt: dict) -> float:
    people = max(1, int(receipt.get("number_of_people") or 1))
    items = receipt.get("line_items") or []
    excluded = any(
        item.get("included_in_arvine", True) and not item.get("included_in_ivado", True)
        for item in items
    )
    if excluded:
        included_total = sum(
            float(item.get("amount") or 0)
            for item in items
            if item.get("included_in_arvine", True) and item.get("included_in_ivado", True)
        )
        if included_total > 0:
            return round(included_total / people, 2)
    return round(float(receipt.get("total") or 0) / people, 2)


def ivado_tax_values(receipt: dict, claim_cad: float) -> dict[str, float]:
    result = {"gst": 0.0, "qst": 0.0, "hst": 0.0, "other": 0.0}
    if not ivado_is_canadian(receipt):
        return result
    original_total = float(receipt.get("total") or 0)
    fx_rate = float(receipt.get("fx_rate") or (1.0 if receipt.get("currency") == "CAD" else 0.0))
    if original_total <= 0 or fx_rate <= 0:
        return result
    claim_ratio = ivado_original_claim_amount(receipt) / original_total
    gst_hst_cad = round(float(receipt.get("gst") or 0) * fx_rate * claim_ratio, 2)
    qst_cad = round(float(receipt.get("qst") or 0) * fx_rate * claim_ratio, 2)
    province = str(receipt.get("province") or "").strip().upper()
    if province in {
        "ON",
        "ONTARIO",
        "NB",
        "NEW BRUNSWICK",
        "NL",
        "NEWFOUNDLAND",
        "NS",
        "NOVA SCOTIA",
        "PE",
        "PEI",
        "PRINCE EDWARD ISLAND",
    }:
        result["hst"] = min(claim_cad, gst_hst_cad)
    else:
        result["gst"] = min(claim_cad, gst_hst_cad)
        result["qst"] = min(max(0.0, claim_cad - result["gst"]), qst_cad)
    return result


def ivado_is_canadian(receipt: dict) -> bool:
    country = str(receipt.get("country") or "").strip().upper()
    province = str(receipt.get("province") or "").strip()
    if country in {"CANADA", "CA", "CAN"}:
        return True
    if country and country not in {"CANADA", "CA", "CAN"}:
        return False
    return bool(province) or str(receipt.get("currency") or "").upper() == "CAD"


def ivado_location(receipt: dict) -> str:
    if not ivado_is_canadian(receipt):
        return "NR - Outside Canada/À l'extérieur du Canada"
    province = str(receipt.get("province") or "").strip().upper()
    labels = {
        "QC": "QC - Quebec",
        "QUEBEC": "QC - Quebec",
        "QUÉBEC": "QC - Quebec",
        "ON": "ON - Ontario",
        "ONTARIO": "ON - Ontario",
        "AB": "AB - Alberta",
        "ALBERTA": "AB - Alberta",
        "BC": "BC - Brit. Columbia/Columbie Brit.",
        "BRITISH COLUMBIA": "BC - Brit. Columbia/Columbie Brit.",
        "MB": "MB - Manitoba",
        "MANITOBA": "MB - Manitoba",
        "NB": "NB - New Brunswick/Nouveau Brunswick",
        "NEW BRUNSWICK": "NB - New Brunswick/Nouveau Brunswick",
        "NL": "NL - Newfoundland/Terre-Neuve",
        "NEWFOUNDLAND": "NL - Newfoundland/Terre-Neuve",
        "NS": "NS - Nova Scotia/Nouvelle Écosse",
        "NOVA SCOTIA": "NS - Nova Scotia/Nouvelle Écosse",
        "PE": "PE - Prince Edward Island/ île du Prince Édouard",
        "PEI": "PE - Prince Edward Island/ île du Prince Édouard",
        "SK": "SK - Saskatchewan",
        "SASKATCHEWAN": "SK - Saskatchewan",
    }
    return labels.get(province, "NR - Outside Canada/À l'extérieur du Canada")


def ivado_category_column(expense_type: object) -> int:
    value = str(expense_type or "").strip().lower()
    if "conference" in value:
        return 10
    if "training" in value:
        return 11
    if "membership" in value or "subscription" in value:
        return 12
    if "meal" in value or "food" in value or "restaurant" in value:
        return 14
    if any(
        token in value
        for token in ("flight", "hotel", "transport", "taxi", "uber", "parking", "travel")
    ):
        return 13
    if "office" in value:
        return 9
    return 16


def readable_paid_by(value: object) -> str:
    return "Traveller personal" if value == "employee_personal" else "Company card"


def readable_claim_program(value: object) -> str:
    return "IVADO-reimbursed trip" if value == "ivado_sponsored" else "Own-company reimbursement"


def style_minimal_report_workbook(workbook: Workbook) -> None:
    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in range(1, ws.max_column + 1):
            letter = get_column_letter(column)
            max_length = max(
                (len(str(cell.value)) for cell in ws[letter] if cell.value not in (None, "")),
                default=10,
            )
            ws.column_dimensions[letter].width = min(max(max_length + 2, 12), 38)
    report_ws = workbook["Expense Report"]
    report_ws["A1"].font = Font(bold=True, size=18, color="1F4E78")
    report_ws.merge_cells("A1:P1")
    style_table_header(report_ws, 7, len(MINIMAL_ARVINE_HEADERS))
    total_row = report_ws.max_row
    for cell in report_ws[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for column in (9, 10, 11, 13, 14, 15, 16):
        for cell in report_ws[get_column_letter(column)][7:]:
            cell.number_format = "#,##0.00;[Red]-#,##0.00"
    accounting_ws = workbook["Accounting Rows"]
    style_table_header(accounting_ws, 1, 7)
    accounting_ws.column_dimensions["A"].width = 15
    accounting_ws.column_dimensions["B"].width = 32
    accounting_ws.column_dimensions["C"].width = 52
    accounting_ws.column_dimensions["D"].width = 35
    accounting_ws.column_dimensions["E"].width = 36
    accounting_ws.column_dimensions["F"].width = 18
    accounting_ws.column_dimensions["G"].width = 30
    for cell in accounting_ws["A"][1:]:
        if hasattr(cell.value, "year"):
            cell.number_format = "yyyy-mm-dd"
    for cell in accounting_ws["F"][1:]:
        cell.number_format = "#,##0.00;[Red]-#,##0.00"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"


def split_manifest_records(records: list[dict]) -> tuple[list[dict], dict]:
    receipts = [record for record in records if record.get("kind") == "receipt"]
    reports = [record for record in records if record.get("kind") == "trip_report"]
    if len(reports) != 1:
        raise ValueError("Expected exactly one trip report record.")
    return receipts, reports[0]
