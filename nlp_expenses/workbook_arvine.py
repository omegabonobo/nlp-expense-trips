from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

from nlp_expenses.models import Expense, NormalizedTransaction
from nlp_expenses.trip_metadata import trip_metadata
from nlp_expenses.trips import source_file_key
from nlp_expenses.workbook_common import (
    append_clean,
    excel_date,
    set_filter_range,
    set_widths,
    style_table_header,
)

CHECK_FILL = PatternFill("solid", fgColor="FFF2CC")

ARVINE_STATEMENT_HEADERS = [
    "transaction_group_id",
    "funding_leg_id",
    "transaction_date",
    "posted_date",
    "provider",
    "account_label",
    "cardholder",
    "description",
    "category",
    "transaction_type",
    "status",
    "direction",
    "match_eligible",
    "purchase_amount",
    "purchase_currency",
    "settlement_amount",
    "settlement_currency",
    "CAD_amount",
    "CAD_completeness",
    "expense_id",
    "suggested_expense_id",
    "match_status",
    "match_confidence",
    "source_file",
    "source_row",
    "normalization_status",
    "review_note",
    "cad_conversion_rate",
    "cad_conversion_week_start",
    "cad_conversion_week_end",
    "cad_conversion_method",
    "cad_conversion_route",
    "cad_conversion_source",
    "cad_conversion_source_urls",
]

ARVINE_SUMMARY_HEADERS = [
    "Invoice Date",
    "Vendor / Customer",
    "Description",
    "Account",
    "Counter-Account",
    "Currency",
    "FX Rate",
    "Subtotal (in currency)",
    "GST/HST (in currency)",
    "QST (in currency)",
    "Total with Tx (in currency)",
    "Subtotal (before tax) CAD",
    "GST/HST CAD",
    "QST CAD",
    "Amount total CAD",
    "GST #",
    "QST #",
]


def write_arvine_summary_sheet(ws, trip_dir: Path, accounting_profile: dict) -> None:
    trip_info = trip_metadata(trip_dir)
    traveller_company = " / ".join(
        value
        for value in (
            trip_info["company"] or accounting_profile["company_legal_name"],
            trip_info["traveller"],
        )
        if value
    )
    purpose_client = " / ".join(
        value for value in (trip_info["business_purpose"], trip_info["client_project"]) if value
    )
    metadata = [
        ("Report date", date.today()),
        ("Trip label", trip_dir.name),
        ("Company / traveller", traveller_company),
        ("Business purpose / client", purpose_client),
        ("Counter-account", accounting_profile["counter_account"]),
        ("Mode", "Own-company reimbursement"),
    ]
    for row, (label, value) in enumerate(metadata, start=1):
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
    ws["D1"] = "Legend"
    ws["D2"] = "Editable input"
    ws["D3"] = "Linked formula"
    ws["D4"] = "Calculated formula"
    ws["D5"] = "Review warning"
    try:
        from nlp_expenses.reconciliation import load_reconciliation_state

        reconciliation_state = load_reconciliation_state(trip_dir) or {}
        coverage_confirmation = reconciliation_state.get("coverage_confirmation") or {}
    except Exception:
        coverage_confirmation = {}
    ws["D6"] = "Statement coverage"
    ws["E6"] = (
        f"Confirmed {coverage_confirmation.get('confirmed_at')}: "
        f"{coverage_confirmation.get('note') or 'No unresolved coverage gaps.'}"
        if coverage_confirmation
        else "Not confirmed"
    )
    ws["D7"] = "Coverage gaps"
    ws["E7"] = coverage_confirmation.get("gap_count", "")
    ws["S1"] = "Applied accounting assumptions"
    profile_snapshot = [
        ("Profile version", accounting_profile["version"]),
        ("Effective date", accounting_profile["effective_date"]),
        ("Reimbursement type", accounting_profile["traveller_reimbursement_type"]),
        ("GST/HST registrant", accounting_profile["gst_hst_registrant"]),
        ("QST registrant", accounting_profile["qst_registrant"]),
        ("Commercial use", accounting_profile["commercial_use_pct"]),
        ("Normal tax recovery", accounting_profile["normal_tax_recovery_pct"]),
        ("Meal tax recovery", accounting_profile["meal_tax_recovery_pct"]),
        ("Normal deduction", accounting_profile["normal_deduction_pct"]),
        ("Meal deduction", accounting_profile["meal_deduction_pct"]),
        ("Tax method", accounting_profile["tax_calculation_method"]),
    ]
    for profile_row, (label, value) in enumerate(profile_snapshot, start=2):
        ws.cell(profile_row, 19, label)
        ws.cell(profile_row, 20, value)
    ws["V1"] = "Trip metadata and policy"
    trip_snapshot = [
        ("Traveller", trip_info["traveller"]),
        ("Company", trip_info["company"]),
        ("Start date", trip_info["start_date"]),
        ("End date", trip_info["end_date"]),
        ("Origin", ", ".join(trip_info["origins"])),
        ("Destinations", ", ".join(trip_info["destinations"])),
        ("Business purpose", trip_info["business_purpose"]),
        ("Client / project", trip_info["client_project"]),
        ("Cost centre", trip_info["cost_centre"]),
        ("Approver", trip_info["approver"]),
        ("Payment method", trip_info["payment_method"]),
        ("Expected accounts", ", ".join(trip_info["expected_accounts"])),
        ("Policy profile", trip_info["policy_profile"]),
    ]
    for trip_row, (label, value) in enumerate(trip_snapshot, start=2):
        ws.cell(trip_row, 22, label)
        ws.cell(trip_row, 23, value)
    ws["V16"] = "Policy controls"
    policy_snapshot = [
        ("Receipt threshold CAD", trip_info["policy"]["receipt_required_threshold"]),
        ("Allowed categories", ", ".join(trip_info["policy"]["allowed_categories"])),
        ("Meal limit CAD", trip_info["policy"]["meal_limit_cad"]),
        ("Alcohol treatment", trip_info["policy"]["alcohol_treatment"]),
        ("Personal expense treatment", trip_info["policy"]["personal_expense_treatment"]),
        ("Mileage rate CAD", trip_info["policy"]["mileage_rate_cad"]),
        ("Per diem CAD", trip_info["policy"]["per_diem_cad"]),
        ("Statement buffer days", trip_info["policy"]["statement_coverage_buffer_days"]),
    ]
    for policy_row, (label, value) in enumerate(policy_snapshot, start=17):
        ws.cell(policy_row, 22, label)
        ws.cell(policy_row, 23, value)
    try:
        from nlp_expenses.reconciliation import reconciliation_view

        policy_warnings = reconciliation_view(trip_dir).get("policy_warnings", [])
    except Exception:
        policy_warnings = []
    ws["V26"] = "Policy warnings / exceptions"
    if policy_warnings:
        for warning_row, warning in enumerate(policy_warnings, start=27):
            ws.cell(warning_row, 22, warning["message"])
            ws.cell(
                warning_row,
                23,
                f"Exception: {warning['exception_note']}"
                if warning["resolved"]
                else "Needs review",
            )
    else:
        ws["V27"] = "No policy warnings"
    ws["S1"].font = Font(bold=True, color="1F4E78")
    ws["V1"].font = Font(bold=True, color="1F4E78")
    ws["V16"].font = Font(bold=True, color="1F4E78")
    ws["V26"].font = Font(bold=True, color="1F4E78")
    ws["D1"].font = Font(bold=True, color="1F4E78")
    ws["D2"].font = Font(color="0070C0")
    ws["D3"].font = Font(color="008000")
    ws["D4"].font = Font(color="000000")
    ws["D5"].fill = CHECK_FILL
    ws["D6"].font = Font(bold=True, color="1F4E78")
    ws["D7"].font = Font(bold=True, color="1F4E78")
    ws["E6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B1"].number_format = "yyyy-mm-dd"
    ws["A8"] = "Trip-level CAD journal"
    ws["A8"].font = Font(size=14, bold=True, color="1F4E78")
    for col, header in enumerate(ARVINE_SUMMARY_HEADERS, start=1):
        ws.cell(9, col, header)

    mapping = accounting_profile["account_mapping"]
    accounts = [
        (
            "Airfare, accommodation, transport and other non-meal expense, excluding recoverable GST/QST",
            mapping["non_meal"],
        ),
        ("Meal deductible portion, excluding recoverable GST/QST", mapping["meal_deductible"]),
        (
            "Meal non-deductible portion, excluding recoverable GST/QST",
            mapping["meal_nondeductible"],
        ),
        ("Recoverable GST/HST paid on travel and meals", mapping["gst_hst_receivable"]),
        ("Recoverable QST paid on travel and meals", mapping["qst_receivable"]),
    ]
    amount_formulas = [
        "=SUMIFS('expense_detail'!$AA$2:$AA$5000,'expense_detail'!$A$2:$A$5000,TRUE,'expense_detail'!$F$2:$F$5000,\"<>meal\")",
        "=SUMIFS('expense_detail'!$AB$2:$AB$5000,'expense_detail'!$A$2:$A$5000,TRUE,'expense_detail'!$F$2:$F$5000,\"meal\")",
        "=SUMIFS('expense_detail'!$AC$2:$AC$5000,'expense_detail'!$A$2:$A$5000,TRUE,'expense_detail'!$F$2:$F$5000,\"meal\")",
        "=SUMIFS('expense_detail'!$Y$2:$Y$5000,'expense_detail'!$A$2:$A$5000,TRUE)",
        "=SUMIFS('expense_detail'!$Z$2:$Z$5000,'expense_detail'!$A$2:$A$5000,TRUE)",
    ]
    for offset, ((description, account), amount_formula) in enumerate(
        zip(accounts, amount_formulas, strict=True), start=10
    ):
        ws.cell(offset, 1, "=$B$1")
        ws.cell(offset, 2, '=IF($B$3="","Expense report – "&$B$2,$B$3)')
        ws.cell(offset, 3, description)
        ws.cell(offset, 4, account)
        ws.cell(offset, 5, "=$B$5")
        ws.cell(offset, 6, "CAD")
        ws.cell(offset, 7, 1)
        if offset <= 12:
            ws.cell(offset, 12, amount_formula)
        elif offset == 13:
            ws.cell(offset, 13, amount_formula)
        else:
            ws.cell(offset, 14, amount_formula)
        ws.cell(offset, 15, amount_formula)
    ws.cell(15, 3, "Journal debit total")
    ws.cell(15, 15, "=SUM(O10:O14)")

    ws["A18"] = "Review checks"
    ws["A18"].font = Font(size=14, bold=True, color="1F4E78")
    checks = [
        ("Journal debits", "=O15", ""),
        (
            "Shareholder reimbursement",
            "=SUMIFS('expense_detail'!$U$2:$U$5000,'expense_detail'!$A$2:$A$5000,TRUE)",
            "",
        ),
        ("Journal balance difference", "=B20-B21", '=IF(ABS(B22)<0.01,"OK","REVIEW")'),
        (
            "Included receipts without a complete CAD amount",
            "=COUNTIFS('expense_detail'!$A$2:$A$5000,TRUE,'expense_detail'!$U$2:$U$5000,\"\")",
            '=IF(B23=0,"OK","REVIEW")',
        ),
        (
            "Included receipts without an acceptable statement match or manual override",
            "=COUNTIFS('expense_detail'!$A$2:$A$5000,TRUE,'expense_detail'!$T$2:$T$5000,\"\",'expense_detail'!$AH$2:$AH$5000,\"<>matched\")",
            '=IF(B24=0,"OK","REVIEW")',
        ),
        (
            "Tax-documentation warnings",
            "=COUNTIFS('expense_detail'!$A$2:$A$5000,TRUE,'expense_detail'!$AI$2:$AI$5000,\"review\")",
            '=IF(B25=0,"OK","REVIEW")',
        ),
        (
            "Receipt-extraction warnings",
            "=COUNTIFS('expense_detail'!$A$2:$A$5000,TRUE,'expense_detail'!$AG$2:$AG$5000,\"<>ok\")",
            '=IF(B26=0,"OK","REVIEW")',
        ),
    ]
    for col, header in enumerate(("Check", "Value", "Status"), start=1):
        ws.cell(19, col, header)
    for row, check in enumerate(checks, start=20):
        for col, value in enumerate(check, start=1):
            ws.cell(row, col, value)

    style_table_header(ws, 9, len(ARVINE_SUMMARY_HEADERS))
    style_table_header(ws, 18 + 1, 3)
    for row in range(10, 16):
        ws.cell(row, 1).number_format = "yyyy-mm-dd"
        for col in range(7, 16):
            ws.cell(row, col).number_format = '#,##0.00;[Red]-#,##0.00;"-"'
    for row in range(10, 15):
        for col in (12, 13, 14, 15):
            if isinstance(ws.cell(row, col).value, str) and ws.cell(row, col).value.startswith("="):
                ws.cell(row, col).font = Font(color="008000")
    for row in range(20, 27):
        ws.cell(row, 2).number_format = '#,##0.00;[Red]-#,##0.00;"-"'
        ws.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 32 if row != 24 else 45
    for row in (21, 23, 24, 25, 26):
        ws.cell(row, 2).font = Font(color="008000")
    input_fill = PatternFill("solid", fgColor="DDEBF7")
    for row in range(1, 7):
        ws.cell(row, 1).font = Font(bold=True, color="1F4E78")
        ws.cell(row, 2).fill = input_fill
        ws.cell(row, 2).font = Font(color="0070C0")
    ws.conditional_formatting.add(
        "C22:C26", FormulaRule(formula=['$C22="REVIEW"'], fill=CHECK_FILL)
    )
    ws.freeze_panes = "A9"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = "A9:Q15"
    set_widths(ws, [32, 28, 58, 30, 30, 12, 12, 18, 18, 16, 19, 22, 16, 14, 20, 16, 16])
    ws.column_dimensions["S"].width = 25
    ws.column_dimensions["T"].width = 34
    ws.column_dimensions["V"].width = 34
    ws.column_dimensions["W"].width = 42


def apply_allocations_to_workbook_transactions(
    transactions: list[NormalizedTransaction],
    allocations_by_group: dict[str, list[dict]],
) -> None:
    for transaction in transactions:
        if not allocations_by_group.get(transaction.transaction_group_id):
            continue
        transaction.match_eligible = False
        transaction.expense_id = ""
        transaction.suggested_expense_id = ""
        transaction.match_status = "allocated_raw"
        transaction.review_note = (
            f"{transaction.review_note.rstrip()} Raw transaction retained; reimbursement is driven by allocation rows."
        ).strip()


def write_arvine_statement_sheet(
    ws,
    transactions: list[NormalizedTransaction],
    allocations_by_group: dict[str, list[dict]] | None = None,
    expenses: list[Expense] | None = None,
) -> None:
    allocations_by_group = allocations_by_group or {}
    expenses_by_file = {
        source_file_key(expense.source_file): expense for expense in (expenses or [])
    }
    ws.append(ARVINE_STATEMENT_HEADERS)
    previous_group = None
    group_fill = None
    for row, transaction in enumerate(transactions, start=2):
        if transaction.transaction_group_id != previous_group:
            group_fill = PatternFill("solid", fgColor="F3F6FA" if row % 2 == 0 else "FFFFFF")
        append_clean(
            ws,
            [
                transaction.transaction_group_id,
                transaction.funding_leg_id,
                excel_date(transaction.transaction_date),
                excel_date(transaction.posted_date),
                transaction.provider,
                transaction.account_label,
                transaction.cardholder,
                transaction.description,
                transaction.category,
                transaction.transaction_type,
                transaction.status,
                transaction.direction,
                transaction.match_eligible,
                transaction.purchase_amount,
                transaction.purchase_currency,
                transaction.settlement_amount,
                transaction.settlement_currency,
                transaction.cad_amount,
                transaction.cad_completeness,
                transaction.expense_id,
                transaction.suggested_expense_id,
                transaction.match_status,
                transaction.match_confidence,
                transaction.source_file.name,
                transaction.source_row,
                transaction.normalization_status,
                transaction.review_note,
                transaction.cad_conversion_rate,
                excel_date(transaction.cad_conversion_week_start),
                excel_date(transaction.cad_conversion_week_end),
                transaction.cad_conversion_method,
                transaction.cad_conversion_route,
                transaction.cad_conversion_source,
                ", ".join(transaction.cad_conversion_source_urls),
            ],
        )
        for col in range(1, len(ARVINE_STATEMENT_HEADERS) + 1):
            ws.cell(row, col).fill = group_fill
        if previous_group is not None and transaction.transaction_group_id != previous_group:
            border = Border(top=Side(style="thin", color="A6A6A6"))
            for col in range(1, len(ARVINE_STATEMENT_HEADERS) + 1):
                ws.cell(row, col).border = border
        source_cell = ws.cell(row, 24)
        source_cell.hyperlink = f"card_statements/{transaction.source_file.name}"
        source_cell.style = "Hyperlink"
        source_cell.font = Font(color="FF0000", underline="single")
        previous_group = transaction.transaction_group_id

    representatives: dict[str, NormalizedTransaction] = {}
    cad_by_group: dict[str, float] = {}
    for transaction in transactions:
        representatives.setdefault(transaction.transaction_group_id, transaction)
        if transaction.cad_amount is not None:
            cad_by_group[transaction.transaction_group_id] = round(
                cad_by_group.get(transaction.transaction_group_id, 0.0) + transaction.cad_amount,
                2,
            )
    for group_id, allocations in allocations_by_group.items():
        representative = representatives.get(group_id)
        if not representative:
            continue
        target_cad = cad_by_group.get(group_id)
        allocation_total = sum(float(item.get("cad_amount") or 0) for item in allocations)
        status = (
            "balanced"
            if target_cad is not None and abs(target_cad - allocation_total) <= 0.01
            else "review"
        )
        for allocation in allocations:
            expense = expenses_by_file.get(str(allocation.get("invoice_file") or ""))
            reimbursable = (
                allocation.get("type") in {"purchase", "refund", "fee"} and expense is not None
            )
            note_parts = [
                f"Allocation {allocation.get('allocation_id')} ({allocation.get('type')}).",
                str(allocation.get("note") or allocation.get("category") or "").strip(),
            ]
            append_clean(
                ws,
                [
                    group_id,
                    f"{group_id}:allocation:{allocation.get('allocation_id')}",
                    excel_date(representative.transaction_date),
                    excel_date(representative.posted_date),
                    "allocation",
                    representative.account_label,
                    representative.cardholder,
                    representative.description,
                    allocation.get("category"),
                    allocation.get("type"),
                    "allocated",
                    "in" if float(allocation.get("cad_amount") or 0) < 0 else "out",
                    reimbursable,
                    allocation.get("original_amount"),
                    representative.purchase_currency,
                    allocation.get("cad_amount"),
                    "CAD",
                    allocation.get("cad_amount"),
                    "complete",
                    expense.expense_id if reimbursable else None,
                    None,
                    "allocation",
                    float(allocation.get("percentage") or 0) / 100,
                    representative.source_file.name,
                    representative.source_row,
                    status,
                    " ".join(part for part in note_parts if part),
                    representative.cad_conversion_rate,
                    excel_date(representative.cad_conversion_week_start),
                    excel_date(representative.cad_conversion_week_end),
                    representative.cad_conversion_method,
                    representative.cad_conversion_route,
                    representative.cad_conversion_source,
                    ", ".join(representative.cad_conversion_source_urls),
                ],
            )
            allocation_row = ws.max_row
            source_cell = ws.cell(allocation_row, 24)
            source_cell.hyperlink = f"card_statements/{representative.source_file.name}"
            source_cell.style = "Hyperlink"
            source_cell.font = Font(color="FF0000", underline="single")

    end_row = max(ws.max_row, 2)
    set_filter_range(ws, len(ARVINE_STATEMENT_HEADERS), end_row)
    ws.freeze_panes = "H2"
    ws.sheet_view.showGridLines = False
    style_table_header(ws, 1, len(ARVINE_STATEMENT_HEADERS))
    for row in range(2, end_row + 1):
        for col in (20, 22):
            ws.cell(row, col).fill = PatternFill("solid", fgColor="DDEBF7")
            ws.cell(row, col).font = Font(color="0070C0")
            ws.cell(row, col).protection = Protection(locked=False)
        for col in (3, 4):
            ws.cell(row, col).number_format = "yyyy-mm-dd"
        for col in (14, 16, 18):
            ws.cell(row, col).number_format = '#,##0.00;[Red]-#,##0.00;"-"'
        ws.cell(row, 28).number_format = "0.00000000"
        for col in (29, 30):
            ws.cell(row, col).number_format = "yyyy-mm-dd"
        ws.cell(row, 23).number_format = "0%"
        ws.cell(row, 27).alignment = Alignment(vertical="top", wrap_text=True)
        for col in (33, 34):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        if any(ws.cell(row, col).value for col in (27, 33, 34)):
            ws.row_dimensions[row].height = 32
    if transactions or allocations_by_group:
        status_validation = DataValidation(
            type="list",
            formula1='"unmatched,suggested,auto,manual,matched,ignored,allocation,allocated_raw"',
        )
        ws.add_data_validation(status_validation)
        status_validation.add(f"V2:V{end_row}")
    ws.conditional_formatting.add(
        f"A2:AH{end_row}",
        FormulaRule(formula=['AND($M2=TRUE,$T2="")'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )
    ws.conditional_formatting.add(
        f"S2:S{end_row}",
        FormulaRule(
            formula=['OR($S2="partial",$S2="incomplete")'],
            fill=PatternFill("solid", fgColor="FCE4D6"),
        ),
    )
    ws.conditional_formatting.add(
        f"Z2:Z{end_row}",
        FormulaRule(formula=['$Z2<>"ok"'], fill=PatternFill("solid", fgColor="FFD966")),
    )
    ws.conditional_formatting.add(
        f"J2:J{end_row}",
        FormulaRule(formula=['$J2="refund"'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    set_widths(
        ws,
        [
            27,
            31,
            14,
            14,
            12,
            14,
            20,
            32,
            18,
            16,
            14,
            12,
            14,
            17,
            15,
            18,
            17,
            15,
            18,
            22,
            22,
            16,
            16,
            28,
            12,
            20,
            55,
            19,
            14,
            14,
            28,
            18,
            48,
            60,
        ],
    )
