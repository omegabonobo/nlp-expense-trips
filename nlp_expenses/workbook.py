from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from nlp_expenses.matching import (
    apply_manual_matches,
    match_normalized_transactions,
    match_transactions,
)
from nlp_expenses.models import Expense, NormalizedTransaction, StatementTransaction
from nlp_expenses.trips import source_file_key
from nlp_expenses.workbook_arvine import (
    apply_allocations_to_workbook_transactions,
    write_arvine_statement_sheet,
    write_arvine_summary_sheet,
)
from nlp_expenses.workbook_arvine_detail import write_arvine_detail_sheet
from nlp_expenses.workbook_common import (
    append_clean,
    configure_arvine_calculation,
    save_workbook_atomic,
    set_filter_range,
)
from nlp_expenses.workbook_reports import (
    build_ivado_claim_workbook as build_ivado_claim_workbook,
)
from nlp_expenses.workbook_reports import (
    build_reimbursement_report_workbook as build_reimbursement_report_workbook,
)

EXPENSE_HEADERS = [
    "expense_id",
    "date",
    "supplier_name",
    "expense_type",
    "amount_in_currency",
    "receipt_total_in_currency",
    "amount_check",
    "currency",
    "number_of_person",
    "corrected_amount_in_currency",
    "CAD_amount_statement",
    "CAD_rate",
    "corrected_CAD",
    "source_file",
    "extraction_status",
    "review_note",
    "include",
    "manual_CAD_override",
    "manual_CAD_note",
    "statement_purchase_amount",
    "statement_purchase_currency",
    "accounting_original_basis",
    "accounting_basis_status",
]
LINE_HEADERS = [
    "expense_id",
    "line_item_id",
    "date",
    "supplier_name",
    "description",
    "amount_in_currency",
    "currency",
    "is_alcohol",
    "included",
    "alcohol_detection_confidence",
    "alcohol_detection_reason",
    "alcohol_matched_term",
    "inclusion_overridden",
    "alcohol_overridden",
    "inclusion_note",
    "source_file",
    "confidence",
    "review_note",
]
STATEMENT_HEADERS = [
    "date",
    "description",
    "amount",
    "expense_id",
    "suggested_expense_id",
    "match_confidence",
    "source_file",
    "foreign_amount",
    "foreign_currency",
]

REVIEW_FILL = PatternFill("solid", fgColor="C00000")
REVIEW_FONT = Font(color="FFFFFF", bold=True)
CHECK_FILL = PatternFill("solid", fgColor="FFF2CC")


def build_workbook(
    trip_dir: Path,
    expenses: list[Expense],
    transactions: list[StatementTransaction],
    output_path: Path | None = None,
    pre_matched: bool = False,
) -> Path:
    if not pre_matched:
        match_transactions(expenses, transactions)
    wb = Workbook()
    expense_ws = wb.active
    expense_ws.title = "expense_list"
    line_ws = wb.create_sheet("expense_line_items")
    statement_ws = wb.create_sheet("card_statements")

    write_expense_sheet(expense_ws, expenses)
    write_line_sheet(line_ws, expenses)
    write_statement_sheet(statement_ws, transactions)

    for ws in [expense_ws, line_ws, statement_ws]:
        style_sheet(ws)
    apply_review_highlights(expense_ws, line_ws, statement_ws)

    out_path = output_path or trip_dir / f"expense_review_{trip_dir.name}.xlsx"
    return save_workbook_atomic(wb, out_path)


def build_arvine_workbook(
    trip_dir: Path,
    expenses: list[Expense],
    transactions: list[NormalizedTransaction],
    output_path: Path | None = None,
    manual_matches: dict[str, str | None] | None = None,
    accounting_profile: dict | None = None,
    transaction_allocations: dict[str, list[dict]] | None = None,
) -> Path:
    """Build Arvine's accounting workbook without changing the IVADO workbook path."""
    match_normalized_transactions(expenses, transactions)
    if manual_matches:
        apply_manual_matches(expenses, transactions, manual_matches)
    transaction_allocations = transaction_allocations or {}
    apply_allocations_to_workbook_transactions(transactions, transaction_allocations)
    workbook = Workbook()
    detail_ws = workbook.active
    detail_ws.title = "expense_detail"
    summary_ws = workbook.create_sheet("expense_summary")
    statements_ws = workbook.create_sheet("card_statements")
    line_ws = workbook.create_sheet("expense_line_items")

    if accounting_profile is None:
        from nlp_expenses.accounting import builtin_accounting_profile

        accounting_profile = builtin_accounting_profile()
    write_arvine_detail_sheet(detail_ws, expenses, accounting_profile)
    write_arvine_summary_sheet(summary_ws, trip_dir, accounting_profile)
    write_arvine_statement_sheet(statements_ws, transactions, transaction_allocations, expenses)
    write_line_sheet(line_ws, expenses)
    style_sheet(line_ws)
    configure_arvine_calculation(workbook)

    out_path = output_path or trip_dir / f"expense_review_{trip_dir.name}.xlsx"
    return save_workbook_atomic(workbook, out_path)


def write_expense_sheet(ws, expenses: list[Expense]) -> None:
    ws.append(EXPENSE_HEADERS)
    for idx, expense in enumerate(expenses, start=2):
        corrected_original_formula = (
            f'=IF($Q{idx}=FALSE,0,IF($F{idx}="","",'
            f"IF(COUNTIFS(expense_line_items!$A:$A,$A{idx},expense_line_items!$I:$I,FALSE,"
            f'expense_line_items!$F:$F,">0")>0,'
            f"SUMIFS(expense_line_items!$F:$F,expense_line_items!$A:$A,$A{idx},"
            f'expense_line_items!$I:$I,TRUE)/IF(OR($I{idx}="",$I{idx}=0),1,$I{idx}),'
            f'$F{idx}/IF(OR($I{idx}="",$I{idx}=0),1,$I{idx}))))'
        )
        line_ratio_formula = (
            f"IF(COUNTIFS(expense_line_items!$A:$A,$A{idx},expense_line_items!$I:$I,FALSE,"
            f'expense_line_items!$F:$F,">0")>0,'
            f'IF(OR($E{idx}="",$E{idx}=0),1,$J{idx}*MAX(1,$I{idx})/$E{idx}),1)'
        )
        append_clean(
            ws,
            [
                expense.expense_id,
                expense.date,
                expense.supplier_name,
                expense.expense_type,
                f'=IF(COUNTIFS(expense_line_items!$A:$A,$A{idx},expense_line_items!$F:$F,">0")=0,"",SUMIFS(expense_line_items!$F:$F,expense_line_items!$A:$A,$A{idx}))',
                expense.amount,
                (
                    f'=IF($F{idx}="","missing",IF($E{idx}="","receipt total used",'
                    f'IF(ABS($E{idx}-$F{idx})<=MAX(0.05,$F{idx}*0.03),"ok","mismatch")))'
                ),
                expense.currency,
                max(1, int(expense.number_of_people or 1)),
                corrected_original_formula,
                (
                    f'=IF($R{idx}<>"",$R{idx},IF(COUNTIF(card_statements!$D:$D,$A{idx})=0,"",'
                    f"SUMIFS(card_statements!$C:$C,card_statements!$D:$D,$A{idx})))"
                ),
                (
                    f'=IF(OR($K{idx}="",$K{idx}=0),"",'
                    f'IF($H{idx}="CAD",1,IF(OR($V{idx}="",$V{idx}=0),"",$K{idx}/$V{idx})))'
                ),
                (
                    f'=IF(OR($J{idx}="",$L{idx}=""),"",'
                    f'IF($W{idx}="statement_person_share",$K{idx}*{line_ratio_formula},'
                    f'IF(OR($W{idx}="statement_receipt_total",$W{idx}="statement_aggregated"),'
                    f"$K{idx}*{line_ratio_formula}/MAX(1,$I{idx}),$J{idx}*$L{idx})))"
                ),
                source_file_key(expense.source_file),
                extraction_status(expense),
                expense.review_note,
                expense.included,
                expense.manual_cad_override,
                expense.manual_cad_note,
                (
                    f'=IF($H{idx}="CAD",'
                    f'IF(COUNTIF(card_statements!$D:$D,$A{idx})=0,"",'
                    f"SUMIFS(card_statements!$C:$C,card_statements!$D:$D,$A{idx})),"
                    f'IF(COUNTIFS(card_statements!$D:$D,$A{idx},card_statements!$H:$H,">0")=0,"",'
                    f"SUMIFS(card_statements!$H:$H,card_statements!$D:$D,$A{idx})))"
                ),
                (
                    f'=IF($T{idx}="","",IF($H{idx}="CAD","CAD",'
                    f'IFERROR(INDEX(card_statements!$I:$I,MATCH($A{idx},card_statements!$D:$D,0)),"")))'
                ),
                (
                    f'=IF($R{idx}<>"",IF(AND($F{idx}<>"",$F{idx}<>0),$F{idx},$E{idx}),'
                    f'IF(OR($F{idx}="",$F{idx}=0),$E{idx},'
                    f'IF(OR($T{idx}="",$T{idx}=0),$F{idx},'
                    f"IF($U{idx}<>$H{idx},$F{idx},"
                    f"IF(OR(COUNTIF(card_statements!$D:$D,$A{idx})>1,"
                    f"ABS($T{idx}-$F{idx})<=MAX(2,ABS($F{idx})*0.08),"
                    f"AND($I{idx}>1,ABS($T{idx}-$F{idx}/$I{idx})"
                    f"<=MAX(2,ABS($F{idx}/$I{idx})*0.08))),$T{idx},$F{idx})))))"
                ),
                (
                    f'=IF($R{idx}<>"","manual_receipt_total",'
                    f'IF(OR($T{idx}="",$T{idx}=0),"receipt_total",'
                    f'IF($U{idx}<>$H{idx},"receipt_fallback_currency",'
                    f'IF(COUNTIF(card_statements!$D:$D,$A{idx})>1,"statement_aggregated",'
                    f'IF(ABS($T{idx}-$F{idx})<=MAX(2,ABS($F{idx})*0.08),"statement_receipt_total",'
                    f"IF(AND($I{idx}>1,ABS($T{idx}-$F{idx}/$I{idx})"
                    f"<=MAX(2,ABS($F{idx}/$I{idx})*0.08)),"
                    f'"statement_person_share","receipt_fallback_mismatch"))))))'
                ),
            ],
        )
    set_filter_range(ws, len(EXPENSE_HEADERS), max(len(expenses) + 1, 2))
    ws.freeze_panes = "A2"
    if expenses:
        validation = DataValidation(
            type="whole", operator="between", formula1="1", formula2="99", allow_blank=False
        )
        validation.error = "Enter a whole number from 1 to 99."
        validation.errorTitle = "Invalid number of persons"
        ws.add_data_validation(validation)
        validation.add(f"I2:I{len(expenses) + 1}")
        include_validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
        ws.add_data_validation(include_validation)
        include_validation.add(f"Q2:Q{len(expenses) + 1}")
        ws.conditional_formatting.add(
            f"K2:K{len(expenses) + 1}",
            CellIsRule(operator="equal", formula=["0"], fill=REVIEW_FILL, font=REVIEW_FONT),
        )
    for col in ["E", "F", "J", "K", "M", "T", "V"]:
        for cell in ws[col][1:]:
            if cell.value is not None:
                cell.number_format = "#,##0.00"
    for cell in ws["L"][1:]:
        if cell.value is not None:
            cell.number_format = "0.00000"


def write_line_sheet(ws, expenses: list[Expense]) -> None:
    ws.append(LINE_HEADERS)
    row_count = 0
    for expense in expenses:
        items = expense.line_items or []
        if not items:
            items = [fallback_line_item(expense)]
        for item_idx, item in enumerate(items, start=1):
            append_clean(
                ws,
                [
                    expense.expense_id,
                    f"{expense.expense_id}-L{item_idx:03d}",
                    expense.date,
                    expense.supplier_name,
                    item.description,
                    item.amount,
                    expense.currency,
                    item.is_alcohol,
                    item.included,
                    item.alcohol_confidence,
                    item.alcohol_reason,
                    item.alcohol_matched_term,
                    item.inclusion_overridden,
                    item.alcohol_overridden,
                    item.inclusion_note,
                    source_file_key(expense.source_file),
                    item.confidence,
                    item.review_note,
                ],
            )
            row_count += 1
    set_filter_range(ws, len(LINE_HEADERS), max(row_count + 1, 2))
    ws.freeze_panes = "A2"
    for cell in ws["F"][1:]:
        if cell.value is not None:
            cell.number_format = "#,##0.00"
    if row_count:
        validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
        validation.error = "Choose TRUE or FALSE."
        validation.errorTitle = "Invalid line-item choice"
        ws.add_data_validation(validation)
        validation.add(f"H2:I{row_count + 1}")


def write_statement_sheet(ws, transactions: list[StatementTransaction]) -> None:
    ws.append(STATEMENT_HEADERS)
    for transaction in transactions:
        append_clean(
            ws,
            [
                transaction.date,
                transaction.description,
                transaction.amount_cad,
                transaction.expense_id,
                transaction.suggested_expense_id,
                transaction.match_confidence or "",
                transaction.source_file.name,
                transaction.foreign_amount,
                transaction.foreign_currency,
            ],
        )
    set_filter_range(ws, len(STATEMENT_HEADERS), max(len(transactions) + 1, 2))
    ws.freeze_panes = "A2"
    for col in ["C", "H"]:
        for cell in ws[col][1:]:
            if cell.value is not None:
                cell.number_format = "#,##0.00"


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    widths = {
        "A": 28,
        "B": 12,
        "C": 30,
        "D": 18,
        "E": 16,
        "F": 20,
        "G": 14,
        "H": 12,
        "I": 16,
        "J": 20,
        "K": 20,
        "L": 14,
        "M": 16,
        "N": 32,
        "O": 12,
        "P": 48,
    }
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = widths.get(letter, 18)


def apply_review_highlights(expense_ws, line_ws, statement_ws) -> None:
    for row in range(2, expense_ws.max_row + 1):
        for col in (2, 3, 4, 5, 6, 8):
            cell = expense_ws.cell(row, col)
            if cell.value in (None, ""):
                mark_review(cell)
        amount_check = expense_ws.cell(row, 7)
        if amount_check.value != "ok":
            amount_check.fill = CHECK_FILL
        status = expense_ws.cell(row, 15)
        if status.value in {"review", "manual"}:
            mark_review(status)
        review_note = expense_ws.cell(row, 16)
        if review_note.value:
            mark_review(review_note)
        expense_ws.cell(row, 9).fill = CHECK_FILL

    for row in range(2, statement_ws.max_row + 1):
        expense_id = statement_ws.cell(row, 4)
        if expense_id.value in (None, ""):
            mark_review(expense_id)

    for row in range(2, line_ws.max_row + 1):
        for col in (5, 6):
            cell = line_ws.cell(row, col)
            if cell.value in (None, ""):
                mark_review(cell)
        review_note = line_ws.cell(row, LINE_HEADERS.index("review_note") + 1)
        if review_note.value:
            review_note.fill = CHECK_FILL


def mark_review(cell) -> None:
    cell.fill = REVIEW_FILL
    cell.font = REVIEW_FONT


def fallback_line_item(expense: Expense):
    class FallbackLineItem:
        description = (
            "Receipt total" if expense.amount is not None else "Receipt total missing - review"
        )
        amount = expense.amount
        is_alcohol = False
        included = True
        alcohol_confidence = 0.0
        alcohol_reason = "no line-item classification available"
        alcohol_matched_term = ""
        inclusion_overridden = False
        alcohol_overridden = False
        inclusion_note = ""
        confidence = 0.0
        review_note = "Generated workbook fallback line item."

    return FallbackLineItem()


def extraction_status(expense: Expense) -> str:
    if expense.amount is None or "No extractable text" in expense.review_note:
        return "manual"
    if not expense.date or not expense.supplier_name or not expense.expense_type:
        return "review"
    if expense.review_note:
        return "review"
    return "ok"
