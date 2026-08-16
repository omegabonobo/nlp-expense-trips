from __future__ import annotations

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from nlp_expenses.models import Expense
from nlp_expenses.trips import source_file_key
from nlp_expenses.workbook_common import (
    append_clean,
    arvine_extraction_status,
    excel_date,
    set_filter_range,
    set_widths,
    style_table_header,
)

ARVINE_DETAIL_HEADERS = [
    "include",
    "expense_id",
    "invoice_date",
    "vendor",
    "description",
    "expense_type",
    "country",
    "province",
    "original_total",
    "currency",
    "GST/HST_in_currency",
    "QST_in_currency",
    "subtotal_excl_GST_QST_in_currency",
    "GST/HST_number",
    "QST_number",
    "tax_deductible_pct",
    "tax_credit_pct",
    "CAD_amount_statement",
    "CAD_statement_status",
    "manual_CAD_override",
    "CAD_amount_used",
    "FX_rate",
    "GST/HST_CAD",
    "QST_CAD",
    "recoverable_GST/HST_CAD",
    "recoverable_QST_CAD",
    "net_expense_CAD",
    "deductible_CAD",
    "non_deductible_CAD",
    "trip_purpose",
    "meal_attendees_client",
    "source_file",
    "extraction_status",
    "statement_match_status",
    "tax_documentation_status",
    "review_note",
    "line_item_review_status",
    "line_item_total",
    "included_line_total",
    "excluded_line_total",
    "claimable_pct",
    "number_of_people",
    "statement_purchase_amount",
    "statement_purchase_currency",
    "accounting_original_basis",
    "accounting_basis_status",
]


def expense_line_review_values(expense: Expense) -> dict:
    amounts = [item.amount for item in expense.line_items if item.amount is not None]
    included = [
        item.amount for item in expense.line_items if item.included and item.amount is not None
    ]
    line_total = (
        expense.line_item_total
        if expense.line_item_total is not None
        else round(sum(amounts), 2)
        if amounts
        else None
    )
    included_total = (
        expense.included_line_total
        if expense.included_line_total is not None
        else round(sum(included), 2)
        if amounts
        else None
    )
    excluded_total = (
        expense.excluded_line_total
        if expense.excluded_line_total is not None
        else round((line_total or 0) - (included_total or 0), 2)
        if amounts
        else None
    )
    if expense.line_item_review_status:
        ratio = expense.claimable_ratio
        status = expense.line_item_review_status
    elif line_total and excluded_total and excluded_total > 0:
        ratio = max(0.0, min(1.0, (included_total or 0) / line_total))
        status = "automatic"
    else:
        ratio = 1.0
        status = "automatic" if amounts else "not_available"
    return {
        "status": status,
        "line_total": line_total,
        "included_total": included_total,
        "excluded_total": excluded_total,
        "claimable_ratio": round(ratio, 8),
    }


def write_arvine_detail_sheet(ws, expenses: list[Expense], accounting_profile: dict) -> None:
    ws.append(ARVINE_DETAIL_HEADERS)
    statement_end = max(5000, len(expenses) * 20 + 100)
    claimable_col = get_column_letter(ARVINE_DETAIL_HEADERS.index("claimable_pct") + 1)
    for row, expense in enumerate(expenses, start=2):
        line_review = expense_line_review_values(expense)
        is_meal = str(expense.expense_type or "").startswith("meal")
        claimable_formula = (
            f"=IF($A{row}=FALSE,0,IF(COUNTIFS(expense_line_items!$A:$A,$B{row},"
            f'expense_line_items!$I:$I,FALSE,expense_line_items!$F:$F,">0")>0,'
            f'IF(OR($AL{row}="",$AL{row}=0),1/MAX(1,$AP{row}),'
            f"MAX(0,MIN(1,$AM{row}/$AL{row}/MAX(1,$AP{row})))),"
            f"1/MAX(1,$AP{row})))"
        )
        claimable_original_formula = (
            f"IF(COUNTIFS(expense_line_items!$A:$A,$B{row},"
            f'expense_line_items!$I:$I,FALSE,expense_line_items!$F:$F,">0")>0,'
            f"$AM{row}/MAX(1,$AP{row}),$I{row}/MAX(1,$AP{row}))"
        )
        commercial_use = accounting_profile["commercial_use_pct"]
        deductible_pct = (
            accounting_profile["meal_deduction_pct"]
            if is_meal
            else accounting_profile["normal_deduction_pct"]
        ) * commercial_use
        base_tax_recovery = (
            accounting_profile["meal_tax_recovery_pct"]
            if is_meal
            else accounting_profile["normal_tax_recovery_pct"]
        ) * commercial_use
        gst_recovery_pct = base_tax_recovery if accounting_profile["gst_hst_registrant"] else 0.0
        qst_recovery_pct = base_tax_recovery if accounting_profile["qst_registrant"] else 0.0
        tax_credit_pct = max(gst_recovery_pct, qst_recovery_pct)
        append_clean(
            ws,
            [
                expense.included,
                expense.expense_id,
                excel_date(expense.date),
                expense.supplier_name,
                expense.description,
                expense.expense_type,
                expense.country,
                expense.province,
                expense.amount,
                expense.currency,
                expense.gst_hst,
                expense.qst,
                f'=IF($I{row}="","",IF(AND($K{row}="",$L{row}=""),$I{row},$I{row}-N($K{row})-N($L{row})))',
                expense.gst_hst_number,
                expense.qst_number,
                deductible_pct,
                tax_credit_pct,
                (
                    f"=IF(COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},'card_statements'!$M$2:$M${statement_end},TRUE)=0,\"\","
                    f"IF(COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},'card_statements'!$M$2:$M${statement_end},TRUE,"
                    f'\'card_statements\'!$S$2:$S${statement_end},"<>complete")>0,"",'
                    f"SUMIFS('card_statements'!$R$2:$R${statement_end},'card_statements'!$T$2:$T${statement_end},$B{row},"
                    f"'card_statements'!$M$2:$M${statement_end},TRUE)))"
                ),
                (
                    f'=IF($T{row}<>"","manual",'
                    f"IF(COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},'card_statements'!$M$2:$M${statement_end},TRUE)=0,\"missing\","
                    f"IF(COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},'card_statements'!$M$2:$M${statement_end},TRUE,"
                    f'\'card_statements\'!$S$2:$S${statement_end},"<>complete")>0,"incomplete","complete")))'
                ),
                expense.manual_cad_override,
                (
                    f'=IF($T{row}<>"",ROUND($T{row}*${claimable_col}{row},2),'
                    f'IF(AND($R{row}<>"",$AS{row}<>""),'
                    f'ROUND(IF(OR($AT{row}="statement_person_share",$AT{row}="statement_person_share_includes_tip"),'
                    f"$R{row}*${claimable_col}{row}*MAX(1,$AP{row}),"
                    f'IF(OR($AT{row}="statement_receipt_total",$AT{row}="statement_includes_tip",$AT{row}="statement_aggregated"),'
                    f"$R{row}*${claimable_col}{row},$R{row}/$AS{row}*{claimable_original_formula})),2),"
                    f'IF($J{row}="CAD",ROUND($I{row}*${claimable_col}{row},2),"")))'
                ),
                (
                    f'=IF($I{row}="","",IF($J{row}="CAD",1,'
                    f'IFERROR(IF($T{row}<>"",$T{row}/$I{row},$R{row}/$AS{row}),"")))'
                ),
                f'=IF(OR($K{row}="",$V{row}=""),0,ROUND($K{row}*$V{row}*${claimable_col}{row},2))',
                f'=IF(OR($L{row}="",$V{row}=""),0,ROUND($L{row}*$V{row}*${claimable_col}{row},2))',
                f'=IF($U{row}="","",ROUND($W{row}*{gst_recovery_pct:.6f},2))',
                f'=IF($U{row}="","",ROUND($X{row}*{qst_recovery_pct:.6f},2))',
                f'=IF($U{row}="","",ROUND($U{row}-N($Y{row})-N($Z{row}),2))',
                f'=IF($AA{row}="","",IF($F{row}="meal",ROUND($AA{row}*$P{row},2),$AA{row}))',
                f'=IF($AA{row}="","",IF($F{row}="meal",$AA{row}-$AB{row},0))',
                expense.business_purpose,
                expense.attendees_client,
                source_file_key(expense.source_file),
                arvine_extraction_status(expense),
                (
                    f"=IF(COUNTIF('card_statements'!$T$2:$T${statement_end},$B{row})=0,\"unmatched\","
                    f"IF(COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},'card_statements'!$V$2:$V${statement_end},\"auto\")+"
                    f"COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},'card_statements'!$V$2:$V${statement_end},\"manual\")+"
                    f"COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},'card_statements'!$V$2:$V${statement_end},\"allocation\")+"
                    f"COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},'card_statements'!$V$2:$V${statement_end},\"matched\")"
                    f'>0,"matched","review"))'
                ),
                expense.tax_documentation_status or "review",
                (
                    f"{expense.review_note.rstrip()} Manual CAD override source: {expense.manual_cad_note}".strip()
                    if expense.manual_cad_override is not None
                    else expense.review_note
                ),
                line_review["status"],
                (
                    f'=IF(COUNTIF(expense_line_items!$A:$A,$B{row})=0,"",'
                    f"SUMIFS(expense_line_items!$F:$F,expense_line_items!$A:$A,$B{row}))"
                ),
                (
                    f'=IF($AL{row}="","",SUMIFS(expense_line_items!$F:$F,'
                    f"expense_line_items!$A:$A,$B{row},expense_line_items!$I:$I,TRUE))"
                ),
                f'=IF($AL{row}="","",$AL{row}-$AM{row})',
                claimable_formula,
                max(1, int(expense.number_of_people or 1)),
                (
                    f"=IF(COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},"
                    f"'card_statements'!$M$2:$M${statement_end},TRUE,"
                    f'\'card_statements\'!$N$2:$N${statement_end},"<>")=0,"",'
                    f"SUMIFS('card_statements'!$N$2:$N${statement_end},"
                    f"'card_statements'!$T$2:$T${statement_end},$B{row},"
                    f"'card_statements'!$M$2:$M${statement_end},TRUE))"
                ),
                (
                    f'=IF($AQ{row}="","",IFERROR(INDEX(\'card_statements\'!$O$2:$O${statement_end},'
                    f"MATCH($B{row},'card_statements'!$T$2:$T${statement_end},0)),\"\"))"
                ),
                (
                    f'=IF(OR($I{row}="",$I{row}=0),"",IF($T{row}<>"",$I{row},'
                    f'IF(AND($AQ{row}<>"",$AQ{row}<>0,$AR{row}=$J{row},'
                    f"OR(COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},"
                    f"'card_statements'!$M$2:$M${statement_end},TRUE)>1,"
                    f"OR(ABS($AQ{row}-$I{row})<=MAX(2,ABS($I{row})*0.08),"
                    f"AND($AQ{row}>$I{row}+MAX(2,ABS($I{row})*0.08),$AQ{row}<=$I{row}*1.35),"
                    f"AND($AP{row}>1,ABS($AQ{row}-$I{row}/$AP{row})"
                    f"<=MAX(2,ABS($I{row}/$AP{row})*0.08)),"
                    f"AND($AP{row}>1,$AQ{row}>$I{row}/$AP{row}+MAX(2,ABS($I{row}/$AP{row})*0.08),"
                    f"$AQ{row}<=$I{row}/$AP{row}*1.35)))),$AQ{row},$I{row})))"
                ),
                (
                    f'=IF(OR($I{row}="",$I{row}=0),"unavailable",IF($T{row}<>"","manual_receipt_total",'
                    f'IF(OR($AQ{row}="",$AQ{row}=0),"receipt_total",'
                    f'IF($AR{row}<>$J{row},"receipt_fallback_currency",'
                    f"IF(COUNTIFS('card_statements'!$T$2:$T${statement_end},$B{row},"
                    f"'card_statements'!$M$2:$M${statement_end},TRUE)>1,\"statement_aggregated\","
                    f'IF(ABS($AQ{row}-$I{row})<=MAX(2,ABS($I{row})*0.08),"statement_receipt_total",'
                    f'IF(AND($AQ{row}>$I{row}+MAX(2,ABS($I{row})*0.08),$AQ{row}<=$I{row}*1.35),"statement_includes_tip",'
                    f"IF(AND($AP{row}>1,ABS($AQ{row}-$I{row}/$AP{row})"
                    f"<=MAX(2,ABS($I{row}/$AP{row})*0.08)),"
                    f'"statement_person_share",IF(AND($AP{row}>1,$AQ{row}>$I{row}/$AP{row}+MAX(2,ABS($I{row}/$AP{row})*0.08),'
                    f'$AQ{row}<=$I{row}/$AP{row}*1.35),"statement_person_share_includes_tip","receipt_fallback_mismatch"))))))))'
                ),
            ],
        )
        source_cell = ws.cell(row, 32)
        source_cell.hyperlink = f"expenses_receipts/{source_file_key(expense.source_file)}"
        source_cell.style = "Hyperlink"
        source_cell.font = Font(color="FF0000", underline="single")

    detail_end = max(len(expenses) + 1, 2)
    set_filter_range(ws, len(ARVINE_DETAIL_HEADERS), detail_end)
    ws.freeze_panes = "D2"
    ws.sheet_view.showGridLines = False
    style_table_header(ws, 1, len(ARVINE_DETAIL_HEADERS))
    add_arvine_detail_validations(ws, detail_end)
    style_arvine_detail_rows(ws, detail_end)
    add_arvine_detail_highlights(ws, detail_end)
    widths = [
        10,
        20,
        13,
        24,
        24,
        14,
        15,
        10,
        15,
        10,
        16,
        14,
        20,
        19,
        18,
        17,
        15,
        18,
        19,
        20,
        17,
        12,
        14,
        13,
        20,
        17,
        17,
        16,
        19,
        25,
        25,
        26,
        18,
        21,
        22,
        48,
        20,
        16,
        18,
        18,
        14,
        18,
        20,
        18,
        20,
        26,
    ]
    set_widths(ws, widths)


def add_arvine_detail_validations(ws, end_row: int) -> None:
    if end_row < 2:
        return
    boolean_validation = DataValidation(type="list", formula1='"TRUE,FALSE"')
    type_validation = DataValidation(type="list", formula1='"flight,hotel,transport,meal,other"')
    currency_validation = DataValidation(
        type="list", formula1='"CAD,USD,AUD,EUR,GBP,IDR,VND,QAR,HKD,CHF"'
    )
    pct_validation = DataValidation(type="decimal", operator="between", formula1="0", formula2="1")
    people_validation = DataValidation(
        type="whole", operator="between", formula1="1", formula2="99"
    )
    for validation, target in [
        (boolean_validation, f"A2:A{end_row}"),
        (type_validation, f"F2:F{end_row}"),
        (currency_validation, f"J2:J{end_row}"),
        (pct_validation, f"P2:Q{end_row}"),
        (people_validation, f"AP2:AP{end_row}"),
    ]:
        ws.add_data_validation(validation)
        validation.add(target)


def style_arvine_detail_rows(ws, end_row: int) -> None:
    input_columns = {1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 14, 15, 16, 17, 20, 30, 31, 35, 36, 42}
    formula_columns = {
        13,
        18,
        19,
        21,
        22,
        23,
        24,
        25,
        26,
        27,
        28,
        29,
        34,
        38,
        39,
        40,
        41,
        43,
        44,
        45,
        46,
    }
    for row in range(2, end_row + 1):
        for col in input_columns:
            ws.cell(row, col).fill = PatternFill("solid", fgColor="DDEBF7")
            ws.cell(row, col).font = Font(color="0070C0")
        for col in formula_columns:
            ws.cell(row, col).font = (
                Font(color="008000")
                if col in {18, 19, 34, 38, 39, 40, 41}
                else Font(color="000000")
            )
        ws.cell(row, 3).number_format = "yyyy-mm-dd"
        for col in (9, 11, 12, 13, 18, 20, 21, 23, 24, 25, 26, 27, 28, 29, 43, 45):
            ws.cell(row, col).number_format = '#,##0.00;[Red]-#,##0.00;"-"'
        for col in (38, 39, 40):
            ws.cell(row, col).number_format = '#,##0.00;[Red]-#,##0.00;"-"'
        for col in (16, 17):
            ws.cell(row, col).number_format = "0%"
        ws.cell(row, 41).number_format = "0.00%"
        ws.cell(row, 22).number_format = "0.000000"
        for col in range(1, len(ARVINE_DETAIL_HEADERS) + 1):
            ws.cell(row, col).alignment = Alignment(
                vertical="top", wrap_text=col in {5, 30, 31, 36}
            )


def add_arvine_detail_highlights(ws, end_row: int) -> None:
    ws.conditional_formatting.add(
        f"S2:S{end_row}",
        FormulaRule(
            formula=['OR($S2="missing",$S2="incomplete")'],
            fill=PatternFill("solid", fgColor="FCE4D6"),
        ),
    )
    ws.conditional_formatting.add(
        f"AH2:AH{end_row}",
        FormulaRule(formula=['$AH2<>"matched"'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )
    ws.conditional_formatting.add(
        f"AI2:AI{end_row}",
        FormulaRule(formula=['$AI2="review"'], fill=PatternFill("solid", fgColor="FCE4D6")),
    )
    ws.conditional_formatting.add(
        f"AG2:AG{end_row}",
        FormulaRule(formula=['$AG2<>"ok"'], fill=PatternFill("solid", fgColor="FCE4D6")),
    )
    ws.conditional_formatting.add(
        f"AK2:AK{end_row}",
        FormulaRule(
            formula=['OR($AK2="review",$AK2="not_available")'],
            fill=PatternFill("solid", fgColor="FFF2CC"),
        ),
    )
    ws.conditional_formatting.add(
        f"AT2:AT{end_row}",
        FormulaRule(
            formula=['OR($AT2="receipt_fallback_mismatch",$AT2="receipt_fallback_currency")'],
            fill=PatternFill("solid", fgColor="FFF2CC"),
        ),
    )
