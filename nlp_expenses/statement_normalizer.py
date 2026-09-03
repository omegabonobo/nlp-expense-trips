from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import defaultdict
from collections.abc import Callable
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from nlp_expenses.fx_rates import FxRateUnavailable, WeeklyCadFxResolver
from nlp_expenses.models import NormalizationResult, NormalizedTransaction, StatementFileReport
from nlp_expenses.storage import write_json_atomic

SUPPORTED_STATEMENT_EXTENSIONS = {".csv", ".xls", ".xlsx"}
STATEMENT_SETTINGS_FILE = ".nlp-expenses-statement-settings.json"
DATE_CONVENTIONS = {"day_first", "month_first"}
SIGN_CONVENTIONS = {"positive_purchase", "negative_purchase", "debit_credit"}
PROVIDER_DATE_CONVENTIONS = {
    "standard": "iso",
    "amex": "month_first",
    "bmo": "month_first",
    "bnc": "day_first",
    "wise": "iso",
}
STANDARD_TRANSACTION_TYPES = {
    "purchase",
    "refund",
    "payment",
    "cashback",
    "cash",
    "fee",
    "transfer",
    "deposit",
}
STANDARD_INCOMING_TYPES = {"refund", "cashback", "deposit"}

GENERIC_FIELD_ALIASES = {
    "transaction_date": {
        "date",
        "transaction_date",
        "transaction_day",
        "purchase_date",
        "activity_date",
        "booking_date",
        "booked_date",
        "trans_date",
    },
    "posted_date": {
        "posted_date",
        "date_posted",
        "date_processed",
        "processing_date",
        "value_date",
        "posting_date",
        "posted_on",
    },
    "description": {
        "description",
        "merchant",
        "merchant_name",
        "merchant_details",
        "details",
        "transaction",
        "transaction_description",
        "card_transaction",
        "name",
        "memo",
        "narrative",
        "payee",
    },
    "amount": {
        "amount",
        "transaction_amount",
        "transaction_value",
        "cad_amount",
        "charges_adjustments",
        "net_amount",
    },
    "debit": {
        "debit",
        "debit_amount",
        "charge",
        "charges",
        "withdrawal",
        "withdrawal_amount",
        "money_out",
    },
    "credit": {
        "credit",
        "credit_amount",
        "deposit_amount",
        "money_in",
    },
    "purchase_amount": {
        "purchase_amount",
        "original_amount",
        "foreign_amount",
        "local_amount",
        "transaction_currency_amount",
    },
    "purchase_currency": {
        "purchase_currency",
        "original_currency",
        "foreign_currency",
        "local_currency",
        "transaction_currency",
    },
    "settlement_amount": {
        "settlement_amount",
        "billing_amount",
        "billed_amount",
        "card_amount",
        "statement_amount",
        "home_amount",
    },
    "settlement_currency": {
        "settlement_currency",
        "billing_currency",
        "billed_currency",
        "card_currency",
        "statement_currency",
        "home_currency",
        "currency",
    },
    "account": {
        "account",
        "account_number",
        "card",
        "card_number",
        "card_no",
        "last_four",
    },
    "cardholder": {"cardholder", "card_member", "cardholder_name", "employee"},
    "transaction_type": {"transaction_type", "type", "debit_credit", "credit_debit"},
    "category": {"category", "merchant_category", "expense_category"},
}

GENERIC_MAPPING_FIELDS = tuple(GENERIC_FIELD_ALIASES)


class StatementNormalizationError(ValueError):
    pass


def list_statement_files(statement_dir: Path) -> list[Path]:
    if not statement_dir.exists():
        return []
    return sorted(
        path for path in statement_dir.iterdir() if path.is_file() and not path.name.startswith(".")
    )


def prepare_statement_file(path: Path) -> dict:
    rows, metadata = read_tabular_rows_with_metadata(path)
    settings = load_statement_settings(path.parent.parent)
    file_override = settings.get("file_import_overrides", {}).get(
        statement_file_fingerprint(path), {}
    )
    override_header = file_override.get("header_index") if isinstance(file_override, dict) else None
    if isinstance(override_header, int) and 0 <= override_header < len(rows):
        provider, header_index = "generic", override_header
    else:
        provider, header_index = detect_provider(rows)
    raw_headers = list(rows[header_index]) if header_index < len(rows) else []
    option_width = max((len(row) for row in rows[:60]), default=len(raw_headers))
    options = [
        {
            "index": index,
            "label": clean_text(raw_headers[index])
            if index < len(raw_headers) and clean_text(raw_headers[index])
            else f"Column {index + 1}",
        }
        for index in range(option_width)
    ]
    plan = {
        "rows": rows,
        "provider": provider,
        "header_index": header_index,
        "encoding": metadata.get("encoding", ""),
        "delimiter": metadata.get("delimiter", ""),
        "mapping": {},
        "mapping_options": options,
        "mapping_confidence": 1.0,
        "mapping_required": False,
        "sign_convention": "",
        "sign_convention_required": False,
        "profile_reused": False,
        "structure_fingerprint": statement_structure_fingerprint(raw_headers),
        "saved_profile": {},
        "errors": [],
    }
    if provider != "generic":
        plan["rows_read"] = count_data_rows(rows, header_index)
        return plan

    profiles = settings.get("import_profiles", {})
    saved_profile = profiles.get(plan["structure_fingerprint"], {})
    if not isinstance(saved_profile, dict):
        saved_profile = {}
    inferred, confidence, ambiguous = infer_generic_column_mapping(rows, header_index)
    saved_mapping = normalize_saved_mapping(saved_profile.get("mapping"), len(raw_headers))
    mapping = saved_mapping or inferred
    plan["mapping"] = mapping
    plan["mapping_confidence"] = 1.0 if saved_mapping else confidence
    plan["profile_reused"] = bool(saved_mapping)
    plan["saved_profile"] = saved_profile

    missing = missing_generic_mapping_fields(mapping)
    if missing:
        plan["mapping_required"] = True
        plan["errors"].append("statement columns need mapping for " + ", ".join(missing) + ".")
    blocking_ambiguous = ambiguous & {"description"}
    if not ({"transaction_date", "posted_date"} & set(mapping)):
        blocking_ambiguous.update(ambiguous & {"transaction_date", "posted_date"})
    if "amount" in ambiguous and not any(
        field in mapping for field in ("debit", "credit", "purchase_amount", "settlement_amount")
    ):
        blocking_ambiguous.add("amount")
    if blocking_ambiguous and not saved_mapping:
        plan["mapping_required"] = True
        plan["errors"].append(
            "multiple columns could represent " + ", ".join(sorted(blocking_ambiguous)) + "."
        )

    saved_sign = clean_text(saved_profile.get("sign_convention"))
    sign, sign_required = infer_sign_convention(raw_headers, mapping, saved_sign)
    plan["sign_convention"] = sign
    plan["sign_convention_required"] = sign_required
    if sign_required:
        plan["errors"].append(
            "purchase signs are ambiguous. Choose whether positive or negative values are purchases."
        )

    plan["rows"] = apply_generic_column_mapping(rows, header_index, mapping)
    plan["rows_read"] = count_data_rows(plan["rows"], header_index)
    return plan


def apply_import_plan_to_report(report: StatementFileReport, plan: dict) -> None:
    report.provider = str(plan.get("provider", ""))
    report.encoding = str(plan.get("encoding", ""))
    report.delimiter = str(plan.get("delimiter", ""))
    report.header_row = int(plan.get("header_index", 0)) + 1
    report.mapping_confidence = float(plan.get("mapping_confidence", 0.0))
    report.column_mapping = dict(plan.get("mapping", {}))
    options = list(plan.get("mapping_options", []))
    labels = {int(item["index"]): str(item["label"]) for item in options}
    report.column_labels = {
        field: labels.get(index, f"Column {index + 1}")
        for field, index in report.column_mapping.items()
    }
    report.mapping_options = options
    report.mapping_required = bool(plan.get("mapping_required"))
    report.sign_convention = str(plan.get("sign_convention", ""))
    report.sign_convention_required = bool(plan.get("sign_convention_required"))
    report.profile_reused = bool(plan.get("profile_reused"))
    report.structure_fingerprint = str(plan.get("structure_fingerprint", ""))
    report.rows_read = int(plan.get("rows_read", 0))


def count_data_rows(rows: list[list[object]], header_index: int) -> int:
    return sum(1 for row in rows[header_index + 1 :] if any(clean_text(value) for value in row))


def statement_structure_fingerprint(headers: list[object]) -> str:
    normalized = [normalize_key(value) for value in headers]
    payload = json.dumps(normalized, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:20]


def normalize_saved_mapping(value: object, width: int) -> dict[str, int]:
    if not isinstance(value, dict):
        return {}
    mapping: dict[str, int] = {}
    for field, index in value.items():
        if field not in GENERIC_MAPPING_FIELDS or isinstance(index, bool):
            continue
        try:
            selected = int(index)
        except (TypeError, ValueError):
            continue
        if 0 <= selected < width:
            mapping[field] = selected
    return mapping


def infer_generic_column_mapping(
    rows: list[list[object]], header_index: int
) -> tuple[dict[str, int], float, set[str]]:
    headers = [normalize_key(value) for value in rows[header_index]]
    mapping: dict[str, int] = {}
    ambiguous: set[str] = set()
    scores: list[float] = []
    for field, aliases in GENERIC_FIELD_ALIASES.items():
        candidates: list[tuple[float, int]] = []
        for index, header in enumerate(headers):
            if not header:
                continue
            if header in aliases:
                score = 1.0
            else:
                matches = [alias for alias in aliases if len(alias) >= 5 and alias in header]
                score = 0.72 if matches else 0.0
            if score:
                score += column_value_score(field, rows, header_index, index)
                candidates.append((score, index))
        if not candidates:
            continue
        candidates.sort(reverse=True)
        best_score, best_index = candidates[0]
        tied = [index for score, index in candidates if abs(score - best_score) < 0.01]
        if len(tied) > 1:
            ambiguous.add(field)
            continue
        mapping[field] = best_index
        scores.append(min(best_score, 1.0))
    required_detected = [field for field in ("transaction_date", "description") if field in mapping]
    amount_detected = any(
        field in mapping
        for field in ("amount", "debit", "credit", "purchase_amount", "settlement_amount")
    )
    confidence_parts = scores + [len(required_detected) / 2, float(amount_detected)]
    confidence = sum(confidence_parts) / len(confidence_parts) if confidence_parts else 0.0
    return mapping, round(confidence, 2), ambiguous


def column_value_score(
    field: str, rows: list[list[object]], header_index: int, column_index: int
) -> float:
    values = [
        row[column_index]
        for row in rows[header_index + 1 : header_index + 16]
        if column_index < len(row) and clean_text(row[column_index])
    ]
    if not values:
        return 0.0
    if field in {"transaction_date", "posted_date"}:
        matches = sum(
            bool(parse_date(value, "day_first") or parse_date(value, "month_first"))
            for value in values
        )
    elif field in {"amount", "debit", "credit", "purchase_amount", "settlement_amount"}:
        matches = sum(parse_number(value) is not None for value in values)
    elif field in {"purchase_currency", "settlement_currency"}:
        matches = sum(currency_code(value) is not None for value in values)
    else:
        matches = sum(bool(clean_text(value)) for value in values)
    return 0.18 * matches / len(values)


def missing_generic_mapping_fields(mapping: dict[str, int]) -> list[str]:
    missing = []
    if "transaction_date" not in mapping and "posted_date" not in mapping:
        missing.append("transaction date")
    if "description" not in mapping:
        missing.append("description")
    if not any(
        field in mapping
        for field in ("amount", "debit", "credit", "purchase_amount", "settlement_amount")
    ):
        missing.append("amount")
    return missing


def infer_sign_convention(
    headers: list[object], mapping: dict[str, int], saved: str = ""
) -> tuple[str, bool]:
    if saved in SIGN_CONVENTIONS:
        return saved, False
    if "debit" in mapping or "credit" in mapping:
        return "debit_credit", False
    amount_fields = [
        field for field in ("amount", "purchase_amount", "settlement_amount") if field in mapping
    ]
    if not amount_fields:
        return "", False
    labels = {field: normalize_key(headers[mapping[field]]) for field in amount_fields}
    explicit_positive = {alias for field in amount_fields for alias in GENERIC_FIELD_ALIASES[field]}
    if all(label in explicit_positive for label in labels.values()):
        return "positive_purchase", False
    return "", True


def apply_generic_column_mapping(
    rows: list[list[object]], header_index: int, mapping: dict[str, int]
) -> list[list[object]]:
    transformed = [list(row) for row in rows]
    headers = [normalize_key(value) for value in transformed[header_index]]
    for field, index in mapping.items():
        while len(headers) <= index:
            headers.append("")
        headers[index] = field
    transformed[header_index] = headers
    return transformed


def transaction_preview(
    transactions: list[NormalizedTransaction], limit: int = 3
) -> list[dict[str, object]]:
    return [
        {
            "source_row": item.source_row,
            "date": item.transaction_date,
            "description": item.description,
            "purchase_amount": item.purchase_amount,
            "purchase_currency": item.purchase_currency,
            "settlement_amount": item.settlement_amount,
            "settlement_currency": item.settlement_currency,
            "transaction_type": item.transaction_type,
        }
        for item in transactions[:limit]
    ]


def preflight_statement_files(paths: list[Path]) -> list[StatementFileReport]:
    reports: list[StatementFileReport] = []
    for path in paths:
        report = StatementFileReport(source_file=path)
        if path.suffix.lower() not in SUPPORTED_STATEMENT_EXTENSIONS:
            report.errors.append(
                f"{path.name}: unsupported statement file extension {path.suffix or '(none)'}. Use CSV, XLS, or XLSX."
            )
            reports.append(report)
            continue
        try:
            plan = prepare_statement_file(path)
            apply_import_plan_to_report(report, plan)
            if plan["errors"]:
                report.errors.extend(f"{path.name}: {message}" for message in plan["errors"])
                reports.append(report)
                continue
            rows = plan["rows"]
            provider = plan["provider"]
            header_index = plan["header_index"]
            date_profile = statement_date_profile(path, provider, rows, header_index, plan)
            apply_date_profile_to_report(report, date_profile)
            if report.date_convention_required:
                report.errors.append(
                    f"{path.name}: dates are ambiguous. Choose day-first or month-first before syncing."
                )
                reports.append(report)
                continue
            transactions = ADAPTERS[provider](
                path,
                rows,
                header_index,
                date_profile["convention"],
                plan,
            )
            report.rows_normalized = len(transactions)
            report.rows_skipped = max(report.rows_read - report.rows_normalized, 0)
            report.preview = transaction_preview(transactions)
            if report.rows_read and not transactions:
                raise StatementNormalizationError(
                    f"{path.name}: recognized as {provider}, but no valid transactions could be normalized."
                )
            report.warnings.extend(mark_transaction_issues(transactions))
            if transactions and not any(
                transaction_is_usable(transaction) for transaction in transactions
            ):
                raise StatementNormalizationError(
                    f"{path.name}: transactions are missing required dates, descriptions, amounts, or currencies."
                )
        except Exception as exc:
            message = str(exc)
            report.errors.append(
                message if message.startswith(f"{path.name}:") else f"{path.name}: {message}"
            )
        reports.append(report)
    return reports


def normalize_statement_files(
    paths: list[Path],
    fx_resolver: WeeklyCadFxResolver | None = None,
) -> NormalizationResult:
    result = NormalizationResult()
    for path in paths:
        report = StatementFileReport(source_file=path)
        result.files.append(report)
        if path.suffix.lower() not in SUPPORTED_STATEMENT_EXTENSIONS:
            message = f"{path.name}: unsupported statement format {path.suffix or '(none)'}; use CSV, XLS, or XLSX."
            report.errors.append(message)
            result.errors.append(message)
            continue
        try:
            plan = prepare_statement_file(path)
            apply_import_plan_to_report(report, plan)
            if plan["errors"]:
                raise StatementNormalizationError("; ".join(plan["errors"]))
            rows = plan["rows"]
            provider = plan["provider"]
            header_index = plan["header_index"]
            date_profile = statement_date_profile(path, provider, rows, header_index, plan)
            apply_date_profile_to_report(report, date_profile)
            if report.date_convention_required:
                raise StatementNormalizationError(
                    "dates are ambiguous. Choose day-first or month-first before syncing."
                )
            adapter = ADAPTERS[provider]
            transactions = adapter(path, rows, header_index, date_profile["convention"], plan)
            report.rows_normalized = len(transactions)
            report.rows_skipped = max(report.rows_read - report.rows_normalized, 0)
            report.preview = transaction_preview(transactions)
            if report.rows_read and not transactions:
                raise StatementNormalizationError(
                    f"recognized as {provider}, but no valid transactions could be normalized."
                )
            report.warnings.extend(mark_transaction_issues(transactions))
            if transactions and not any(
                transaction_is_usable(transaction) for transaction in transactions
            ):
                raise StatementNormalizationError(
                    "transactions are missing required dates, descriptions, amounts, or currencies."
                )
            if report.warnings:
                result.warnings.append(
                    f"{path.name}: {len(report.warnings)} normalized rows require review."
                )
            result.transactions.extend(transactions)
        except Exception as exc:
            message = f"{path.name}: {exc}"
            report.errors.append(message)
            result.errors.append(message)

    if result.errors:
        return result

    result.transactions = deduplicate_transactions(result.transactions, result)
    annotate_exact_cad_conversions(result.transactions)
    if fx_resolver:
        apply_weekly_cad_rates(result.transactions, fx_resolver, result)
    assign_cad_completeness(result.transactions)
    result.transactions.sort(
        key=lambda tx: (
            tx.transaction_date or "9999-99-99",
            tx.provider,
            tx.transaction_group_id,
            tx.source_file.name,
            tx.source_row,
        )
    )
    return result


def read_tabular_rows(path: Path) -> list[list[object]]:
    rows, _metadata = read_tabular_rows_with_metadata(path)
    return rows


def read_tabular_rows_with_metadata(path: Path) -> tuple[list[list[object]], dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        text, encoding = decode_csv_text(path.read_bytes())
        delimiter = detect_delimiter(text)
        lines = [unwrap_quoted_delimited_line(line, delimiter) for line in text.splitlines()]
        try:
            rows = [list(row) for row in csv.reader(lines, delimiter=delimiter)]
        except csv.Error as exc:
            raise StatementNormalizationError(f"CSV quoting is malformed: {exc}.") from exc
        return rows, {"encoding": encoding, "delimiter": delimiter}
    if suffix == ".xlsx":
        workbook = load_workbook(path, data_only=True, read_only=True)
        return [list(row) for row in workbook.active.iter_rows(values_only=True)], {
            "encoding": "spreadsheet",
            "delimiter": "",
        }
    if suffix == ".xls":
        try:
            import xlrd
        except Exception as exc:  # pragma: no cover - dependency is declared
            raise StatementNormalizationError("XLS support requires xlrd.") from exc
        workbook = xlrd.open_workbook(str(path))
        sheet = workbook.sheet_by_index(0)
        return [
            [sheet.cell_value(row, col) for col in range(sheet.ncols)] for row in range(sheet.nrows)
        ], {"encoding": "spreadsheet", "delimiter": ""}
    raise StatementNormalizationError(
        f"Unsupported statement format: {suffix or '(none)'}. Use CSV, XLS, or XLSX."
    )


def decode_csv_text(payload: bytes) -> tuple[str, str]:
    if payload.startswith((b"\xff\xfe", b"\xfe\xff")):
        return payload.decode("utf-16"), "utf-16"
    if payload.startswith(b"\xef\xbb\xbf"):
        return payload.decode("utf-8-sig"), "utf-8-sig"
    if b"\x00" in payload[:200]:
        try:
            return payload.decode("utf-16"), "utf-16"
        except UnicodeDecodeError:
            pass
    try:
        return payload.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return payload.decode("cp1252"), "cp1252"


def detect_delimiter(text: str) -> str:
    lines = [line for line in text.splitlines()[:40] if line.strip()]
    if not lines:
        return ","
    candidates = [",", ";", "\t"]
    sample = "\n".join(lines)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        if dialect.delimiter in candidates:
            return dialect.delimiter
    except csv.Error:
        pass
    scores = {}
    for delimiter in candidates:
        counts = [len(next(csv.reader([line], delimiter=delimiter))) for line in lines]
        common = max(set(counts), key=counts.count)
        scores[delimiter] = (sum(count == common and count > 1 for count in counts), common)
    return max(candidates, key=lambda delimiter: scores[delimiter])


def unwrap_quoted_delimited_line(line: str, delimiter: str) -> str:
    """Undo exports that quote an entire semicolon/tab-delimited record."""

    if delimiter == "," or not (line.startswith('"') and line.endswith('"')):
        return line
    try:
        outer_row = next(csv.reader([line], delimiter=","))
    except (csv.Error, StopIteration):
        return line
    if len(outer_row) == 1 and delimiter in outer_row[0]:
        return outer_row[0]
    return line


def detect_provider(rows: list[list[object]]) -> tuple[str, int]:
    for index, row in enumerate(rows[:40]):
        headers = {normalize_key(value) for value in row if str(value or "").strip()}
        if {
            "transaction_date",
            "description",
            "purchase_amount",
            "purchase_currency",
        } <= headers:
            return "standard", index
        if {"date", "date_processed", "description", "card_member", "account", "amount"} <= headers:
            return "amex", index
        if {
            "first_bank_card",
            "transaction_type",
            "date_posted",
            "transaction_amount",
            "description",
        } <= headers:
            return "bmo", index
        if {"date", "description", "category", "debit", "credit"} <= headers and bool(
            headers & {"card_number", "balance"}
        ):
            return "bnc", index
        if {
            "id",
            "status",
            "direction",
            "created_on",
            "source_amount_after_fees",
            "source_currency",
            "target_amount_after_fees",
            "target_currency",
        } <= headers:
            return "wise", index
        if has_generic_signature(headers):
            return "generic", index
    candidate = generic_header_candidate(rows)
    if candidate is not None:
        return "generic", candidate
    raise StatementNormalizationError(
        "Statement columns are not recognized; choose a CSV with a tabular header row."
    )


def generic_header_candidate(rows: list[list[object]]) -> int | None:
    aliases = set().union(*GENERIC_FIELD_ALIASES.values())
    candidates: list[tuple[float, int]] = []
    for index, row in enumerate(rows[:60]):
        values = [clean_text(value) for value in row]
        nonempty = [value for value in values if value]
        if len(nonempty) < 2:
            continue
        normalized = [normalize_key(value) for value in nonempty]
        alias_hits = sum(value in aliases for value in normalized)
        text_like = sum(
            parse_number(value) is None and not numeric_date_parts(value) for value in nonempty
        )
        following = [
            candidate
            for candidate in rows[index + 1 : index + 5]
            if any(clean_text(value) for value in candidate)
        ]
        width_consistency = sum(len(candidate) >= len(row) - 1 for candidate in following)
        score = alias_hits * 10 + text_like + width_consistency * 2 + min(len(nonempty), 10) / 10
        candidates.append((score, index))
    if not candidates:
        return None
    candidates.sort(reverse=True)
    return candidates[0][1]


def has_generic_signature(headers: set[str]) -> bool:
    has_date = bool(
        headers & {"date", "transaction_date", "posted_date", "date_posted", "date_processed"}
    )
    has_description = bool(headers & {"description", "merchant", "details", "transaction", "name"})
    has_amount = bool(
        headers
        & {"amount", "cad_amount", "debit", "credit", "charge", "withdrawal", "charges_adjustments"}
    )
    return has_date and has_description and has_amount


def rows_as_dicts(
    rows: list[list[object]], header_index: int
) -> list[tuple[int, dict[str, object]]]:
    headers = [normalize_key(value) for value in rows[header_index]]
    result: list[tuple[int, dict[str, object]]] = []
    for zero_index, row in enumerate(rows[header_index + 1 :], start=header_index + 1):
        if not any(str(value or "").strip() for value in row):
            continue
        result.append(
            (
                zero_index + 1,
                {
                    headers[index]: row[index] if index < len(row) else None
                    for index in range(len(headers))
                    if headers[index]
                },
            )
        )
    return result


def statement_date_profile(
    path: Path,
    provider: str,
    rows: list[list[object]],
    header_index: int,
    plan: dict | None = None,
) -> dict:
    date_values = statement_date_values(rows, header_index)
    if provider != "generic":
        convention = PROVIDER_DATE_CONVENTIONS[provider]
        return {
            "convention": convention,
            "required": False,
            "samples": parsed_date_samples(date_values, convention),
        }

    inferred = infer_generic_date_convention(date_values)
    settings = load_statement_settings(path.parent.parent)
    saved = settings.get("date_conventions", {}).get(statement_file_fingerprint(path), {})
    saved_convention = saved.get("convention") if isinstance(saved, dict) else None
    if not saved_convention and plan:
        profile = plan.get("saved_profile") or {}
        saved_convention = profile.get("date_convention")
    convention = inferred or (saved_convention if saved_convention in DATE_CONVENTIONS else "")
    ambiguous = has_ambiguous_numeric_dates(date_values)
    required = ambiguous and not convention
    return {
        "convention": convention or ("iso" if not ambiguous else ""),
        "required": required,
        "samples": parsed_date_samples(date_values, convention),
    }


def statement_date_values(rows: list[list[object]], header_index: int) -> list[object]:
    values: list[object] = []
    for _source_row, row in rows_as_dicts(rows, header_index):
        value = first_value(
            row,
            [
                "date",
                "transaction_date",
                "posted_date",
                "date_posted",
                "date_processed",
                "created_on",
            ],
        )
        if value not in (None, ""):
            values.append(value)
    return values


def infer_generic_date_convention(values: list[object]) -> str:
    day_first_proof = False
    month_first_proof = False
    for value in values:
        parts = numeric_date_parts(value)
        if not parts:
            continue
        first, second, _year = parts
        if first > 12 and second <= 12:
            day_first_proof = True
        elif second > 12 and first <= 12:
            month_first_proof = True
    if day_first_proof and month_first_proof:
        raise StatementNormalizationError(
            "numeric dates contain conflicting day-first and month-first evidence."
        )
    if day_first_proof:
        return "day_first"
    if month_first_proof:
        return "month_first"
    return ""


def has_ambiguous_numeric_dates(values: list[object]) -> bool:
    return any(
        parts and parts[0] <= 12 and parts[1] <= 12
        for parts in (numeric_date_parts(value) for value in values)
    )


def numeric_date_parts(value: object) -> tuple[int, int, int] | None:
    if isinstance(value, (datetime, int, float)):
        return None
    match = re.match(
        r"^\s*(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2}|\d{4})(?:\D|$)",
        clean_text(value),
    )
    if not match:
        return None
    first, second, year = (int(part) for part in match.groups())
    if year < 100:
        year += 2000
    return first, second, year


def parsed_date_samples(values: list[object], convention: str, limit: int = 3) -> list[str]:
    samples = []
    for value in values[:limit]:
        raw = clean_text(value)
        parsed = parse_date(value, convention) if convention else None
        samples.append(f"{raw} → {parsed or 'choice required'}")
    return samples


def apply_date_profile_to_report(report: StatementFileReport, profile: dict) -> None:
    report.date_convention = str(profile.get("convention", ""))
    report.date_convention_required = bool(profile.get("required"))
    report.date_samples = [str(value) for value in profile.get("samples", [])]


def set_statement_date_convention(
    trip_dir: Path,
    filename: str,
    convention: str,
) -> None:
    if convention not in DATE_CONVENTIONS:
        raise ValueError("Choose day-first or month-first.")
    if not filename or filename != Path(filename).name:
        raise ValueError("Invalid statement filename.")
    statements_dir = (trip_dir / "card_statements").resolve()
    path = (statements_dir / filename).resolve()
    if path.parent != statements_dir or not path.is_file():
        raise FileNotFoundError("Statement file was not found.")
    rows = read_tabular_rows(path)
    provider, _header_index = detect_provider(rows)
    if provider != "generic":
        raise ValueError(
            "This provider has a fixed date convention and does not need a manual choice."
        )
    settings = load_statement_settings(trip_dir)
    choices = settings.setdefault("date_conventions", {})
    choices[statement_file_fingerprint(path)] = {
        "filename": filename,
        "convention": convention,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    save_statement_settings(trip_dir, settings)


def set_statement_import_profile(
    trip_dir: Path,
    filename: str,
    *,
    header_row: int,
    mapping: dict[str, object],
    sign_convention: str,
    date_convention: str = "",
) -> None:
    if not filename or filename != Path(filename).name:
        raise ValueError("Invalid statement filename.")
    statements_dir = (trip_dir / "card_statements").resolve()
    path = (statements_dir / filename).resolve()
    if path.parent != statements_dir or not path.is_file():
        raise FileNotFoundError("Statement file was not found.")
    if path.suffix.lower() != ".csv":
        raise ValueError("Manual import profiles are available for CSV statements only.")
    rows, _metadata = read_tabular_rows_with_metadata(path)
    header_index = header_row - 1
    if header_index < 0 or header_index >= len(rows):
        raise ValueError("Header row is outside the CSV file.")
    headers = list(rows[header_index])
    selected_mapping = normalize_saved_mapping(mapping, len(headers))
    missing = missing_generic_mapping_fields(selected_mapping)
    if missing:
        raise ValueError("Map " + ", ".join(missing) + " before saving.")
    if sign_convention not in SIGN_CONVENTIONS:
        raise ValueError("Choose whether positive or negative values are purchases.")
    if sign_convention == "debit_credit" and not ({"debit", "credit"} & set(selected_mapping)):
        raise ValueError("Separate debit/credit signs require a debit or credit column.")
    if date_convention and date_convention not in DATE_CONVENTIONS:
        raise ValueError("Choose day-first or month-first.")

    structure = statement_structure_fingerprint(headers)
    settings = load_statement_settings(trip_dir)
    settings["version"] = 2
    profiles = settings.setdefault("import_profiles", {})
    profiles[structure] = {
        "headers": [clean_text(value) for value in headers],
        "mapping": selected_mapping,
        "sign_convention": sign_convention,
        "date_convention": date_convention,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    overrides = settings.setdefault("file_import_overrides", {})
    overrides[statement_file_fingerprint(path)] = {
        "filename": filename,
        "header_index": header_index,
        "structure_fingerprint": structure,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    if date_convention:
        settings.setdefault("date_conventions", {})[statement_file_fingerprint(path)] = {
            "filename": filename,
            "convention": date_convention,
            "updated_at": datetime.now().isoformat(timespec="seconds"),
        }
    save_statement_settings(trip_dir, settings)


def statement_file_fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_statement_settings(trip_dir: Path) -> dict:
    path = trip_dir / STATEMENT_SETTINGS_FILE
    if not path.exists():
        return statement_settings_defaults()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return statement_settings_defaults()
    if not isinstance(data, dict):
        return statement_settings_defaults()
    data.setdefault("version", 1)
    data.setdefault("date_conventions", {})
    data.setdefault("import_profiles", {})
    data.setdefault("file_import_overrides", {})
    return data


def statement_settings_defaults() -> dict:
    return {
        "version": 2,
        "date_conventions": {},
        "import_profiles": {},
        "file_import_overrides": {},
    }


def save_statement_settings(trip_dir: Path, settings: dict) -> None:
    path = trip_dir / STATEMENT_SETTINGS_FILE
    write_json_atomic(path, settings)


def normalize_amex(
    path: Path,
    rows: list[list[object]],
    header_index: int,
    date_convention: str,
    _plan: dict | None = None,
) -> list[NormalizedTransaction]:
    transactions: list[NormalizedTransaction] = []
    for source_row, row in rows_as_dicts(rows, header_index):
        amount = parse_number(row.get("amount"))
        if amount is None:
            continue
        description = clean_text(row.get("description"))
        transaction_type = classify_credit(description) if amount < 0 else "purchase"
        match_eligible = transaction_type in {"purchase", "refund"}
        signed_amount = abs(amount) if transaction_type == "purchase" else -abs(amount)
        foreign_amount, foreign_currency = parse_foreign_value(
            first_value(row, ["foreign_spend_amount", "foreign_amount", "original_amount"])
        )
        if foreign_amount is not None:
            foreign_amount = (
                abs(foreign_amount) if transaction_type == "purchase" else -abs(foreign_amount)
            )
        date = parse_date(row.get("date"), date_convention)
        posted = parse_date(row.get("date_processed"), date_convention)
        account = mask_account(row.get("account"))
        group_id = source_group_id("amex", path, source_row)
        transactions.append(
            NormalizedTransaction(
                source_file=path,
                source_row=source_row,
                provider="amex",
                transaction_group_id=group_id,
                funding_leg_id=f"{group_id}:1",
                transaction_date=date,
                posted_date=posted,
                account_label=account,
                cardholder=clean_text(row.get("card_member")),
                description=description,
                transaction_type=transaction_type,
                status="completed",
                direction="out" if signed_amount >= 0 else "in",
                match_eligible=match_eligible,
                purchase_amount=(foreign_amount if foreign_amount is not None else signed_amount)
                if match_eligible
                else None,
                purchase_currency=(foreign_currency or "CAD") if match_eligible else None,
                settlement_amount=signed_amount,
                settlement_currency="CAD",
                cad_amount=signed_amount,
                cad_completeness="complete" if match_eligible else "not_applicable",
                match_status="unmatched" if match_eligible else "ignored",
            )
        )
    return transactions


def normalize_bmo(
    path: Path,
    rows: list[list[object]],
    header_index: int,
    date_convention: str,
    _plan: dict | None = None,
) -> list[NormalizedTransaction]:
    transactions: list[NormalizedTransaction] = []
    for source_row, row in rows_as_dicts(rows, header_index):
        amount = parse_number(row.get("transaction_amount"))
        if amount is None:
            continue
        description = clean_text(row.get("description"))
        bank_type = clean_text(row.get("transaction_type")).upper()
        if bank_type == "DEBIT":
            transaction_type = "purchase"
            signed_amount = abs(amount)
            match_eligible = True
        else:
            transaction_type = classify_credit(description)
            signed_amount = -abs(amount)
            match_eligible = transaction_type == "refund"
        foreign_amount, foreign_currency = parse_foreign_value(
            first_value(row, ["foreign_spend_amount", "foreign_amount", "original_amount"])
        )
        if foreign_amount is not None:
            foreign_amount = (
                abs(foreign_amount) if transaction_type == "purchase" else -abs(foreign_amount)
            )
        account = mask_account(row.get("first_bank_card"))
        group_id = source_group_id("bmo", path, source_row)
        transactions.append(
            NormalizedTransaction(
                source_file=path,
                source_row=source_row,
                provider="bmo",
                transaction_group_id=group_id,
                funding_leg_id=f"{group_id}:1",
                transaction_date=parse_date(row.get("date_posted"), date_convention),
                posted_date=parse_date(row.get("date_posted"), date_convention),
                account_label=account,
                description=description,
                transaction_type=transaction_type,
                status="completed",
                direction="out" if signed_amount >= 0 else "in",
                match_eligible=match_eligible,
                purchase_amount=(foreign_amount if foreign_amount is not None else signed_amount)
                if match_eligible
                else None,
                purchase_currency=(foreign_currency or "CAD") if match_eligible else None,
                settlement_amount=signed_amount,
                settlement_currency="CAD",
                cad_amount=signed_amount,
                cad_completeness="complete" if match_eligible else "not_applicable",
                match_status="unmatched" if match_eligible else "ignored",
            )
        )
    return transactions


def normalize_bnc(
    path: Path,
    rows: list[list[object]],
    header_index: int,
    date_convention: str,
    _plan: dict | None = None,
) -> list[NormalizedTransaction]:
    transactions: list[NormalizedTransaction] = []
    for source_row, row in rows_as_dicts(rows, header_index):
        debit = parse_number(row.get("debit")) or 0.0
        credit = parse_number(row.get("credit")) or 0.0
        if debit == 0 and credit == 0:
            continue
        description = clean_text(row.get("description"))
        if debit > 0:
            transaction_type = classify_bnc_debit(
                description,
                clean_text(row.get("category")),
            )
            signed_amount = abs(debit)
            match_eligible = transaction_type == "purchase"
        else:
            transaction_type = classify_credit(description)
            signed_amount = -abs(credit)
            match_eligible = transaction_type == "refund"
        foreign_amount, foreign_currency = parse_foreign_value(
            first_value(row, ["foreign_spend_amount", "foreign_amount", "original_amount"])
        )
        if foreign_amount is not None:
            foreign_amount = (
                abs(foreign_amount) if transaction_type == "purchase" else -abs(foreign_amount)
            )
        group_id = source_group_id("bnc", path, source_row)
        transactions.append(
            NormalizedTransaction(
                source_file=path,
                source_row=source_row,
                provider="bnc",
                transaction_group_id=group_id,
                funding_leg_id=f"{group_id}:1",
                transaction_date=parse_date(row.get("date"), date_convention),
                posted_date=parse_date(row.get("date"), date_convention),
                account_label=(
                    mask_account(row.get("card_number"))
                    or ("BNC debit" if "balance" in row else "")
                ),
                description=description,
                category=clean_text(row.get("category")),
                transaction_type=transaction_type,
                status="completed",
                direction="out" if signed_amount >= 0 else "in",
                match_eligible=match_eligible,
                purchase_amount=(foreign_amount if foreign_amount is not None else signed_amount)
                if match_eligible
                else None,
                purchase_currency=(foreign_currency or "CAD") if match_eligible else None,
                settlement_amount=signed_amount,
                settlement_currency="CAD",
                cad_amount=signed_amount,
                cad_completeness="complete" if match_eligible else "not_applicable",
                match_status="unmatched" if match_eligible else "ignored",
            )
        )
    return transactions


def normalize_wise(
    path: Path,
    rows: list[list[object]],
    header_index: int,
    date_convention: str,
    _plan: dict | None = None,
) -> list[NormalizedTransaction]:
    transactions: list[NormalizedTransaction] = []
    for source_row, row in rows_as_dicts(rows, header_index):
        raw_id = clean_text(row.get("id"))
        if not raw_id.startswith("CARD_TRANSACTION-"):
            continue
        source_amount = parse_number(row.get("source_amount_after_fees"))
        target_amount = parse_number(row.get("target_amount_after_fees"))
        source_fee = parse_number(row.get("source_fee_amount")) or 0.0
        if source_amount is None and target_amount is None:
            continue
        status = clean_text(row.get("status")).lower()
        direction = clean_text(row.get("direction")).lower()
        is_refund = status == "refunded" or direction == "in"
        sign = -1.0 if is_refund else 1.0
        category = clean_text(row.get("category"))
        is_cash = category.lower() == "cash"
        if is_refund:
            transaction_type = "refund"
        elif is_cash:
            transaction_type = "cash"
        else:
            transaction_type = "purchase"
        source_currency = currency_code(row.get("source_currency"))
        target_currency = currency_code(row.get("target_currency"))
        settlement_amount = (
            sign * ((source_amount or 0.0) + source_fee) if source_amount is not None else None
        )
        purchase_amount = sign * target_amount if target_amount is not None else None
        has_transaction_amount = any(
            amount not in (None, 0, -0.0) for amount in (purchase_amount, settlement_amount)
        )
        match_eligible = (
            status != "cancelled"
            and transaction_type in {"purchase", "refund"}
            and has_transaction_amount
        )
        if source_currency == "CAD" and settlement_amount is not None:
            cad_amount = settlement_amount
        elif target_currency == "CAD" and purchase_amount is not None:
            cad_amount = purchase_amount
        else:
            cad_amount = None
        fingerprint = stable_hash(
            raw_id,
            source_currency,
            source_amount,
            source_fee,
            target_currency,
            target_amount,
            status,
            direction,
        )
        review_note = ""
        if status == "cancelled":
            review_note = "Cancelled Wise activity is retained for audit."
        elif status == "refunded" and direction != "in":
            review_note = (
                "Wise status is REFUNDED; amounts are normalized as a negative refund "
                "even though the exported direction is not IN."
            )
        transactions.append(
            NormalizedTransaction(
                source_file=path,
                source_row=source_row,
                provider="wise",
                transaction_group_id=raw_id,
                funding_leg_id=f"{raw_id}:{fingerprint}",
                transaction_date=parse_date(row.get("created_on"), date_convention),
                posted_date=parse_date(row.get("finished_on"), date_convention),
                account_label="Wise multi-currency",
                cardholder=clean_text(row.get("source_name")),
                description=clean_text(row.get("target_name")),
                category=category,
                transaction_type=transaction_type,
                status=status,
                direction=direction,
                match_eligible=match_eligible,
                purchase_amount=purchase_amount,
                purchase_currency=target_currency,
                settlement_amount=settlement_amount,
                settlement_currency=source_currency,
                cad_amount=cad_amount,
                cad_completeness="incomplete" if match_eligible else "not_applicable",
                match_status="unmatched" if match_eligible else "ignored",
                normalization_status="review"
                if status not in {"completed", "refunded", "cancelled"}
                else "ok",
                review_note=review_note,
            )
        )
    return transactions


def normalize_generic(
    path: Path,
    rows: list[list[object]],
    header_index: int,
    date_convention: str,
    plan: dict | None = None,
) -> list[NormalizedTransaction]:
    transactions: list[NormalizedTransaction] = []
    sign_convention = str((plan or {}).get("sign_convention") or "positive_purchase")
    for source_row, row in rows_as_dicts(rows, header_index):
        description = clean_text(
            first_value(row, ["description", "merchant", "details", "transaction", "name"])
        )
        raw_date = first_value(
            row, ["date", "transaction_date", "posted_date", "date_posted", "date_processed"]
        )
        if not parse_date(raw_date, date_convention) and normalize_key(description) in {
            "balance",
            "closing_balance",
            "ending_balance",
            "opening_balance",
            "statement_total",
            "total",
            "totals",
        }:
            continue
        debit = parse_number(first_value(row, ["debit", "charge", "withdrawal"]))
        credit = parse_number(first_value(row, ["credit"]))
        amount = parse_number(
            first_value(row, ["settlement_amount", "amount", "cad_amount", "charges_adjustments"])
        )
        raw_purchase_value = first_value(
            row, ["purchase_amount", "foreign_amount", "original_amount"]
        )
        purchase_amount = parse_number(raw_purchase_value)
        if amount is None:
            amount = purchase_amount
        special_credit_type = classify_credit(description)
        special_credit_type = special_credit_type if special_credit_type != "refund" else None
        raw_type = normalize_key(first_value(row, ["transaction_type", "type"]))
        explicit_types = {
            "purchase": "purchase",
            "debit": "purchase",
            "charge": "purchase",
            "refund": "refund",
            "credit": special_credit_type or "refund",
            "payment": "payment",
            "cashback": "cashback",
            "deposit": "deposit",
            "fee": "fee",
            "cash": "cash",
            "withdrawal": "cash",
            "transfer": "transfer",
        }
        if debit not in (None, 0):
            transaction_type = "purchase"
            signed_amount = abs(debit)
        elif credit not in (None, 0):
            transaction_type = special_credit_type or "refund"
            signed_amount = -abs(credit)
        elif raw_type in explicit_types and amount is not None:
            transaction_type = explicit_types[raw_type]
            signed_amount = signed_amount_for_type(amount, transaction_type)
        elif amount is not None:
            if special_credit_type:
                transaction_type = special_credit_type
                signed_amount = -abs(amount)
            else:
                purchase_sign = (
                    amount >= 0 if sign_convention == "positive_purchase" else amount < 0
                )
                transaction_type = "purchase" if purchase_sign else "refund"
                signed_amount = abs(amount) if purchase_sign else -abs(amount)
        else:
            continue
        date = parse_date(raw_date, date_convention)
        settlement_currency = (
            currency_code(first_value(row, ["settlement_currency", "currency"])) or "CAD"
        )
        purchase_currency = currency_code(
            first_value(row, ["purchase_currency", "foreign_currency", "original_currency"])
        )
        foreign_amount, embedded_foreign_currency = parse_foreign_value(
            first_value(
                row,
                ["foreign_spend_amount", "foreign_amount", "original_amount"],
            )
            or raw_purchase_value
        )
        if purchase_amount is None:
            purchase_amount = foreign_amount
        if purchase_currency is None:
            purchase_currency = embedded_foreign_currency or settlement_currency
        match_eligible = transaction_type in {"purchase", "refund"}
        if purchase_amount is not None:
            purchase_amount = (
                abs(purchase_amount) if transaction_type == "purchase" else -abs(purchase_amount)
            )
        group_id = source_group_id("generic", path, source_row)
        transactions.append(
            NormalizedTransaction(
                source_file=path,
                source_row=source_row,
                provider="generic",
                transaction_group_id=group_id,
                funding_leg_id=f"{group_id}:1",
                transaction_date=date,
                posted_date=parse_date(
                    first_value(row, ["posted_date", "date_posted", "date_processed"]),
                    date_convention,
                )
                or date,
                account_label=mask_account(
                    first_value(row, ["account", "account_number", "card_number"])
                ),
                cardholder=clean_text(first_value(row, ["cardholder", "card_member"])),
                description=description,
                transaction_type=transaction_type,
                status="completed",
                direction="out" if signed_amount >= 0 else "in",
                match_eligible=match_eligible,
                purchase_amount=(purchase_amount if purchase_amount is not None else signed_amount)
                if match_eligible
                else None,
                purchase_currency=purchase_currency if match_eligible else None,
                settlement_amount=signed_amount,
                settlement_currency=settlement_currency,
                cad_amount=signed_amount if settlement_currency == "CAD" else None,
                cad_completeness="complete"
                if match_eligible and settlement_currency == "CAD"
                else ("incomplete" if match_eligible else "not_applicable"),
                match_status="unmatched" if match_eligible else "ignored",
            )
        )
    return transactions


def normalize_standard(
    path: Path,
    rows: list[list[object]],
    header_index: int,
    date_convention: str,
    _plan: dict | None = None,
) -> list[NormalizedTransaction]:
    """Normalize the documented provider-neutral CSV template."""

    transactions: list[NormalizedTransaction] = []
    for source_row, row in rows_as_dicts(rows, header_index):
        raw_amount = parse_number(row.get("purchase_amount"))
        raw_cad_amount = parse_number(row.get("cad_amount"))
        if raw_amount is None or (raw_amount == 0 and raw_cad_amount in (None, 0)):
            continue
        raw_type = normalize_key(row.get("transaction_type"))
        invalid_type = bool(raw_type and raw_type not in STANDARD_TRANSACTION_TYPES)
        if invalid_type:
            transaction_type = "other"
        elif raw_type:
            transaction_type = raw_type
        else:
            transaction_type = "refund" if raw_amount < 0 else "purchase"

        if transaction_type == "purchase":
            signed_amount = abs(raw_amount)
        elif transaction_type in STANDARD_INCOMING_TYPES:
            signed_amount = -abs(raw_amount)
        else:
            signed_amount = raw_amount
        if raw_cad_amount is None:
            signed_cad_amount = None
        elif transaction_type == "purchase":
            signed_cad_amount = abs(raw_cad_amount)
        elif transaction_type in STANDARD_INCOMING_TYPES:
            signed_cad_amount = -abs(raw_cad_amount)
        else:
            signed_cad_amount = raw_cad_amount

        purchase_currency = currency_code(row.get("purchase_currency"))
        if signed_cad_amount is None and purchase_currency == "CAD":
            signed_cad_amount = signed_amount
        match_eligible = transaction_type in {"purchase", "refund"}
        group_id = source_group_id("standard", path, source_row)
        normalization_status = "review" if invalid_type else "ok"
        review_note = (
            f"Unsupported standard CSV transaction_type '{clean_text(row.get('transaction_type'))}'."
            if invalid_type
            else ""
        )
        transactions.append(
            NormalizedTransaction(
                source_file=path,
                source_row=source_row,
                provider="standard",
                transaction_group_id=group_id,
                funding_leg_id=f"{group_id}:1",
                transaction_date=parse_date(row.get("transaction_date"), date_convention),
                posted_date=(
                    parse_date(row.get("posted_date"), date_convention)
                    or parse_date(row.get("transaction_date"), date_convention)
                ),
                account_label=clean_text(row.get("account")),
                cardholder=clean_text(row.get("cardholder")),
                description=clean_text(row.get("description")),
                category=clean_text(row.get("category")),
                transaction_type=transaction_type,
                status="completed",
                direction="out" if signed_amount >= 0 else "in",
                match_eligible=match_eligible,
                purchase_amount=signed_amount if match_eligible else None,
                purchase_currency=purchase_currency if match_eligible else None,
                settlement_amount=(
                    signed_cad_amount if signed_cad_amount is not None else signed_amount
                ),
                settlement_currency=("CAD" if signed_cad_amount is not None else purchase_currency),
                cad_amount=signed_cad_amount,
                cad_completeness=(
                    "complete"
                    if match_eligible and signed_cad_amount is not None
                    else "incomplete"
                    if match_eligible
                    else "not_applicable"
                ),
                match_status="unmatched" if match_eligible else "ignored",
                normalization_status=normalization_status,
                review_note=review_note,
            )
        )
    return transactions


ADAPTERS: dict[
    str,
    Callable[[Path, list[list[object]], int, str, dict | None], list[NormalizedTransaction]],
] = {
    "standard": normalize_standard,
    "amex": normalize_amex,
    "bmo": normalize_bmo,
    "bnc": normalize_bnc,
    "wise": normalize_wise,
    "generic": normalize_generic,
}


def deduplicate_transactions(
    transactions: list[NormalizedTransaction], result: NormalizationResult
) -> list[NormalizedTransaction]:
    deduplicated: list[NormalizedTransaction] = []
    strong_seen: set[tuple[object, ...]] = set()
    weak_groups: dict[tuple[object, ...], list[NormalizedTransaction]] = defaultdict(list)
    for transaction in transactions:
        strong_key = (
            transaction.provider,
            transaction.transaction_group_id,
            transaction.funding_leg_id,
            transaction.transaction_date,
            transaction.settlement_amount,
            transaction.settlement_currency,
            transaction.purchase_amount,
            transaction.purchase_currency,
        )
        if transaction.provider == "wise":
            if strong_key in strong_seen:
                result.warnings.append(
                    f"Removed duplicate statement leg {transaction.funding_leg_id} from {transaction.source_file.name}."
                )
                report = next(
                    (item for item in result.files if item.source_file == transaction.source_file),
                    None,
                )
                if report:
                    report.rows_skipped += 1
                    report.rows_normalized = max(report.rows_normalized - 1, 0)
                continue
            strong_seen.add(strong_key)
        else:
            weak_key = (
                transaction.provider,
                transaction.account_label,
                transaction.transaction_date,
                normalize_description(transaction.description),
                transaction.settlement_amount,
                transaction.settlement_currency,
            )
            weak_groups[weak_key].append(transaction)
        deduplicated.append(transaction)

    for duplicate_group in weak_groups.values():
        if len(duplicate_group) < 2:
            continue
        evidence = ", ".join(
            f"{transaction.source_file.name} row {transaction.source_row}"
            for transaction in duplicate_group
        )
        warning = f"Possible duplicate statement transactions: {evidence}."
        result.warnings.append(warning)
        for transaction in duplicate_group:
            transaction.normalization_status = "possible_duplicate"
            transaction.review_note = append_note(
                transaction.review_note,
                "Possible duplicate: same provider, account, date, description, amount, and currency.",
            )
            report = next(
                (item for item in result.files if item.source_file == transaction.source_file), None
            )
            if report and warning not in report.warnings:
                report.warnings.append(warning)
    return deduplicated


def assign_cad_completeness(transactions: list[NormalizedTransaction]) -> None:
    groups: dict[str, list[NormalizedTransaction]] = defaultdict(list)
    for transaction in transactions:
        groups[transaction.transaction_group_id].append(transaction)
    for legs in groups.values():
        eligible = [leg for leg in legs if leg.match_eligible]
        if not eligible:
            status = "not_applicable"
        elif all(leg.cad_amount is not None for leg in eligible):
            status = "complete"
        elif any(leg.cad_amount is not None for leg in eligible):
            status = "partial"
        else:
            status = "incomplete"
        for leg in legs:
            leg.cad_completeness = status
            if leg.match_eligible and status in {"partial", "incomplete"}:
                leg.review_note = append_note(
                    leg.review_note,
                    "No complete CAD amount is available; use a manual CAD override after matching.",
                )


def annotate_exact_cad_conversions(transactions: list[NormalizedTransaction]) -> None:
    """Record why a transaction already has an exact CAD amount."""

    for transaction in transactions:
        if transaction.cad_amount is None or not transaction.match_eligible:
            continue
        basis_amount = transaction.purchase_amount
        basis_currency = transaction.purchase_currency
        if basis_amount not in (None, 0) and basis_currency:
            transaction.cad_conversion_rate = round(
                abs(float(transaction.cad_amount)) / abs(float(basis_amount)),
                10,
            )
            transaction.cad_conversion_route = (
                "CAD→CAD" if basis_currency == "CAD" else f"{basis_currency}→CAD"
            )
        elif transaction.settlement_currency == "CAD":
            transaction.cad_conversion_rate = 1.0
            transaction.cad_conversion_route = "CAD→CAD"
        if transaction.purchase_currency == "CAD":
            transaction.cad_conversion_method = "statement_purchase_cad"
            transaction.cad_conversion_source = (
                f"{transaction.provider.upper()} statement purchase amount in CAD"
            )
        elif transaction.settlement_currency == "CAD":
            transaction.cad_conversion_method = "statement_exact_cad_settlement"
            transaction.cad_conversion_source = (
                f"{transaction.provider.upper()} statement exact CAD settlement"
            )


def apply_weekly_cad_rates(
    transactions: list[NormalizedTransaction],
    fx_resolver: WeeklyCadFxResolver,
    result: NormalizationResult,
) -> None:
    """Enrich any provider's unresolved foreign purchase with a weekly CAD rate."""

    warned: set[tuple[str, str, str]] = set()
    reports_by_path = {report.source_file: report for report in result.files}
    for transaction in transactions:
        if not transaction.match_eligible or transaction.cad_amount is not None:
            continue
        if (
            transaction.purchase_amount is None
            or not transaction.purchase_currency
            or not transaction.transaction_date
        ):
            continue
        try:
            rate = fx_resolver.resolve(
                transaction.purchase_currency,
                transaction.transaction_date,
            )
        except FxRateUnavailable as exc:
            key = (
                transaction.purchase_currency,
                transaction.transaction_date,
                str(exc),
            )
            note = f"Weekly CAD conversion unavailable: {exc}"
            transaction.review_note = append_note(transaction.review_note, note)
            if key not in warned:
                warned.add(key)
                warning = f"{transaction.source_file.name} row {transaction.source_row}: {note}"
                result.warnings.append(warning)
                report = reports_by_path.get(transaction.source_file)
                if report:
                    report.warnings.append(warning)
            continue
        transaction.cad_amount = round(
            float(transaction.purchase_amount) * rate.cad_per_unit,
            2,
        )
        transaction.cad_conversion_rate = rate.cad_per_unit
        transaction.cad_conversion_week_start = rate.week_start
        transaction.cad_conversion_week_end = rate.week_end
        transaction.cad_conversion_method = rate.method
        transaction.cad_conversion_route = rate.route
        transaction.cad_conversion_source = rate.source
        transaction.cad_conversion_source_urls = list(rate.source_urls)


def normalize_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").lstrip("\ufeff").strip().lower()).strip("_")


def clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def currency_code(value: object) -> str | None:
    text = clean_text(value).upper()
    match = re.search(r"\b([A-Z]{3})\b", text)
    if match:
        return match.group(1)
    symbols = {
        "C$": "CAD",
        "CA$": "CAD",
        "US$": "USD",
        "A$": "AUD",
        "AU$": "AUD",
        "€": "EUR",
        "£": "GBP",
        "¥": "JPY",
        "$": "CAD",
    }
    return symbols.get(text)


def parse_number(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value).replace("\u00a0", "").replace("\u202f", "")
    negative = text.startswith("(") and text.endswith(")")
    if re.search(r"\bCR\s*$", text, flags=re.IGNORECASE):
        negative = True
    text = re.sub(r"\b(?:CR|DR)\s*$", "", text, flags=re.IGNORECASE)
    text = text.strip("()").strip()
    text = re.sub(r"[^0-9,.'+\-]", "", text).replace("'", "")
    if not text or text in {"+", "-"}:
        return None
    comma_positions = [match.start() for match in re.finditer(",", text)]
    dot_positions = [match.start() for match in re.finditer(r"\.", text)]
    if comma_positions and dot_positions:
        decimal = "," if comma_positions[-1] > dot_positions[-1] else "."
        other = "." if decimal == "," else ","
        decimals = len(text) - text.rfind(decimal) - 1
        if decimals in {1, 2}:
            text = text.replace(other, "").replace(decimal, ".")
        else:
            text = text.replace(",", "").replace(".", "")
    elif comma_positions:
        decimals = len(text) - comma_positions[-1] - 1
        if decimals in {1, 2}:
            text = (
                text.replace(",", ".")
                if len(comma_positions) == 1
                else text.replace(",", "", len(comma_positions) - 1).replace(",", ".")
            )
        else:
            text = text.replace(",", "")
    elif len(dot_positions) > 1:
        decimals = len(text) - dot_positions[-1] - 1
        if decimals in {1, 2}:
            head, tail = text.rsplit(".", 1)
            text = head.replace(".", "") + "." + tail
        else:
            text = text.replace(".", "")
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


def parse_date(value: object, convention: str = "") -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and value > 20000:
        base = datetime(1899, 12, 30)
        return (base + timedelta(days=int(value))).strftime("%Y-%m-%d")
    text = clean_text(value)
    for candidate in (text, text[:10], text[:11]):
        for format_string in (
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%Y.%m.%d",
            "%Y%m%d",
            "%d-%b-%y",
            "%d-%B-%y",
            "%d %b %y",
            "%d %B %y",
            "%d %b %Y",
            "%d %B %Y",
            "%d %b. %Y",
            "%d.%m.%Y",
            "%d.%m.%y",
            "%Y-%m-%d %H:%M:%S",
        ):
            try:
                return datetime.strptime(candidate, format_string).strftime("%Y-%m-%d")
            except ValueError:
                continue
    parts = numeric_date_parts(text)
    if parts:
        first, second, year = parts
        selected = convention
        if first > 12 and second <= 12:
            selected = "day_first"
        elif second > 12 and first <= 12:
            selected = "month_first"
        if selected == "day_first":
            day, month = first, second
        elif selected == "month_first":
            month, day = first, second
        else:
            return None
        try:
            return datetime(year, month, day).strftime("%Y-%m-%d")
        except ValueError:
            return None
    return None


def mask_account(value: object) -> str:
    text = clean_text(value).strip("'")
    digits = re.sub(r"\D", "", text)
    return f"••••{digits[-4:]}" if digits else ""


def source_group_id(provider: str, path: Path, source_row: int) -> str:
    return f"{provider}:{path.stem}:{source_row}"


def stable_hash(*values: object) -> str:
    payload = "|".join(str(value or "") for value in values)
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()[:12]


def first_value(row: dict[str, object], keys: list[str]) -> object | None:
    for key in keys:
        if row.get(key) not in (None, ""):
            return row[key]
    return None


def parse_foreign_value(value: object) -> tuple[float | None, str | None]:
    if value in (None, ""):
        return None, None
    text = clean_text(value).upper().replace(",", "")
    match = re.search(r"(-?\d+(?:\.\d+)?)\s+([A-Z ]+)", text)
    if not match:
        return None, None
    currency_name = match.group(2).strip()
    currencies = {
        "AUSTRALIAN DOLLAR": "AUD",
        "CANADIAN DOLLAR": "CAD",
        "US DOLLAR": "USD",
        "INDONESIAN RUPIAH": "IDR",
        "VIETNAMESE DONG": "VND",
        "QATARI RIYAL": "QAR",
        "HONG KONG DOLLAR": "HKD",
        "SWISS FRANC": "CHF",
    }
    return float(match.group(1)), currencies.get(currency_name, currency_name[:3])


def normalize_description(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def transaction_is_usable(transaction: NormalizedTransaction) -> bool:
    return bool(
        transaction.transaction_date
        and transaction.description
        and transaction.settlement_amount is not None
        and transaction.settlement_currency
        and (
            not transaction.match_eligible
            or (transaction.purchase_amount is not None and transaction.purchase_currency)
        )
    )


def mark_transaction_issues(transactions: list[NormalizedTransaction]) -> list[str]:
    warnings: list[str] = []
    for transaction in transactions:
        missing: list[str] = []
        if not transaction.transaction_date:
            missing.append("transaction date")
        if not transaction.description:
            missing.append("description")
        if transaction.settlement_amount is None:
            missing.append("settlement amount")
        if not transaction.settlement_currency:
            missing.append("settlement currency")
        if transaction.match_eligible and transaction.purchase_amount is None:
            missing.append("purchase amount")
        if transaction.match_eligible and not transaction.purchase_currency:
            missing.append("purchase currency")
        if not missing:
            continue
        transaction.normalization_status = "review"
        note = f"Normalization review: missing {', '.join(missing)}."
        transaction.review_note = append_note(transaction.review_note, note)
        warnings.append(f"row {transaction.source_row}: {note}")
    return warnings


def classify_credit(description: str) -> str:
    if re.search(r"cashback|cash back|reward|rebate", description, re.I):
        return "cashback"
    if re.search(r"payment", description, re.I):
        return "payment"
    if re.search(r"deposit", description, re.I):
        return "deposit"
    if re.search(r"transfer", description, re.I):
        return "transfer"
    return "refund"


def signed_amount_for_type(amount: float, transaction_type: str) -> float:
    if transaction_type == "purchase" or transaction_type in {"fee", "cash"}:
        return abs(amount)
    if transaction_type in STANDARD_INCOMING_TYPES or transaction_type == "payment":
        return -abs(amount)
    return amount


def classify_bnc_debit(description: str, category: str) -> str:
    description_text = description.lower()
    category_text = category.lower()
    evidence = f"{description_text} {category_text}"
    if re.search(r"credit card payment|mastercard payment|payment received", evidence):
        return "payment"
    if category_text == "cash" or re.search(r"\b(?:abm|atm)\s+withdrawal\b", description_text):
        return "cash"
    if category_text in {"fee", "fees", "bank fee"} or re.search(
        r"\b(?:monthly|transaction|service|card purchase)\s+fees?\b",
        description_text,
    ):
        return "fee"
    return "purchase"


def append_note(existing: str, note: str) -> str:
    return f"{existing} {note}".strip() if existing else note
